from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_incremental_highlight import (  # noqa: E402
    _bare_whitespace_run_cells,
    _harden_workbook_whitespace_runs,
    _next_output_path,
    _term_spans,
    _verify_saved_workbook,
    _workbook_text_snapshot,
    render_cell,
)


def _runs(value):
    return [
        (part.text, part.font.color.rgb if isinstance(part, TextBlock) and part.font.color else None)
        for part in value
    ]


class IncrementalHighlightTests(unittest.TestCase):
    def test_blue_glossary_wins_over_red_change(self):
        value = render_cell(
            "Use SmartThings now", base_font=None,
            red_spans=[[4, 15]], blue_spans=[[4, 15]],
        )
        self.assertIsInstance(value, CellRichText)
        self.assertEqual(str(value), "Use SmartThings now")
        self.assertEqual(_runs(value), [("Use ", None), ("SmartThings", "000000FF"), (" now", None)])

    def test_renderer_accepts_explicit_opaque_blue(self):
        value = render_cell(
            "Use SmartThings", base_font=None, red_spans=[], blue_spans=[[4, 15]], blue_color="FF0000FF",
        )
        self.assertEqual(_runs(value), [("Use ", None), ("SmartThings", "FF0000FF")])

    def test_non_overlapping_change_remains_red(self):
        value = render_cell(
            "Use SmartThings now", base_font=None,
            red_spans=[[16, 19]], blue_spans=[[4, 15]],
        )
        self.assertEqual(_runs(value), [("Use ", None), ("SmartThings ", "000000FF"), ("now", "00FF0000")])

    def test_terms_are_case_insensitive_and_unmatched_are_reported(self):
        spans, missing = _term_spans("Use smartthings", ["SmartThings", "Galaxy"])
        self.assertEqual(spans, [[4, 15]])
        self.assertEqual(missing, ["Galaxy"])

    def test_output_never_reuses_existing_file_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook = Path(tmp) / "story.xlsx"
            workbook.touch()
            first = _next_output_path(workbook, "20260803_120000")
            first.touch()
            second = _next_output_path(workbook, "20260803_120000")
            self.assertNotEqual(first, second)
            self.assertTrue(second.name.endswith("_2.xlsx"))

    def test_composed_runs_survive_xlsx_round_trip_with_spaces(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            wb = openpyxl.Workbook()
            cell = wb.active["C7"]
            cell.value = render_cell(
                "Use SmartThings now", base_font=cell.font,
                red_spans=[[16, 19]], blue_spans=[[4, 15]],
            )
            wb.save(path)
            restored = openpyxl.load_workbook(path, rich_text=True)
            value = restored.active["C7"].value
            self.assertEqual(str(value), "Use SmartThings now")
            self.assertEqual(_runs(value), [("Use ", None), ("SmartThings ", "000000FF"), ("now", "00FF0000")])

    def test_workbook_wide_whitespace_hardening_protects_unmodified_rich_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "AE(아랍에메리트)"
            base_font = InlineFont()
            ws["C7"].value = CellRichText([
                TextBlock(text="الانتقال", font=base_font),
                TextBlock(text=" ", font=base_font),
                TextBlock(text="إلى", font=base_font),
                TextBlock(text=" ", font=base_font),
                TextBlock(text="Settings", font=base_font),
            ])
            ws["C8"].value = render_cell(
                "Use SmartThings now", base_font=ws["C8"].font,
                red_spans=[[16, 19]], blue_spans=[[4, 15]],
            )
            expected = _workbook_text_snapshot(wb)

            self.assertEqual(_bare_whitespace_run_cells(wb), ["AE(아랍에메리트)!C7"])
            self.assertEqual(_harden_workbook_whitespace_runs(wb), ["AE(아랍에메리트)!C7"])
            self.assertEqual(_workbook_text_snapshot(wb), expected)
            self.assertEqual(_bare_whitespace_run_cells(wb), [])
            self.assertEqual(_runs(ws["C8"].value), [("Use ", None), ("SmartThings ", "000000FF"), ("now", "00FF0000")])

            wb.save(path)
            _verify_saved_workbook(path, expected)
            restored = openpyxl.load_workbook(path, rich_text=True)
            self.assertEqual(str(restored["AE(아랍에메리트)"]["C7"].value), "الانتقال إلى Settings")
            self.assertEqual(_bare_whitespace_run_cells(restored), [])

    def test_whitespace_only_rich_text_cell_is_not_a_hard_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "space.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet"
            ws["A1"].value = CellRichText([TextBlock(text=" ", font=InlineFont())])
            expected = _workbook_text_snapshot(wb)

            self.assertEqual(_harden_workbook_whitespace_runs(wb), [])
            self.assertEqual(_bare_whitespace_run_cells(wb), [])
            self.assertEqual(_workbook_text_snapshot(wb), expected)
            wb.save(path)
            _verify_saved_workbook(path, expected)


if __name__ == "__main__":
    unittest.main()
