from __future__ import annotations

import sys
import tempfile
import unittest
import subprocess
import json
from pathlib import Path
from unittest.mock import AsyncMock, patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from agent_sheet_review import build_packet  # noqa: E402


class AgentSheetReviewTests(unittest.IsolatedAsyncioTestCase):
    async def test_builds_one_sheet_packet_without_mutating_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "story.xlsx"
            wb = openpyxl.Workbook()
            us = wb.active
            us.title = "US(미국)"
            us["C5"] = "story-1"
            us["C7"] = "Welcome home"
            co = wb.create_sheet("CO(콜롬비아)")
            co["C5"] = "story-1"
            co["C7"] = "Te damos la bienvenida"
            wb.save(path)
            packet = await build_packet(path, "CO(콜롬비아)", semantic_rag_budget=3)
            self.assertEqual(packet["source_sheet"], "US(미국)")
            self.assertEqual(packet["semantic_rag_budget"], 3)
            self.assertEqual(openpyxl.load_workbook(path)["CO(콜롬비아)"]["C7"].value, "Te damos la bienvenida")

    async def test_five_role_escalation_is_off_unless_approved(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "story.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "US(미국)"; wb.create_sheet("CO(콜롬비아)"); wb.save(path)
            default = await build_packet(path, "CO(콜롬비아)")
            self.assertEqual(default["review_mode"], "staged_cell_sheet_lead")
            self.assertEqual(default["subagent_roles"], [])
            approved = await build_packet(path, "CO(콜롬비아)", multi_agent=True)
            self.assertEqual(approved["review_mode"], "multi_agent")
            self.assertEqual(len(approved["subagent_roles"]), 5)

    async def test_packet_id_is_stable_and_snapshots_reviewed_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "story.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "US(미국)"
            co = wb.create_sheet("CO(콜롬비아)"); co["C7"] = "Hola"; wb.save(path)
            first = await build_packet(path, "CO(콜롬비아)")
            second = await build_packet(path, "CO(콜롬비아)")
            self.assertEqual(first["packet_id"], second["packet_id"])
            self.assertEqual(first["cell_snapshot"], {"C7": "Hola"})

    async def test_rejects_negative_semantic_budget(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "story.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "US(미국)"; wb.create_sheet("CO(콜롬비아)"); wb.save(path)
            with self.assertRaises(ValueError):
                await build_packet(path, "CO(콜롬비아)", semantic_rag_budget=-1)

    async def test_injects_existing_hard_rule_evidence_without_recalculation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "story.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "US(미국)"; wb.create_sheet("CO(콜롬비아)"); wb.save(path)
            evidence = [{"cell": "C7", "hard_rule_issues": ["[괄호 오류]"], "sentence_case_report": None, "simple_case_fix": None}]
            with patch("agent_sheet_review._hard_rule_evidence",
                       new=AsyncMock(return_value=(evidence, "[HARD CONSTRAINTS]\n", [], ["문법/유창성"]))):
                packet = await build_packet(path, "CO(콜롬비아)", glossary=Path(tmp) / "g.csv", app_root=Path(tmp))
            self.assertEqual(packet["deterministic_evidence"], evidence)
            # The app's own authority wording travels with the evidence rather than
            # being paraphrased by the role prompts.
            self.assertEqual(packet["hard_constraint_preamble"], "[HARD CONSTRAINTS]\n")
        self.assertEqual(packet["hard_rule_policy"], "resolver_card_required; proposals_must_be_validated_before_merge")

    async def test_raw_mode_keeps_legacy_inspection_available(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "story.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "US(미국)"; co = wb.create_sheet("CO(콜롬비아)")
            co["C7"] = "Hola"; wb.save(path)
            result = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "agent_sheet_review.py"), str(path),
                 "--sheet", "CO(콜롬비아)", "--raw", "--json"],
                check=True, capture_output=True, text=True,
            )
            self.assertEqual(json.loads(result.stdout)["sheets"]["CO(콜롬비아)"]["groups"][0]["fields"]["title"]["text"], "Hola")
            self.assertEqual(openpyxl.load_workbook(path)["CO(콜롬비아)"]["C7"].value, "Hola")

    async def test_filters_english_first_candidate_overlay_to_the_current_story(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Story_001.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "US(미국)"; wb.create_sheet("CO(콜롬비아)"); wb.save(path)
            overlay = Path(tmp) / "overlay.json"
            overlay.write_text(json.dumps({"occurrences": [
                {"story": "001", "sheet": "CO(콜롬비아)", "cell": "C7", "source_term": "Energy"},
                {"story": "002", "sheet": "CO(콜롬비아)", "cell": "C7", "source_term": "Save"},
            ]}), encoding="utf-8")
            packet = await build_packet(path, "CO(콜롬비아)", candidate_overlay=overlay)
            self.assertEqual([item["source_term"] for item in packet["candidate_overlay"]], ["Energy"])
            self.assertEqual(packet["candidate_overlay_status"], "available")


if __name__ == "__main__":
    unittest.main()
