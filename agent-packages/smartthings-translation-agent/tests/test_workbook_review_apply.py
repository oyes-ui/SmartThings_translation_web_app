from __future__ import annotations

import sys
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_review_apply import (  # noqa: E402
    _apply_decisions,
    _parse_c_range,
    _parse_columns,
    _snapshot_protected,
    _validate_decisions,
    _verify_protected,
    decisions_from_report_changes,
)


class WorkbookReviewApplyTests(unittest.TestCase):
    def _workbook(self):
        wb = openpyxl.Workbook()
        target = wb.active
        target.title = "FR(프랑스)"
        target["C8"] = "current"
        target["F8"] = "reviewer"
        target["H8"] = "native comment"
        target["C10"] = "current partial"
        target["F10"] = "ignored full proposal"
        target["C11"] = "keep"
        protected = wb.create_sheet("US(미국)")
        protected["C8"] = "source"
        protected["I3"] = "must remain"
        protected.merge_cells("A1:B1")
        return wb

    def test_accept_partial_hold_apply_only_approved_c_values(self):
        wb = self._workbook()
        decisions = _validate_decisions(
            wb,
            [
                {"sheet": "FR(프랑스)", "cell": "C8", "decision": "accept", "source_sheet": "US(미국)"},
                {"sheet": "FR(프랑스)", "cell": "C10", "decision": "partial", "final_value": "approved partial"},
                {"sheet": "FR(프랑스)", "cell": "C11", "decision": "hold"},
            ],
            {"US(미국)"},
            "C7:C28",
        )
        records, changes = _apply_decisions(wb, decisions)
        self.assertEqual(wb["FR(프랑스)"]["C8"].value, "reviewer")
        self.assertEqual(wb["FR(프랑스)"]["C10"].value, "approved partial")
        self.assertEqual(wb["FR(프랑스)"]["C11"].value, "keep")
        self.assertEqual([entry["decision"] for entry in records], ["accept", "partial", "hold"])
        self.assertEqual([entry["decision"] for entry in changes], ["accept", "partial"])

    def test_protected_sheet_and_wrong_source_are_rejected(self):
        wb = self._workbook()
        with self.assertRaises(ValueError):
            _validate_decisions(
                wb, [{"sheet": "US(미국)", "cell": "C8", "decision": "accept"}], {"US(미국)"}, "C7:C28"
            )
        with self.assertRaises(ValueError):
            _validate_decisions(
                wb,
                [{"sheet": "FR(프랑스)", "cell": "C8", "decision": "accept", "source_sheet": "KR(한국)"}],
                {"US(미국)"}, "C7:C28"
            )

    def test_protected_snapshot_allows_explicit_review_column_removal_only(self):
        wb = self._workbook()
        before = _snapshot_protected(wb, {"US(미국)"}, (5, 8))
        for ws in wb.worksheets:
            ws.delete_cols(5, 4)
        after = _snapshot_protected(wb, {"US(미국)"})
        self.assertTrue(_verify_protected(before, after)["passed"])
        wb["US(미국)"]["C8"] = "changed"
        with self.assertRaises(RuntimeError):
            _verify_protected(before, _snapshot_protected(wb, {"US(미국)"}))

    def test_only_c_range_and_column_range_are_accepted(self):
        self.assertEqual(_parse_c_range("C7:C28"), (7, 28))
        self.assertEqual(_parse_columns("E:H"), (5, 8))
        with self.assertRaises(ValueError):
            _parse_c_range("B7:C28")


class ReportManifestConversionTests(unittest.TestCase):
    """report_format_spec.md approval manifest -> apply decisions."""

    CHANGES = [
        {
            "finding_id": "FR-C8", "sheet": "FR(프랑스)", "cell": "C8",
            "before": "current", "after": "ai proposal",
            "rule_ids": ["fr-tone-001"], "approval_status": "approved",
        },
        {
            "finding_id": "FR-C10", "sheet": "FR(프랑스)", "cell": "C10",
            "before": "current partial", "after": "rejected text",
            "rule_ids": [], "approval_status": "rejected",
        },
        {
            "finding_id": "FR-C11", "sheet": "FR(프랑스)", "cell": "C11",
            "before": "keep", "after": "not approved yet",
            "rule_ids": [], "approval_status": "pending_approval",
        },
    ]

    def test_only_approved_rows_convert(self):
        decisions = decisions_from_report_changes(self.CHANGES)
        self.assertEqual(len(decisions), 1)
        self.assertEqual(decisions[0]["cell"], "C8")
        self.assertEqual(decisions[0]["decision"], "accept")
        self.assertEqual(decisions[0]["final_value"], "ai proposal")
        self.assertEqual(decisions[0]["expected_before"], "current")
        self.assertEqual(decisions[0]["basis"], "fr-tone-001")

    def test_v2_context_does_not_change_changes_contract(self):
        manifest = {"manifest_schema_version": 2, "review_context": {"sheet_reviews": []}, "changes": self.CHANGES}
        decisions = decisions_from_report_changes(manifest["changes"])
        self.assertEqual(len(decisions), 1)
        self.assertEqual(decisions[0]["finding_id"], "FR-C8")

    def test_approved_row_requires_before_and_after(self):
        with self.assertRaises(ValueError):
            decisions_from_report_changes(
                [{"sheet": "FR(프랑스)", "cell": "C8", "before": "x", "approval_status": "approved"}]
            )
        with self.assertRaises(ValueError):
            decisions_from_report_changes(
                [{"sheet": "FR(프랑스)", "cell": "C8", "after": "y", "approval_status": "approved"}]
            )


class AiProposalApplyTests(unittest.TestCase):
    """An approved AI proposal applies without any F column."""

    def _workbook(self):
        wb = openpyxl.Workbook()
        target = wb.active
        target.title = "FR(프랑스)"
        target["C8"] = "current"          # deliberately no F8
        wb.create_sheet("US(미국)")["C8"] = "source"
        return wb

    def _decide(self, wb, raw):
        return _validate_decisions(wb, raw, {"US(미국)"}, "C7:C28")

    def test_accept_with_final_value_needs_no_reviewer_column(self):
        wb = self._workbook()
        decisions = self._decide(wb, decisions_from_report_changes([{
            "sheet": "FR(프랑스)", "cell": "C8", "before": "current",
            "after": "ai proposal", "rule_ids": ["fr-tone-001"],
            "approval_status": "approved",
        }]))
        records, changes = _apply_decisions(wb, decisions)
        self.assertEqual(wb["FR(프랑스)"]["C8"].value, "ai proposal")
        self.assertEqual(len(changes), 1)
        self.assertEqual(records[0]["rule_ids"], ["fr-tone-001"])
        self.assertTrue(records[0]["before_verified"])

    def test_drift_since_review_aborts_without_writing(self):
        wb = self._workbook()
        wb["FR(프랑스)"]["C8"] = "someone edited this after the review"
        decisions = self._decide(wb, [{
            "sheet": "FR(프랑스)", "cell": "C8", "decision": "accept",
            "final_value": "ai proposal", "expected_before": "current",
        }])
        with self.assertRaises(ValueError) as ctx:
            _apply_decisions(wb, decisions)
        self.assertIn("검수 이후 변경", str(ctx.exception))
        self.assertEqual(wb["FR(프랑스)"]["C8"].value, "someone edited this after the review")

    def test_no_op_final_value_is_rejected(self):
        wb = self._workbook()
        decisions = self._decide(wb, [{
            "sheet": "FR(프랑스)", "cell": "C8", "decision": "accept", "final_value": "current",
        }])
        with self.assertRaises(ValueError):
            _apply_decisions(wb, decisions)


if __name__ == "__main__":
    unittest.main()
