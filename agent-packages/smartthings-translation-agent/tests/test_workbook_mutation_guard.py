from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_mutation_guard import (  # noqa: E402
    clone_rich_value,
    copy_row_layout,
    delete_rows_with_manifest,
    save_verified_atomic,
    snapshot_axis_diff,
    snapshot_path,
)


class WorkbookMutationGuardTests(unittest.TestCase):
    def _book(self, path: Path, value: str = "source") -> Path:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "CO(콜롬비아)"
        sheet["C7"] = value
        book.save(path)
        book.close()
        return path

    def test_clone_rich_value_keeps_independent_runs(self):
        value = CellRichText(
            TextBlock(InlineFont(color="0000FF"), "Smart"),
            TextBlock(InlineFont(b=True), "Things"),
        )
        cloned = clone_rich_value(value)
        self.assertEqual(str(cloned), "SmartThings")
        self.assertIsNot(cloned, value)
        self.assertEqual(str(cloned[0].font.color.rgb), str(value[0].font.color.rgb))
        self.assertIsNot(cloned[0].font, value[0].font)

    def test_cross_workbook_row_layout_copies_semantics_not_style_ids(self):
        source = openpyxl.Workbook()
        target = openpyxl.Workbook()
        source.active.row_dimensions[7].height = 41
        source.active["C7"].font = Font(bold=True, color="00FF0000")
        source.active["C7"].fill = PatternFill("solid", fgColor="0000FF00")
        copy_row_layout(source.active, target.active, 7)
        self.assertEqual(target.active.row_dimensions[7].height, 41)
        self.assertTrue(target.active["C7"].font.bold)
        self.assertEqual(str(target.active["C7"].fill.fgColor.rgb), "0000FF00")
        source.close()
        target.close()

    def test_physical_row_delete_records_values_styles_dimensions_and_moves_following_rows(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["B15"] = "section"
        sheet["C15"] = "text"
        sheet["E15"] = "review note"
        sheet["C15"].fill = PatternFill("solid", fgColor="00FFFF00")
        sheet.row_dimensions[15].height = 90
        sheet["C16"] = "next"
        sheet.row_dimensions[16].height = 22
        deleted = delete_rows_with_manifest(sheet, 15, 1)
        self.assertEqual(deleted["values"]["E15"], "review note")
        self.assertIn("C15", deleted["styled_cells"])
        self.assertEqual(deleted["row_dimensions"]["15"]["height"], 90)
        self.assertEqual(sheet["C15"].value, "next")
        self.assertEqual(sheet.row_dimensions[15].height, 22)
        self.assertNotIn(16, sheet.row_dimensions)
        book.close()

    def test_snapshot_separates_value_style_and_layout_axes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self._book(root / "source.xlsx")
            styled = root / "styled.xlsx"
            book = openpyxl.load_workbook(source)
            book.active["C7"].font = Font(bold=True)
            book.save(styled)
            book.close()
            style_diff = snapshot_axis_diff(snapshot_path(source), snapshot_path(styled))
            self.assertNotIn("values", style_diff)
            self.assertIn("styles", style_diff)

            structured = root / "structured.xlsx"
            book = openpyxl.load_workbook(styled)
            book.active["C15"] = "section"
            book.active.row_dimensions[15].height = 90
            book.save(structured)
            book.close()
            structure_diff = snapshot_axis_diff(snapshot_path(styled), snapshot_path(structured))
            self.assertIn("values", structure_diff)
            self.assertIn("layout", structure_diff)

    def test_failed_verification_never_replaces_existing_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = self._book(root / "output.xlsx", "old")
            before = output.read_bytes()
            book = openpyxl.Workbook()
            book.active["C7"] = "new"

            def fail(_path: Path):
                raise RuntimeError("verification failed")

            with self.assertRaises(RuntimeError):
                save_verified_atomic(book, output, fail, overwrite=True)
            self.assertEqual(output.read_bytes(), before)
            self.assertFalse(list(root.glob(".output.verify-*.xlsx")))

            def pass_check(path: Path):
                reopened = openpyxl.load_workbook(path)
                try:
                    self.assertEqual(reopened.active["C7"].value, "new")
                finally:
                    reopened.close()
                return {"status": "verified"}

            result = save_verified_atomic(book, output, pass_check, overwrite=True)
            book.close()
            self.assertEqual(result["status"], "verified")
            verified = openpyxl.load_workbook(output)
            self.assertEqual(verified.active["C7"].value, "new")
            verified.close()


if __name__ == "__main__":
    unittest.main()
