from __future__ import annotations

import sys
import tempfile
import unittest
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_apply_edits import _load_edits, apply_edits  # noqa: E402


class WorkbookApplyEditsTests(unittest.TestCase):
    def _source(self, directory: str) -> Path:
        path = Path(directory) / "source.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CO(콜롬비아)"
        ws["C2"] = "Hola"
        ws["D2"] = "unchanged"
        ws["C3"] = "=SUM(1,2)"
        wb.save(path)
        return path

    def test_applies_approved_change_to_copy_and_verifies_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self._source(tmp)
            result = apply_edits(source, [{
                "sheet": "CO(콜롬비아)", "cell": "C2",
                "before": "Hola", "after": "Hola, ¿cómo estás?",
            }])
            self.assertEqual(result["status"], "ok")
            self.assertEqual(openpyxl.load_workbook(source)["CO(콜롬비아)"]["C2"].value, "Hola")
            revised = openpyxl.load_workbook(result["revised"])
            self.assertEqual(revised["CO(콜롬비아)"]["C2"].value, "Hola, ¿cómo estás?")
            self.assertEqual(revised["CO(콜롬비아)"]["D2"].value, "unchanged")
            self.assertTrue(Path(result["change_log"]).is_file())

    def test_approved_live_manifest_loads_as_safe_edits(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self._source(tmp)
            manifest = Path(tmp) / "approved.json"
            manifest.write_text(__import__("json").dumps({
                "surface": "chatgpt_excel", "mode": "draft", "approval": "approved",
                "changes": [{
                    "sheet": "CO(콜롬비아)", "cell": "C2", "before": "Hola",
                    "after": "Buenas", "verification": "verified",
                }],
            }, ensure_ascii=False), encoding="utf-8")
            result = apply_edits(source, _load_edits(str(manifest)))
            self.assertEqual(result["status"], "ok")
            self.assertEqual(openpyxl.load_workbook(result["revised"])["CO(콜롬비아)"]["C2"].value, "Buenas")

    def test_second_edit_chains_to_parent_revision(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self._source(tmp)
            first = apply_edits(source, [{
                "sheet": "CO(콜롬비아)", "cell": "C2", "before": "Hola", "after": "Buenas",
            }])
            second = apply_edits(Path(first["revised"]), [{
                "sheet": "CO(콜롬비아)", "cell": "C2", "before": "Buenas", "after": "Saludos",
            }])
            second_revision = json.loads(Path(second["revision_manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(second_revision["parent_revision"], first["revision_id"])

    def test_preview_and_safety_gates_do_not_write_a_copy(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self._source(tmp)
            preview = apply_edits(source, [{
                "sheet": "CO(콜롬비아)", "cell": "C2",
                "before": "Hola", "after": "Buenas",
            }], dry_run=True)
            self.assertEqual(preview["status"], "preview")
            stale = apply_edits(source, [{
                "sheet": "CO(콜롬비아)", "cell": "C2",
                "before": "stale", "after": "Buenas",
            }])
            formula = apply_edits(source, [{
                "sheet": "CO(콜롬비아)", "cell": "C3",
                "before": "=SUM(1,2)", "after": "=SUM(2,2)",
            }])
            self.assertEqual(stale["status"], "aborted")
            self.assertEqual(formula["status"], "aborted")
            self.assertFalse(list(Path(tmp).glob("*_revised_*.xlsx")))


if __name__ == "__main__":
    unittest.main()
