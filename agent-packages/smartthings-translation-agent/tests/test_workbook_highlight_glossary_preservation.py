from __future__ import annotations

import tempfile
import sys
import unittest
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_highlight_glossary import _restore_out_of_scope_rich_text  # noqa: E402
from workbook_mutation_guard import rich_value_signature  # noqa: E402


class HighlightScopePreservationTests(unittest.TestCase):
    def test_restores_only_rich_text_outside_processed_range(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            highlighted = Path(tmp) / "highlighted.xlsx"

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "KR(한국)"
            ws["A1"] = CellRichText(
                TextBlock(InlineFont(rFont="Arial", color="FF0000"), "keep"),
                TextBlock(InlineFont(rFont="Arial"), " this"),
            )
            ws["C13"] = CellRichText(
                TextBlock(InlineFont(rFont="Arial", color="0000FF"), "target")
            )
            wb.save(source)

            changed = openpyxl.load_workbook(source, rich_text=True)
            changed["KR(한국)"]["A1"].value = CellRichText(
                TextBlock(InlineFont(rFont="Arial"), "keep this")
            )
            changed["KR(한국)"]["C13"].value = CellRichText(
                TextBlock(InlineFont(rFont="Arial", color="00FF00"), "target")
            )
            changed.save(highlighted)
            changed.close()

            result = _restore_out_of_scope_rich_text(
                source, highlighted, "C13:C13", {"KR(한국)"}
            )

            before = openpyxl.load_workbook(source, rich_text=True)
            after = openpyxl.load_workbook(highlighted, rich_text=True)
            try:
                self.assertEqual(result["restored_cells"], ["KR(한국)!A1"])
                self.assertEqual(
                    rich_value_signature(before["KR(한국)"]["A1"].value),
                    rich_value_signature(after["KR(한국)"]["A1"].value),
                )
                self.assertNotEqual(
                    rich_value_signature(before["KR(한국)"]["C13"].value),
                    rich_value_signature(after["KR(한국)"]["C13"].value),
                )
            finally:
                before.close()
                after.close()


if __name__ == "__main__":
    unittest.main()
