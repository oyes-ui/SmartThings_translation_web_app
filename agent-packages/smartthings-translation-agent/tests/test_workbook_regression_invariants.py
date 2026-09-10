"""Synthetic accident regressions; no ignored writer or real workbook dependency.

Authorities: each fixture is its own value/layout/format baseline. Changes are
explicit below. Files are temporary test artifacts, never delivery workbooks.
Expected properties are checked after reopening, independently of save callbacks.
"""
from pathlib import Path
import sys
import unittest
from tempfile import TemporaryDirectory

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from workbook_mutation_guard import (
    clone_rich_value, copy_row_layout, delete_rows_with_manifest,
    save_verified_atomic, semantic_workbook_snapshot, snapshot_axis_diff,
    slice_rich_value,
)


def minimal_book():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "US(미국)"
    sheet["C7"] = " 013  keep spaces "
    sheet["C8"] = "=1+2"
    sheet["C9"] = CellRichText(
        TextBlock(InlineFont(color="0000FF"), "Device"),
        TextBlock(InlineFont(b=True), " 7"), " suffix ")
    sheet.row_dimensions[12].height = 22
    sheet.sheet_properties.tabColor = "FFFF00"
    sheet["C7"].comment = Comment("review", "reviewer")
    return book


class WorkbookRegressionInvariants(unittest.TestCase):
    def roundtrip(self, book, check):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "fixture.xlsx"
            def verify(candidate):
                reopened = openpyxl.load_workbook(candidate, rich_text=True, data_only=False)
                try:
                    check(reopened)
                finally:
                    reopened.close()
            save_verified_atomic(book, path, verify)
            verify(path)

    def test_rich_clone_roundtrip_preserves_runs_fonts_and_whitespace(self):
        book = minimal_book()
        self.addCleanup(book.close)
        book.active["C9"] = clone_rich_value(book.active["C9"].value)
        def check(actual):
            value = actual.active["C9"].value
            self.assertEqual(str(value), "Device 7 suffix ")
            self.assertEqual(len(value), 3)
            self.assertEqual(value[0].font.color.rgb, "000000FF")
            self.assertTrue(value[1].font.b)
            self.assertEqual(value[2], " suffix ")
        self.roundtrip(book, check)

    def test_cross_run_edit_keeps_surviving_character_fonts(self):
        book = minimal_book()
        self.addCleanup(book.close)
        original = book.active["C9"].value
        # Only offset 6 (the space before 7) is approved for removal.
        edited = CellRichText()
        for start, end in ((0, 6), (7, len(str(original)))):
            edited.extend(slice_rich_value(original, start, end))
        book.active["C9"] = edited
        def check(actual):
            value = actual.active["C9"].value
            self.assertEqual(str(value), "Device7 suffix ")
            self.assertEqual(value[0].font.color.rgb, "000000FF")
            self.assertEqual(value[1].text, "7")
            self.assertTrue(value[1].font.b)
            self.assertEqual(value[2], " suffix ")
        self.roundtrip(book, check)
        self.assertEqual(str(original), "Device 7 suffix ")

    def test_noop_roundtrip_keeps_five_axes_and_literal_source_values(self):
        book = minimal_book()
        self.addCleanup(book.close)
        baseline = semantic_workbook_snapshot(book)
        def check(actual):
            self.assertEqual(actual.active["C7"].value, " 013  keep spaces ")
            self.assertEqual(actual.active["C8"].value, "=1+2")
            self.assertEqual(actual.active.row_dimensions[12].height, 22)
            self.assertEqual(snapshot_axis_diff(baseline, semantic_workbook_snapshot(actual)), [])
        self.roundtrip(book, check)

    def test_wrong_layout_authority_is_detected_without_value_changes(self):
        base, other = minimal_book(), minimal_book()
        self.addCleanup(base.close)
        self.addCleanup(other.close)
        baseline = semantic_workbook_snapshot(base)
        other.active.row_dimensions[12].height = 77
        other.active["B12"].font = Font(bold=True)
        copy_row_layout(other.active, base.active, 12)
        self.assertEqual(set(snapshot_axis_diff(baseline, semantic_workbook_snapshot(base))),
                         {"layout", "styles"})

    def test_delete_section_removes_styled_tail_and_preserves_separator(self):
        book = minimal_book()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.row_dimensions[14].height = 31
        sheet["B14"].border = Border(bottom=Side(style="thin"))
        sheet["B15"] = "//section_deleted"
        sheet["C15"] = "deleted text"
        sheet["E15"] = "review note"
        sheet["C18"].border = Border(bottom=Side(style="thick"))
        sheet.row_dimensions[18].height = 90
        manifest = delete_rows_with_manifest(sheet, 15, 4)
        self.assertEqual(manifest["values"]["E15"], "review note")
        self.assertIn("C18", manifest["styled_cells"])
        def check(actual):
            target = actual.active
            self.assertEqual(target.max_row, 14)
            self.assertEqual(target.row_dimensions[14].height, 31)
            self.assertEqual(target["B14"].border.bottom.style, "thin")
            self.assertFalse(any(row > 14 for row in target.row_dimensions))
            self.assertFalse(any(row > 14 for row, col in target._cells))
        self.roundtrip(book, check)

    def test_cross_workbook_styles_survive_distinct_style_tables(self):
        source, target = minimal_book(), minimal_book()
        self.addCleanup(source.close)
        self.addCleanup(target.close)
        source.active["C12"].font = Font(bold=True, color="FF0000")
        target.active["A1"].font = Font(italic=True, color="00FF00")
        copy_row_layout(source.active, target.active, 12)
        def check(actual):
            self.assertTrue(actual.active["C12"].font.bold)
            self.assertEqual(actual.active["C12"].font.color.rgb, "00FF0000")
            self.assertTrue(actual.active["A1"].font.italic)
            self.assertEqual(actual.active["C7"].value, " 013  keep spaces ")
        self.roundtrip(target, check)

    def test_five_axes_detect_individual_accident_mutations(self):
        cases = [
            ("values", lambda s: setattr(s["C7"], "value", "013 keep spaces")),
            ("layout", lambda s: setattr(s.row_dimensions[12], "height", 77)),
            ("layout", lambda s: setattr(s.sheet_properties, "tabColor", "FF0000")),
            ("styles", lambda s: setattr(s["B14"], "border", Border(bottom=Side(style="thin")))),
            ("rich_text", lambda s: setattr(s["C9"], "value", str(s["C9"].value))),
            ("annotations", lambda s: setattr(s["C7"], "comment", None)),
        ]
        for axis, mutate in cases:
            with self.subTest(axis=axis):
                book = minimal_book()
                try:
                    baseline = semantic_workbook_snapshot(book)
                    mutate(book.active)
                    self.assertIn(axis, snapshot_axis_diff(baseline, semantic_workbook_snapshot(book)))
                finally:
                    book.close()

    def test_rejected_candidate_does_not_publish_new_output(self):
        book = minimal_book()
        self.addCleanup(book.close)
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "delivery.xlsx"
            def reject(candidate):
                self.assertTrue(candidate.exists())
                raise ValueError("unexpected layout diff")
            with self.assertRaisesRegex(ValueError, "unexpected layout"):
                save_verified_atomic(book, path, reject)
            self.assertFalse(path.exists())
            self.assertEqual(list(Path(tmp).iterdir()), [])
