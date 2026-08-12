from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from agent_review_contract import SPECIALIST_ROLES  # noqa: E402

PACKET_ID = "pkt0000000000001"


def _write_case(tmp: Path, roles=SPECIALIST_ROLES, after="Buenas"):
    """A workbook, a packet and one opinion file per role."""
    workbook = tmp / "story.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "CO(콜롬비아)"; ws["C10"] = "Hola"; wb.save(workbook)

    packet = tmp / "packet.json"
    packet.write_text(json.dumps({
        "kind": "agent_sheet_review_packet", "packet_id": PACKET_ID,
        "workbook_name": "story.xlsx", "target_sheet": "CO(콜롬비아)", "source_sheet": "US(미국)",
        "cell_snapshot": {"C10": "Hola"}, "semantic_rag_budget": 0,
        "review_mode": "multi_agent",
    }, ensure_ascii=False), encoding="utf-8")

    opinions_dir = tmp / "full_review"; opinions_dir.mkdir()
    supporters = {"grammar_fluency", "localization_tone"}
    for role in roles:
        entries = [{"role": role, "sheet": "CO(콜롬비아)", "cell": "C10", "finding_id": f"filler-{role}",
                    "stance": "review", "after": None, "constraint_status": "pass"}]
        if role in supporters:
            entries.append({"role": role, "sheet": "CO(콜롬비아)", "cell": "C10", "finding_id": "co-c10",
                            "stance": "support", "after": after, "rule_ids": ["co-1"],
                            "reason": f"{role} 근거", "constraint_status": "pass"})
        (opinions_dir / f"{role}.json").write_text(
            json.dumps({"role": role, "packet_id": PACKET_ID, "opinions": entries}, ensure_ascii=False),
            encoding="utf-8")
    return workbook, packet, opinions_dir


def _run(tmp: Path, workbook, packet, opinions_dir, report_id="r"):
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "agent_sheet_merge.py"),
         "--packet", str(packet), "--opinions-dir", str(opinions_dir),
         "--workbook", str(workbook), "--report-id", report_id, "--output-dir", str(tmp / "out")],
        capture_output=True, text=True,
    )
    return result, json.loads(result.stdout)


class ResolverGateRequirementTests(unittest.TestCase):
    def test_a_packet_with_cards_refuses_to_merge_without_the_resolver(self):
        """Evidence to validate a proposal exists, so shipping unvalidated text is refused."""
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            workbook, packet, opinions_dir = _write_case(tmp)
            payload = json.loads(packet.read_text(encoding="utf-8"))
            payload["deterministic_evidence_status"] = "available"
            payload["deterministic_evidence"] = [
                {"cell": "C10", "row_type": "description", "constraint_card": {"terms": []}},
            ]
            packet.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result, out = _run(tmp, workbook, packet, opinions_dir)
            self.assertEqual(result.returncode, 2)
            self.assertIn("--glossary", out["error"])

    def test_a_packet_without_cards_reports_the_gate_did_not_run(self):
        """Nothing to validate against is legitimate, but it must be visible in the report."""
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            result, payload = _run(tmp, *_write_case(tmp))
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(payload["resolver_gate"], "not_run")
            markdown = Path(payload["report"]).read_text(encoding="utf-8")
            self.assertIn("resolver로 재검증되지 않았습니다", markdown)


class AgentSheetMergeTests(unittest.TestCase):
    def test_full_role_set_produces_applyable_changes(self):
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            result, payload = _run(tmp, *_write_case(tmp))
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(payload["sheet_status"], "completed")
            self.assertEqual(payload["changes"], 1)
            manifest = json.loads(Path(payload["manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(manifest["changes"][0]["approval_status"], "pending_approval")
            self.assertEqual(manifest["changes"][0]["before"], "Hola")
            self.assertEqual(manifest["changes"][0]["supporting_roles"],
                             ["grammar_fluency", "localization_tone"])

    def test_missing_role_file_yields_incomplete_and_no_changes(self):
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            workbook, packet, opinions_dir = _write_case(tmp)
            (opinions_dir / "semantic_fidelity.json").unlink()
            _, payload = _run(tmp, workbook, packet, opinions_dir)
            self.assertEqual(payload["sheet_status"], "incomplete")
            self.assertEqual(payload["changes"], 0)
            self.assertIn("semantic_fidelity", payload["missing_roles"])

    def test_original_workbook_is_never_modified(self):
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            workbook, packet, opinions_dir = _write_case(tmp)
            _run(tmp, workbook, packet, opinions_dir)
            self.assertEqual(openpyxl.load_workbook(workbook)["CO(콜롬비아)"]["C10"].value, "Hola")

    def test_merge_refuses_a_sheet_not_approved_for_five_roles(self):
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            workbook, packet, opinions_dir = _write_case(tmp)
            payload = json.loads(packet.read_text(encoding="utf-8"))
            packet.write_text(json.dumps({**payload, "review_mode": "lead_2pass"}), encoding="utf-8")
            result, body = _run(tmp, workbook, packet, opinions_dir)
            self.assertEqual(result.returncode, 2)
            self.assertIn("--multi-agent", body["error"])

    def test_findings_carry_their_content_row_type(self):
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            _, payload = _run(tmp, *_write_case(tmp))
            manifest = json.loads(Path(payload["manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(manifest["changes"][0]["row_type"], "title")  # C10
            self.assertEqual(manifest["row_type_counts"]["title"]["changes"], 1)
            self.assertIn("콘텐츠 유형별 finding", Path(payload["report"]).read_text(encoding="utf-8"))

    def test_role_name_mismatch_in_a_file_is_an_error(self):
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            workbook, packet, opinions_dir = _write_case(tmp)
            path = opinions_dir / "grammar_fluency.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            path.write_text(json.dumps({**payload, "role": "localization_tone"}), encoding="utf-8")
            result, body = _run(tmp, workbook, packet, opinions_dir)
            self.assertEqual(result.returncode, 2)
            self.assertIn("역할 불일치", body["error"])


if __name__ == "__main__":
    unittest.main()
