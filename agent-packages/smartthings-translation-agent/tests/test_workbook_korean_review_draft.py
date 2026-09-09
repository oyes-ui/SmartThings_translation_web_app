from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_korean_review_draft import build_workbook, verify_output  # noqa: E402


class KoreanReviewDraftTests(unittest.TestCase):
    def test_only_scoped_f_cells_change_and_blue_overlays_red(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            output = Path(tmp) / "output.xlsx"
            workbook = openpyxl.Workbook()
            korean = workbook.active
            korean.title = "KR(한국)"
            other = workbook.create_sheet("US(미국)")
            blue = InlineFont(color="FF0000FF")
            korean["C7"].value = CellRichText([
                TextBlock(text="스마트싱스", font=blue),
                "에 연결되어있는 기기",
            ])
            korean["F7"] = "x"
            korean["C8"] = "일반 기기"
            korean["F8"] = "x"
            korean["C10"] = "제외 문안"
            korean["F10"] = "제외 유지"
            other["C7"].value = CellRichText([TextBlock(text="Untouched", font=blue)])
            workbook.save(source)

            spec = {
                "target_rows": [7, 8],
                "glossary_terms": {"F8": ["기기"]},
                "edits": [{
                    "cell": "F7",
                    "before": "스마트싱스에 연결되어있는 기기",
                    "after": "스마트싱스에 연결되어 있는 기기",
                    "blue_terms": ["연결되어 있는"],
                }],
            }
            revised = build_workbook(source, spec)
            revised.save(output)
            revised.close()
            result = verify_output(output, source, spec)
            self.assertEqual(result["status"], "verified")

            restored = openpyxl.load_workbook(output, rich_text=True)
            self.assertEqual(str(restored["KR(한국)"]["F7"].value), "스마트싱스에 연결되어 있는 기기")
            runs = [
                (part.text, part.font.color.rgb if isinstance(part, TextBlock) and part.font.color else None)
                for part in restored["KR(한국)"]["F7"].value
            ]
            self.assertIn(("스마트싱스", "FF0000FF"), runs)
            self.assertTrue(any(text == "연결되어 있는" and color == "FF0000FF" for text, color in runs))
            self.assertEqual(restored["KR(한국)"]["F10"].value, "제외 유지")
            f8_runs = [
                (part.text, part.font.color.rgb if isinstance(part, TextBlock) and part.font.color else None)
                for part in restored["KR(한국)"]["F8"].value
            ]
            self.assertTrue(any(text == "기기" and color == "FF0000FF" for text, color in f8_runs))
            self.assertEqual(str(restored["US(미국)"]["C7"].value), "Untouched")
            restored.close()


if __name__ == "__main__":
    unittest.main()
