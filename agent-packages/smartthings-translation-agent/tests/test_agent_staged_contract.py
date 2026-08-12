from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from agent_staged_contract import (gate_stage, merge_staged_reviews, validate_cell_review,
                                   validate_lead_review, validate_sheet_review)
from review_report_builder import build_review_artifacts


PACKET = {
    "packet_id": "pkt-stage-1", "target_sheet": "CO(콜롬비아)",
    "cell_snapshot": {"C7": "Actual", "C8": "Actual 2"},
    "deterministic_evidence": [
        {"cell": "C7", "source_text": "Movie mode", "target_text": "Modo de video",
         "constraint_card": {"terms": []}},
        {"cell": "C8", "source_text": "A safe home", "target_text": "Un hogar seguro",
         "constraint_card": {"terms": []}},
    ],
    "activation_candidates": [
        {"story": "001", "cell": "C8", "source_term": "Safe", "confirmed": False},
    ],
}


def cell_payload():
    return {"kind": "cell_review", "packet_id": "pkt-stage-1", "status": "completed",
            "stop_reason": "complete", "cells": [
                {"cell": "C7", "status": "needs_revision", "after": "Modo película", "reason": "natural",
                 "rule_ids": [], "used_prior_cell_context": False, "prior_cell_refs": [], "prior_cell_influence": ""},
                {"cell": "C8", "status": "pass", "after": None, "reason": "ok", "rule_ids": [],
                 "used_prior_cell_context": True, "prior_cell_refs": ["C7"],
                 "prior_cell_influence": "CTA 길이만 비교"},
            ]}


def sheet_payload():
    return {"kind": "sheet_consistency_review", "packet_id": "pkt-stage-1", "status": "completed",
            "stop_reason": "complete", "issues": [
                {"finding_id": "sheet-cta", "affected_cells": ["C7"], "canonical_pattern": "간결한 CTA",
                 "reason": "통일", "rule_ids": [],
                 "proposals": [{"cell": "C7", "after": "Modo película", "rule_ids": []}]},
            ]}


def lead_payload():
    return {"kind": "lead_review", "packet_id": "pkt-stage-1", "status": "completed",
            "stop_reason": "complete", "decisions": [
                {"cell": "C7", "finding_id": "lead-c7", "status": "needs_revision", "after": "Modo película",
                 "reason": "두 근거 통합", "rule_ids": [], "basis_refs": ["cell:C7", "sheet:sheet-cta"]},
                {"cell": "C8", "finding_id": "lead-c8", "status": "pass", "after": None,
                 "reason": "유지", "rule_ids": [], "basis_refs": ["cell:C8"]},
            ]}


class StagedContractTests(unittest.TestCase):
    def test_cell_review_requires_measurable_prior_cell_references(self):
        payload = cell_payload(); del payload["cells"][0]["used_prior_cell_context"]
        with self.assertRaises(ValueError):
            validate_cell_review(payload, PACKET)

    def test_sheet_agent_must_cover_each_affected_cell_with_a_full_proposal(self):
        payload = sheet_payload(); payload["issues"][0]["proposals"] = []
        with self.assertRaises(ValueError):
            validate_sheet_review(payload, PACKET)

    def test_lead_cannot_cite_a_sheet_issue_that_does_not_cover_the_cell(self):
        payload = lead_payload(); payload["decisions"][1]["basis_refs"] = ["sheet:sheet-cta"]
        with self.assertRaises(ValueError):
            validate_lead_review(payload, PACKET, cell_payload(), sheet_payload())

    def test_missing_evidence_fails_closed(self):
        gated = gate_stage(cell_payload(), "cell_review", PACKET, lambda cell, after: None)
        self.assertEqual(gated["cells"][0]["status"], "blocked")
        self.assertEqual(gated["cells"][0]["resolver_disposition"], "missing_constraint_evidence")

    def test_unconfirmed_candidate_routes_only_matching_missing_target_to_activation_review(self):
        payload = cell_payload(); payload["cells"][0]["after"] = None; payload["cells"][0]["status"] = "pass"
        payload["cells"][1]["after"] = "hogar"; payload["cells"][1]["status"] = "needs_revision"
        gated = gate_stage(payload, "cell_review", PACKET, lambda cell, after: {
            "status": "blocked", "blocked": True, "normalized": after,
            "violations": [{"source_term": "Safe", "reason": "missing_glossary_target", "expected": "Safe"}],
            "review": [],
        })
        self.assertEqual(gated["cells"][1]["status"], "glossary_activation_review")

    def test_non_activation_violation_remains_blocked(self):
        gated = gate_stage(cell_payload(), "cell_review", PACKET, lambda cell, after: {
            "status": "blocked", "blocked": True, "normalized": after,
            "violations": [{"source_term": "Movie mode", "reason": "missing_glossary_target",
                            "expected": "Modo de video"}], "review": [],
        })
        self.assertEqual(gated["cells"][0]["status"], "blocked")
        self.assertEqual(gated["cells"][0]["resolver_disposition"], "invalidated_by_hard_constraint")

    def test_empty_activation_candidates_keep_hard_constraint_blocked(self):
        packet = {**PACKET, "activation_candidates": []}
        gated = gate_stage(cell_payload(), "cell_review", packet, lambda cell, after: {
            "status": "blocked", "blocked": True, "normalized": after,
            "violations": [{"source_term": "Movie mode", "reason": "missing_glossary_target",
                            "expected": "Modo de video"}], "review": [],
        })
        self.assertEqual(gated["cells"][0]["status"], "blocked")
        self.assertEqual(gated["cells"][0]["resolver_disposition"], "invalidated_by_hard_constraint")
        self.assertEqual(gated["cells"][0]["resolver_violations"][0]["expected"], "Modo de video")

    def test_final_merge_gates_the_lead_and_records_anchor_metrics(self):
        cell = gate_stage(cell_payload(), "cell_review", PACKET, lambda cell, after: {
            "status": "pass", "blocked": False, "normalized": after, "violations": [], "review": []})
        sheet = gate_stage(sheet_payload(), "sheet_consistency_review", PACKET, lambda cell, after: {
            "status": "pass", "blocked": False, "normalized": after, "violations": [], "review": []})
        merged = merge_staged_reviews(PACKET, cell, sheet, lead_payload(), lambda cell, after: {
            "status": "pass", "blocked": False, "normalized": after, "violations": [], "review": []})
        self.assertEqual(len(merged.proposals), 1)
        self.assertEqual(merged.anchoring_metrics["referenced_cells"], 1)
        self.assertEqual(merged.stage_reviews["lead_review"]["kind"], "lead_review")

    def test_full_audit_report_contains_all_three_stages_and_anchor_rate(self):
        passing = lambda cell, after: {"status": "pass", "blocked": False, "normalized": after,
                                       "violations": [], "review": []}
        cell = gate_stage(cell_payload(), "cell_review", PACKET, passing)
        sheet = gate_stage(sheet_payload(), "sheet_consistency_review", PACKET, passing)
        merged = merge_staged_reviews(PACKET, cell, sheet, lead_payload(), passing)
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "story.xlsx"
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "CO(콜롬비아)"
            ws["C7"] = "Actual"; ws["C8"] = "Actual 2"; wb.save(path)
            manifest, markdown = build_review_artifacts(path, merged, report_id="staged", source_file_id="story")
        self.assertIn("## 셀 검수", markdown)
        self.assertIn("시트 일관성 의견", markdown)
        self.assertIn("리드 최종 판정", markdown)
        self.assertIn("참조율: 50.0%", markdown)
        self.assertEqual(manifest["review_context"]["anchoring_metrics"]["referenced_cells"], 1)


if __name__ == "__main__":
    unittest.main()
