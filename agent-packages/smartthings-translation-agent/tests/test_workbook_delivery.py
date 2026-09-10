"""Default command integration, synthetic Excel only; zero external API calls.

Source workbook is value/structure/format authority. Only explicit edits,
review-column deletion and C7:C28 glossary runs are permitted to differ.
"""
import asyncio
import json
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Border, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import workbook_delivery as delivery
from workbook_apply_edits import apply_edits
from workbook_contract import atomic_json, file_sha256, InputDrift
from workbook_verifier import facts, diff_facts

APP = Path(__file__).resolve().parents[3]


@pytest.fixture
def job(tmp_path):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in ("KR(한국)", "US(미국)", "DE(독일)"):
        ws = book.create_sheet(name)
        ws["C5"] = "story_test"
        ws["C7"] = "SmartThings old"
        ws["C8"] = "  Keep space  "
        ws["F7"] = "SmartThings neu"
        ws["H7"] = "review note"
        ws["A12"].border = Border(bottom=Side(style="thin"))
        ws["D9"] = CellRichText(TextBlock(InlineFont(color="FF0000"), "outside"), " scope")
        ws["D10"].comment = Comment("retain", "reviewer")
        ws.row_dimensions[12].height = 21
    source = tmp_path / "source.xlsx"
    book.save(source)
    book.close()
    glossary = tmp_path / "glossary.csv"
    glossary.write_text("Key,규칙,한국어,영어_미국,독어_독일\nKey,Rule,ko_KR,en_US,de_DE\nLng,Rule,Lng,Lng,Lng\nSmartThings,,SmartThings,SmartThings,SmartThings\n")
    approval = tmp_path / "approved.json"
    atomic_json(approval, {"protected_sheets": ["KR(한국)", "US(미국)"], "changes": [
        {"sheet": "DE(독일)", "cell": "C7", "before": "SmartThings old", "after": "SmartThings neu", "approval_status": "approved"},
        {"sheet": "DE(독일)", "cell": "C8", "before": "  Keep space  ", "after": "unapproved", "approval_status": "pending_approval"}]})
    return SimpleNamespace(workbook=str(source), glossary=str(glossary), app_root=str(APP),
        edits=json.dumps([{"sheet": "DE(독일)", "cell": "C7", "before": "SmartThings old", "after": "SmartThings neu"}]),
        approval_manifest=str(approval), output=str(tmp_path / "accepted.xlsx"),
        protected_sheets="KR(한국),US(미국)", drop_review_columns=None, delivery_sheets="DE(독일)",
        cell_range="C7:C28", dry_run=False, activation_manifest=None, result_manifest=None)


def run(job, mode):
    with patch.object(delivery.ap, "maybe_reexec_with_app_venv"):
        return asyncio.run(delivery.run_delivery(job, mode))


@pytest.mark.parametrize("mode", ["story", "review"])
def test_default_delivery_batch_and_resume(job, mode):
    original = file_sha256(job.workbook)
    result = run(job, mode)
    assert result["status"] == "ok", result
    assert result == run(job, mode)
    assert file_sha256(job.workbook) == original
    manifest = json.loads(Path(result["batch_manifest"]).read_text())
    assert set(manifest["outputs"]) == {"edited", "final"}
    book = openpyxl.load_workbook(result["final"], rich_text=True)
    try:
        assert str(book["DE(독일)"]["C7"].value) == "SmartThings neu"
        assert str(book["DE(독일)"]["C8"].value) == "  Keep space  "
        for name in book.sheetnames:
            assert isinstance(book[name]["C7"].value, CellRichText)
            assert isinstance(book[name]["D9"].value, CellRichText)
            assert book[name]["D10"].comment.text == "retain"
    finally:
        book.close()
    if mode == "review":
        assert result["decision_counts"] == {"accept": 1, "partial": 0, "hold": 0}


def test_preview_and_glossary_drift_never_publish(job):
    job.dry_run = True
    assert run(job, "story")["status"] == "preview"
    assert not (Path(job.workbook).parent / "verified").exists()
    Path(job.glossary).write_text(Path(job.glossary).read_text() + "New,,New,New,New\n")
    job.dry_run = False
    with pytest.raises(InputDrift):
        run(job, "story")
    assert not (Path(job.workbook).parent / "verified").exists()


def test_review_column_removal_preserves_unrelated_formatting(job):
    job.drop_review_columns = "E:H"
    result = run(job, "review")
    assert result["status"] == "ok", result
    book = openpyxl.load_workbook(result["final"], rich_text=True)
    try:
        for ws in book:
            assert ws.max_column == 4
            assert ws.row_dimensions[12].height == 21
            assert ws["A12"].border.bottom.style == "thin"
            assert ws["D10"].comment.text == "retain"
    finally:
        book.close()


def test_final_failure_never_exposes_acceptance(job):
    job.dry_run = True
    run(job, "review")
    job.dry_run = False
    original = delivery._execute_build
    def corrupt(contract, artifact, last_failure):
        book = original(contract, artifact, last_failure)
        if artifact["id"] == "final":
            book["KR(한국)"]["C8"] = "unauthorized"
        return book
    with patch.object(delivery, "_execute_build", side_effect=corrupt):
        result = run(job, "review")
    assert result["status"] != "ok"
    assert not Path(job.output).exists()
    assert not list((Path(job.output).parent / "verified").glob("*/*.xlsx"))


def test_default_edit_preview_pins_original(job):
    edits = json.loads(job.edits)
    assert apply_edits(Path(job.workbook), edits, dry_run=True)["status"] == "preview"
    book = openpyxl.load_workbook(job.workbook)
    book["DE(독일)"]["C8"] = "drift"
    book.save(job.workbook)
    book.close()
    assert apply_edits(Path(job.workbook), edits)["status"] == "aborted"
    assert not (Path(job.workbook).parent / "verified").exists()


def test_shifted_columns_and_dimension_projection(job):
    book = openpyxl.load_workbook(job.workbook, rich_text=True)
    for ws in book:
        ws["J7"] = "shifted note"
        ws.column_dimensions["J"].width = 29
        ws.column_dimensions["A"].width = 18
        ws.merge_cells("J9:K9")
        ws["J9"] = "merged"
    book.save(job.workbook)
    book.close()
    job.drop_review_columns = "E:H"
    result = run(job, "review")
    assert result["status"] == "ok", result
    book = openpyxl.load_workbook(result["final"], rich_text=True)
    try:
        for ws in book:
            assert ws["F7"].value == "shifted note"
            assert ws.column_dimensions["F"].width == 29
            assert "F9:G9" in ws.merged_cells
    finally:
        book.close()


def test_outcomes_only_records_human_decisions(job):
    result = run(job, "review")
    entries = [json.loads(x) for x in Path(result["outcome_ledger"]).read_text().splitlines()]
    assert len(entries) == 1
    assert entries[0]["cell"] == "C7"
    assert entries[0]["decision"] == "approved"
    assert entries[0]["work_id"] == result["work_id"]


def test_cli_uses_same_default_delivery(job):
    import subprocess
    script = Path(delivery.__file__).with_name("workbook_story_apply.py")
    process = subprocess.run([sys.executable, str(script), job.workbook, job.edits,
        "--delivery-sheets", job.delivery_sheets, "--glossary", job.glossary,
        "--app-root", job.app_root, "--json"], capture_output=True, text=True)
    assert process.returncode == 0, process.stdout + process.stderr
    result = json.loads(process.stdout)
    assert result["artifact_status"] == "delivery"
    assert Path(result["batch_manifest"]).is_file()


def test_comment_box_relocation_requires_separate_plan(job):
    book = openpyxl.load_workbook(job.workbook, rich_text=True)
    book["DE(독일)"]["F7"].comment = Comment("box", "reviewer")
    book.save(job.workbook)
    book.close()
    job.drop_review_columns = "E:H"
    with pytest.raises(ValueError, match="VML"):
        run(job, "review")
    assert not (Path(job.workbook).parent / "verified").exists()
