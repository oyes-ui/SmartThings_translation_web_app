from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from workbook_manifest import create_edit_revision, diff_spans, ensure_baseline, workbook_snapshot  # noqa: E402


class WorkbookManifestTests(unittest.TestCase):
    def _book(self, directory: str) -> Path:
        path = Path(directory) / "story.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CO(콜롬비아)"
        ws["C7"] = "Hola mundo"
        wb.save(path)
        return path

    def test_baseline_is_external_and_is_created_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = self._book(tmp)
            before = book.read_bytes()
            baseline, path = ensure_baseline(book)
            again, again_path = ensure_baseline(book)
            self.assertEqual(book.read_bytes(), before)
            self.assertEqual(path, again_path)
            self.assertEqual(baseline["workbook_id"], again["workbook_id"])
            self.assertEqual(path.parent.parent.name, ".st-history")
            self.assertEqual(baseline["source_file"], "story.xlsx")

    def test_revision_contains_diff_and_no_absolute_artifact_paths(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self._book(tmp)
            revised = Path(tmp) / "story_revised.xlsx"
            wb = openpyxl.load_workbook(source)
            wb["CO(콜롬비아)"]["C7"] = "Hola, mundo"
            wb.save(revised)
            payload, path = create_edit_revision(source, revised, [{
                "sheet": "CO(콜롬비아)", "cell": "C7", "old_value": "Hola mundo", "new_value": "Hola, mundo",
            }])
            saved = json.loads(path.read_text(encoding="utf-8"))
            self.assertEqual(payload["revision_id"], saved["revision_id"])
            self.assertEqual(saved["source_file"], source.name)
            self.assertEqual(saved["revised_file"], revised.name)
            self.assertNotIn(tmp, json.dumps(saved, ensure_ascii=False))
            self.assertTrue(saved["changes"][0]["diff"]["red_spans"])

    def test_delete_marks_both_adjacent_words_or_one_at_edge(self):
        middle = diff_spans("Turn on the light", "Turn the light")
        self.assertEqual([value.strip() for value in middle["deleted_text"]], ["on"])
        self.assertEqual(middle["red_spans"], [[0, 4], [5, 8]])
        edge = diff_spans("Turn on", "on")
        self.assertEqual(edge["red_spans"], [[0, 2]])

    def test_snapshot_records_independent_layout_style_and_rich_text_fingerprints(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = self._book(tmp)
            snapshot = workbook_snapshot(book)
            legacy_values = hashlib.sha256("CO(콜롬비아)\0C7\0Hola mundo".encode("utf-8")).hexdigest()
            self.assertEqual(snapshot["values_sha256"], legacy_values)
            self.assertNotEqual(snapshot["values_sha256"], snapshot["semantic_values_sha256"])
            self.assertNotEqual(snapshot["structure_sha256"], snapshot["layout_sha256"])
            self.assertIn("styles_sha256", snapshot)
            self.assertIn("rich_text_sha256", snapshot)
            self.assertIn("annotations_sha256", snapshot)


if __name__ == "__main__":
    unittest.main()
