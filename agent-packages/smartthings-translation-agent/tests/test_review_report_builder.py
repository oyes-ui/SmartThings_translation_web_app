from __future__ import annotations
import sys, tempfile, unittest
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(ROOT/"scripts"))
from agent_review_contract import SPECIALIST_ROLES, merge_subjective_opinions
from review_report_builder import build_review_artifacts

PACKET_ID="pkt0000000000001"

def _packet(snapshot=None):
 return {"packet_id":PACKET_ID,"target_sheet":"CO(콜롬비아)","cell_snapshot":snapshot or {"C10":"Hola"}}

def _complete(after="Buenas", cell="C10", supporters=("grammar_fluency","localization_tone"), snapshot=None):
 """Opinions from every role, with two of them supporting one change."""
 base={"sheet":"CO(콜롬비아)","cell":cell,"constraint_status":"pass","packet_id":PACKET_ID}
 opinions=[{**base,"role":r,"finding_id":f"filler-{r}","stance":"review","after":None} for r in SPECIALIST_ROLES]
 opinions+=[{**base,"role":r,"finding_id":"co-c10","stance":"support","after":after,"rule_ids":["co-1"],
             "reason":f"{r} 근거"} for r in supporters]
 return merge_subjective_opinions(opinions,packet=_packet(snapshot))

def _workbook(tmp, value="Hola", cell="C10"):
 path=Path(tmp)/"s.xlsx"; wb=openpyxl.Workbook(); ws=wb.active; ws.title="CO(콜롬비아)"; ws[cell]=value; wb.save(path)
 return path

class ReviewReportBuilderTests(unittest.TestCase):
 def test_builds_pending_report_without_mutating_source(self):
  with tempfile.TemporaryDirectory() as tmp:
   path=_workbook(tmp)
   manifest, report=build_review_artifacts(path,_complete(),report_id="review-1",source_file_id="story-1")
   self.assertEqual(manifest["changes"][0]["approval_status"],"pending_approval")
   self.assertIn("#### 제안 번역문",report)
   self.assertEqual(openpyxl.load_workbook(path)["CO(콜롬비아)"]["C10"].value,"Hola")

 def test_report_shows_supporting_roles(self):
  with tempfile.TemporaryDirectory() as tmp:
   manifest, report=build_review_artifacts(_workbook(tmp),_complete(),report_id="r",source_file_id="s")
   self.assertEqual(manifest["changes"][0]["supporting_roles"],["grammar_fluency","localization_tone"])
   self.assertIn("supporting_roles:",report); self.assertIn("- grammar_fluency",report)

 def test_dissenting_opinions_survive_into_the_markdown(self):
  with tempfile.TemporaryDirectory() as tmp:
   base={"sheet":"CO(콜롬비아)","cell":"C10","constraint_status":"pass","packet_id":PACKET_ID,
         "finding_id":"co-c10","after":"Buenas"}
   opinions=[{**base,"role":r,"finding_id":f"filler-{r}","stance":"review","after":None} for r in SPECIALIST_ROLES]
   opinions+=[{**base,"role":"grammar_fluency","stance":"support","reason":"문법상 필요"},
              {**base,"role":"semantic_fidelity","stance":"oppose","reason":"원문 의미가 바뀐다"}]
   merged=merge_subjective_opinions(opinions,packet=_packet())
   _, report=build_review_artifacts(_workbook(tmp),merged,report_id="r",source_file_id="s")
   self.assertIn("원문 의미가 바뀐다",report); self.assertIn("`semantic_fidelity` (oppose)",report)

 def test_rejects_a_free_form_proposal_list(self):
  with tempfile.TemporaryDirectory() as tmp:
   with self.assertRaises(TypeError):
    build_review_artifacts(_workbook(tmp),[{"sheet":"CO(콜롬비아)","cell":"C10","after":"Buenas"}],
                           report_id="r",source_file_id="s")

 def test_incomplete_sheet_produces_no_changes_and_says_so(self):
  with tempfile.TemporaryDirectory() as tmp:
   base={"sheet":"CO(콜롬비아)","cell":"C10","constraint_status":"pass","packet_id":PACKET_ID,
         "finding_id":"co-c10","after":"Buenas","stance":"support","rule_ids":["co-1"]}
   merged=merge_subjective_opinions([{**base,"role":"grammar_fluency"},{**base,"role":"localization_tone"}],
                                    packet=_packet())
   manifest, report=build_review_artifacts(_workbook(tmp),merged,report_id="r",source_file_id="s")
   self.assertEqual(manifest["changes"],[]); self.assertEqual(manifest["sheet_status"],"incomplete")
   self.assertIn("시트 검수 미완료",report); self.assertIn("semantic_fidelity",report)

 def test_cell_edited_since_review_is_queued_as_drift(self):
  with tempfile.TemporaryDirectory() as tmp:
   # Reviewed against "Hola" but the workbook now holds something else.
   merged=_complete(snapshot={"C10":"Hola"})
   path=_workbook(tmp,value="Hola editado")
   manifest, _=build_review_artifacts(path,merged,report_id="r",source_file_id="s")
   self.assertEqual(manifest["changes"],[])
   self.assertTrue(any(i["reason"]=="source_drift" for i in manifest["review_context"]["human_review_queue"]))

 def test_accepts_a_review_with_no_proposed_edits(self):
  with tempfile.TemporaryDirectory() as tmp:
   base={"sheet":"CO(콜롬비아)","cell":"C10","constraint_status":"pass","packet_id":PACKET_ID}
   merged=merge_subjective_opinions(
    [{**base,"role":r,"finding_id":f"none-{r}","stance":"review","after":None} for r in SPECIALIST_ROLES],
    packet=_packet())
   manifest, report=build_review_artifacts(_workbook(tmp),merged,report_id="review-empty",source_file_id="story-1",
                                           rag_usage={"semantic_budget":2,"semantic_used":0})
   self.assertEqual(manifest["manifest_schema_version"],2); self.assertEqual(manifest["changes"],[])
   self.assertIn("제안 없음",report); self.assertIn("semantic: 0 / 2",report)

 def test_anchoring_warning_reaches_the_report(self):
  with tempfile.TemporaryDirectory() as tmp:
   merged=_complete(supporters=("semantic_fidelity","story_and_ui_coherence"))
   object.__setattr__(merged,"anchoring",[{"role":"story_and_ui_coherence","echoes_role":"semantic_fidelity",
                                           "matched":8,"direction":"subset","detail":"독립 지지 0건"}])
   _, report=build_review_artifacts(_workbook(tmp),merged,report_id="r",source_file_id="s")
   self.assertIn("관점 독립성 경고",report); self.assertIn("독립 지지 0건",report)

 def test_deterministic_proposal_needs_no_consensus(self):
  with tempfile.TemporaryDirectory() as tmp:
   base={"sheet":"CO(콜롬비아)","cell":"C10","constraint_status":"pass","packet_id":PACKET_ID}
   merged=merge_subjective_opinions(
    [{**base,"role":r,"finding_id":f"none-{r}","stance":"review","after":None} for r in SPECIALIST_ROLES],
    packet=_packet(),
    deterministic_proposals=[{"finding_id":"hard-1","sheet":"CO(콜롬비아)","cell":"C10","after":"Buenas",
                              "rule_ids":["glossary"],"reason":"용어집 불일치"}])
   manifest, _=build_review_artifacts(_workbook(tmp),merged,report_id="r",source_file_id="s")
   self.assertEqual(len(manifest["changes"]),1)
   self.assertEqual(manifest["changes"][0]["origin"],"deterministic_hard_rule")

if __name__=="__main__": unittest.main()
