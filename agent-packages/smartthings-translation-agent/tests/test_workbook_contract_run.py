"""Contract/run tests use synthetic inputs only; no API or delivery workbooks.

Value, structure and formatting authorities are source.xlsx unless a test
explicitly supplies an independent authority. Each test states its allowed diff.
"""
import json
import os
import sys
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import workbook_run as run
from workbook_contract import (ContractError, InputDrift, atomic_json, capture_authority,
    check_inputs, digest, file_sha256, load_contract, save_contract, validate_contract)
from workbook_verifier import (VerificationError, facts, diff_facts, verify_artifact, validate_record)
from workbook_mutation_guard import delete_cols_with_manifest, save_verified_atomic
from workbook_apply_edits import (prepare_edit_contract, apply_edits_contract,
    contract_writer_identity, _build_contract_edits)


@pytest.fixture
def planned(tmp_path):
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "US"
    sheet["C7"] = " 013 old "
    sheet["C8"] = "=1+2"
    sheet["D7"] = CellRichText(TextBlock(InlineFont(color="0000FF"), "Brand"), " tail ")
    sheet.row_dimensions[5].height = 17
    sheet["B12"].border = Border(bottom=Side(style="thin"))
    sheet["E7"].comment = Comment("retain", "reviewer")
    book.save(source)
    book.close()
    edits = [{"sheet": "US", "cell": "C7", "before": " 013 old ", "after": " 013 new "}]
    preview = prepare_edit_contract(source, edits, work_root=tmp_path / "work",
                                    delivery_root=tmp_path / "delivery", work_id="pilot")
    contract_path = Path(preview["contract"])
    approval = tmp_path / "approval.json"
    atomic_json(approval, {"approved": True, "approved_by": "test-user",
                           "contract_sha256": preview["contract_sha256"]})
    return source, edits, contract_path, approval


def execute(planned, build=_build_contract_edits):
    return run.execute_run(planned[2], build, writer=contract_writer_identity(), approval_path=planned[3])


def edit_contract(planned, mutate):
    path = planned[2]
    contract = load_contract(path)
    mutate(contract)
    validate_contract(contract)
    atomic_json(path, contract)
    atomic_json(planned[3], {"approved": True, "approved_by": "test-user", "contract_sha256": file_sha256(path)})
    return contract


def candidate(planned, tmp_path, mutate=None):
    contract = load_contract(planned[2])
    book = _build_contract_edits(contract, contract["artifacts"][0], None)
    try:
        if mutate:
            mutate(book)
        path = tmp_path / "candidate.xlsx"
        book.save(path)
        return path
    finally:
        book.close()


def test_pilot_end_to_end_and_resume_keeps_original_and_revision(planned):
    original = file_sha256(planned[0])
    first = apply_edits_contract(*planned[:2], planned[2], planned[3])
    assert first["status"] == "ok"
    assert first["artifact_status"] == "draft"
    second = apply_edits_contract(*planned[:2], planned[2], planned[3])
    assert first == second
    assert file_sha256(planned[0]) == original
    book = openpyxl.load_workbook(first["revised"], rich_text=True, data_only=False)
    assert book["US"]["C7"].value == " 013 new "
    assert book["US"]["C8"].value == "=1+2"
    assert isinstance(book["US"]["D7"].value, CellRichText)
    assert book["US"].row_dimensions[5].height == 17
    book.close()
    assert len(list((Path(first["revision_manifest"]).parent).glob("*.json"))) == 1


def test_no_approval_never_calls_writer_or_publishes(planned):
    result = run.execute_run(planned[2], lambda *args: pytest.fail("unapproved writer"), writer=contract_writer_identity())
    assert result["status"] == "awaiting_approval"
    contract = load_contract(planned[2])
    assert not Path(contract["delivery_root"]).exists()
    assert execute(planned)["status"] == "completed"


@pytest.mark.parametrize("axis,mutate", [
    ("values", lambda b: setattr(b["US"]["C8"], "value", "wrong")),
    ("layout", lambda b: setattr(b["US"].row_dimensions[5], "height", 77)),
    ("styles", lambda b: setattr(b["US"]["B12"], "font", Font(bold=True))),
    ("rich_text", lambda b: setattr(b["US"]["D7"], "value", str(b["US"]["D7"].value))),
    ("annotations", lambda b: setattr(b["US"]["E7"], "comment", None)),
])
def test_verifier_rejects_each_undeclared_axis(planned, tmp_path, axis, mutate):
    path = candidate(planned, tmp_path, mutate)
    with pytest.raises(VerificationError) as error:
        verify_artifact(planned[2], "edited", path)
    assert error.value.report["checks"][axis] == "unexpected_diff"
    assert any(x["axis"] == axis and x["path"][0] == "sheets" for x in error.value.report["unexpected"])


def test_missing_declared_diff_is_not_a_pass(planned):
    with pytest.raises(VerificationError) as error:
        verify_artifact(planned[2], "edited", planned[0])
    assert error.value.report["checks"]["values"] == "missing_declared_diff"


def test_exact_property_allowance_does_not_allow_another_row(planned, tmp_path):
    path = candidate(planned, tmp_path, lambda b: setattr(b["US"].row_dimensions[5], "height", 21))
    allowed = [x for x in diff_facts(facts(planned[0]), facts(path)) if x["axis"] == "layout"]
    edit_contract(planned, lambda c: c["artifacts"][0]["allowed_diffs"].extend(allowed))
    assert verify_artifact(planned[2], "edited", path)["status"] == "passed"
    book = openpyxl.load_workbook(path, rich_text=True)
    book["US"].row_dimensions[6].height = 99
    book.save(path); book.close()
    with pytest.raises(VerificationError):
        verify_artifact(planned[2], "edited", path)


def test_separate_structure_authority_and_noop_contract(planned, tmp_path):
    other = tmp_path / "layout.xlsx"
    book = openpyxl.load_workbook(planned[0], rich_text=True)
    book["US"].row_dimensions[5].height = 44
    book.save(other); book.close()
    def change(c):
        c["authorities"]["layout"] = capture_authority(other)
        c["artifacts"][0]["authorities"]["structure"] = {"authority": "layout", "file": "."}
    contract = edit_contract(planned, change)
    path = candidate(planned, tmp_path, lambda b: setattr(b["US"].row_dimensions[5], "height", 44))
    assert verify_artifact(planned[2], "edited", path)["status"] == "passed"
    contract["no_op_files"] = ["edited"]
    with pytest.raises(ContractError):
        validate_contract(contract)


def test_noop_wrong_authority_is_rejected(planned, tmp_path):
    def noop(c):
        c["artifacts"][0]["allowed_diffs"] = []
        c["no_op_files"] = ["edited"]
    edit_contract(planned, noop)
    assert verify_artifact(planned[2], "edited", planned[0])["status"] == "passed"
    bad = candidate(planned, tmp_path, lambda b: setattr(b["US"].row_dimensions[5], "height", 99))
    with pytest.raises(VerificationError):
        verify_artifact(planned[2], "edited", bad)


def test_directory_inventory_detects_add_remove_and_change(tmp_path):
    directory = tmp_path / "inputs"; directory.mkdir()
    a = directory / "a.txt"; a.write_text("a")
    contract = {"authorities": {"input": capture_authority(directory)}}
    check_inputs(contract)
    b = directory / "b.txt"; b.write_text("b")
    with pytest.raises(InputDrift): check_inputs(contract)
    b.unlink(); a.write_text("changed")
    with pytest.raises(InputDrift): check_inputs(contract)
    a.unlink()
    with pytest.raises(InputDrift): check_inputs(contract)


@pytest.mark.parametrize("change", [
    lambda c: c.update(staging_root=c["delivery_root"]),
    lambda c: c.update(staging_root=str(Path(c["delivery_root"]) / "child")),
    lambda c: c["artifacts"][0].update(output="../escape.xlsx"),
    lambda c: c["artifacts"][0]["authorities"].pop("formatting"),
    lambda c: c["writer"].update(cost="api"),
])
def test_invalid_contract_rejected(planned, change):
    contract = load_contract(planned[2]); change(contract)
    with pytest.raises(ContractError): validate_contract(contract)


def test_record_rejects_tampering_and_old_verifier(planned, tmp_path):
    path = candidate(planned, tmp_path)
    record = verify_artifact(planned[2], "edited", path)
    validate_record(planned[2], "edited", path, record)
    for mutate in (lambda r: r["checks"].pop("styles"), lambda r: r.update(verifier_version="old"),
                   lambda r: r.update(contract_sha256="other"), lambda r: r.update(artifact_id="other")):
        bad = json.loads(json.dumps(record)); mutate(bad)
        with pytest.raises(ContractError): validate_record(planned[2], "edited", path, bad)
    path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises(ContractError): validate_record(planned[2], "edited", path, record)


def test_source_drift_blocks_without_writing(planned):
    planned[0].write_bytes(planned[0].read_bytes() + b"drift")
    result = execute(planned, lambda *args: pytest.fail("drifted input used"))
    assert result["status"] == "blocked"
    assert result["error"]["kind"] == "input_drift"


def test_contract_change_invalidates_resume_and_approval(planned):
    execute_result = run.execute_run(planned[2], _build_contract_edits, writer=contract_writer_identity())
    assert execute_result["status"] == "awaiting_approval"
    edit_contract(planned, lambda c: c.update(recovery={"max_attempts": 2}))
    with pytest.raises(ContractError, match="Contract changed"):
        execute(planned)


def test_limited_format_recovery_uses_fresh_inputs(planned):
    calls = []
    def build(c, a, failure):
        calls.append(failure)
        book = _build_contract_edits(c, a, failure)
        if len(calls) == 1: book["US"].row_dimensions[5].height = 88
        return book
    state = execute(planned, build)
    assert state["status"] == "completed"
    assert state["attempts"]["edited"] == 2
    assert calls[1]["kind"] == "format_or_structure_mismatch"
    events = [json.loads(x)["event"] for x in (planned[2].parent / "events.jsonl").read_text().splitlines()]
    assert "recovery_scheduled" in events


def test_same_failure_stops_and_resume_does_not_reset_budget(planned):
    calls = []
    def build(c, a, f):
        calls.append(1)
        book = _build_contract_edits(c, a, f); book["US"].row_dimensions[5].height = 88
        return book
    assert execute(planned, build)["status"] == "failed"
    assert execute(planned, build)["status"] == "failed"
    assert len(calls) == 2
    assert not (Path(load_contract(planned[2])["delivery_root"]) / "pilot").exists()


def test_different_failures_still_hit_total_attempt_limit(planned):
    calls = []
    def build(c, a, f):
        calls.append(1)
        book = _build_contract_edits(c, a, f); book["US"].row_dimensions[5].height = 80+len(calls)
        return book
    state = execute(planned, build)
    assert state["status"] == "failed"
    assert len(calls) == 3


def test_value_mismatch_is_not_automatically_retried(planned):
    def build(c, a, f):
        book = _build_contract_edits(c, a, f); book["US"]["C8"] = "unexpected"
        return book
    result = execute(planned, build)
    assert result["status"] == "blocked"
    assert result["attempts"]["edited"] == 1


def test_crash_after_staging_resumes_without_regenerating(planned):
    original = run._checkpoint
    def crash(root, state, event, **detail):
        if event == "artifact_verified": raise KeyboardInterrupt()
        return original(root, state, event, **detail)
    with patch.object(run, "_checkpoint", crash), pytest.raises(KeyboardInterrupt):
        execute(planned)
    # Delete the non-authoritative cache; journal and staged data suffice.
    (planned[2].parent / "state.json").unlink()
    state = execute(planned, lambda *args: pytest.fail("must reuse independently reverified staging"))
    assert state["status"] == "completed"
    assert state["attempts"]["edited"] == 1


def test_crash_after_publish_reconciles_completion(planned):
    original = run._checkpoint
    def crash(root, state, event, **detail):
        if event == "completed": raise KeyboardInterrupt()
        return original(root, state, event, **detail)
    with patch.object(run, "_checkpoint", crash), pytest.raises(KeyboardInterrupt): execute(planned)
    assert execute(planned, lambda *args: pytest.fail("already published"))["status"] == "completed"


def test_batch_failure_publishes_nothing(planned):
    def add(c):
        second = json.loads(json.dumps(c["artifacts"][0])); second.update(id="second", output="second.xlsx")
        c["artifacts"].append(second)
    edit_contract(planned, add)
    def build(c, a, f):
        book = _build_contract_edits(c, a, f)
        if a["id"] == "second": book["US"]["C8"] = "bad"
        return book
    state = execute(planned, build)
    assert state["status"] == "blocked"
    assert state["verified"] == ["edited"]
    assert not (Path(load_contract(planned[2])["delivery_root"]) / "pilot").exists()


def test_input_drift_immediately_before_promotion_is_rejected(planned):
    original = run.promote_verified
    def drift(contract_path, records):
        planned[0].write_bytes(planned[0].read_bytes() + b"drift")
        return original(contract_path, records)
    with patch.object(run, "promote_verified", drift): state = execute(planned)
    assert state["status"] == "blocked"
    assert not (Path(load_contract(planned[2])["delivery_root"]) / "pilot").exists()


def test_guard_cannot_save_contract_directly_to_delivery(planned):
    c = load_contract(planned[2]); book = openpyxl.Workbook()
    try:
        with pytest.raises(ValueError, match="staging"):
            save_verified_atomic(book, Path(c["delivery_root"]) / "bad.xlsx", lambda p: {}, contract=c)
    finally: book.close()


def test_validation_conditions_are_compared_not_just_ranges(planned, tmp_path):
    path = candidate(planned, tmp_path)
    book = openpyxl.load_workbook(path, rich_text=True)
    validation = DataValidation(type="whole", operator="between", formula1=1, formula2=10)
    validation.add("A1:A5"); book["US"].add_data_validation(validation)
    book.save(path); book.close()
    first = facts(path)
    book = openpyxl.load_workbook(path, rich_text=True)
    book["US"].data_validations.dataValidation[0].formula2 = 20
    book.save(path); book.close()
    assert any(x["path"][-1] == "data_validations" for x in diff_facts(first, facts(path)))


def test_column_deletion_reindexes_dimensions_merges_and_audits(tmp_path):
    book = openpyxl.Workbook(); sheet = book.active
    sheet["C2"] = "deleted"; sheet["C2"].comment = Comment("review", "author")
    sheet["F2"] = "keep"; sheet.column_dimensions["F"].width = 31
    sheet.merge_cells("F4:G4"); sheet["F4"] = "merged"
    removed = delete_cols_with_manifest(sheet, 3, 2)
    assert removed["annotations"]["C2"]["text"] == "review"
    path = tmp_path / "deleted.xlsx"; book.save(path); book.close()
    reopened = openpyxl.load_workbook(path)
    assert reopened.active["D2"].value == "keep"
    assert reopened.active.column_dimensions["D"].width == 31
    assert str(next(iter(reopened.active.merged_cells.ranges))) == "D4:E4"
    assert reopened.active["D4"].value == "merged"
    reopened.close()


def test_column_deletion_rejects_crossing_merge_without_mutation():
    book = openpyxl.Workbook(); sheet = book.active
    sheet.merge_cells("B2:D2"); sheet["B2"] = "keep"
    with pytest.raises(ValueError): delete_cols_with_manifest(sheet, 3, 1)
    assert sheet["B2"].value == "keep"
    assert str(next(iter(sheet.merged_cells.ranges))) == "B2:D2"
    book.close()


def test_separate_process_crash_and_resume(planned):
    import subprocess
    scripts = str(Path(__file__).resolve().parents[1] / "scripts")
    program = """
import os, sys
sys.path.insert(0, sys.argv[1])
import workbook_run as run
from workbook_apply_edits import _build_contract_edits, contract_writer_identity
original = run._checkpoint
def crash(root, state, event, **detail):
    if event == 'artifact_verified': os._exit(73)
    return original(root, state, event, **detail)
run._checkpoint = crash
run.execute_run(sys.argv[2], _build_contract_edits, writer=contract_writer_identity(), approval_path=sys.argv[3])
"""
    crashed = subprocess.run([sys.executable, "-c", program, scripts, str(planned[2]), str(planned[3])], capture_output=True)
    assert crashed.returncode == 73, crashed.stderr.decode()
    state = execute(planned, lambda *args: pytest.fail("persisted staged file should be reverified"))
    assert state["status"] == "completed"
    assert state["attempts"]["edited"] == 1


def test_live_process_lock_blocks_duplicate_execution(planned):
    import subprocess
    scripts = str(Path(__file__).resolve().parents[1] / "scripts")
    with run.exclusive_lock(planned[2].parent / ".run.lock"):
        child = subprocess.run([sys.executable, "-c", """
import sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from workbook_run import exclusive_lock
from workbook_contract import ContractError
try:
    with exclusive_lock(Path(sys.argv[2])): sys.exit(2)
except ContractError:
    sys.exit(0)
""", scripts, str(planned[2].parent / ".run.lock")], capture_output=True)
        assert child.returncode == 0, child.stderr.decode()
    assert execute(planned)["status"] == "completed"


def test_cli_execute_and_resume(planned):
    import subprocess
    script = Path(__file__).resolve().parents[1] / "scripts/workbook_apply_edits.py"
    args = [sys.executable, str(script), str(planned[0]), json.dumps(planned[1]),
            "--run-contract", str(planned[2]), "--approval", str(planned[3]), "--json"]
    first = subprocess.run(args, capture_output=True, text=True)
    assert first.returncode == 0, first.stdout + first.stderr
    second = subprocess.run(args, capture_output=True, text=True)
    assert second.returncode == 0, second.stdout + second.stderr
    assert json.loads(first.stdout) == json.loads(second.stdout)


def test_stale_approval_does_not_run(planned):
    atomic_json(planned[3], {"approved": True, "approved_by": "test-user", "contract_sha256": "stale"})
    state = execute(planned, lambda *a: pytest.fail("stale approval"))
    assert state["status"] == "awaiting_approval"


def test_writer_version_change_is_blocked(planned):
    identity = contract_writer_identity(); identity["version"] = "new-version"
    state = run.execute_run(planned[2], lambda *a: pytest.fail("wrong writer"), writer=identity, approval_path=planned[3])
    assert state["status"] == "blocked"


def test_transient_local_io_retries_once(planned):
    import errno
    calls = []
    def build(c, a, f):
        calls.append(1)
        if len(calls) == 1: raise OSError(errno.EAGAIN, "temporary local busy")
        return _build_contract_edits(c, a, f)
    assert execute(planned, build)["status"] == "completed"
    assert len(calls) == 2


def test_completed_output_tampering_never_silently_repairs(planned):
    state = execute(planned)
    output = Path(state["result"]["outputs"]["edited"])
    changed = output.read_bytes()+b"tampered"
    output.write_bytes(changed)
    assert execute(planned)["status"] == "blocked"
    assert output.read_bytes() == changed


def test_batch_publish_interruption_exposes_no_partial_version(planned):
    def add(c):
        second = json.loads(json.dumps(c["artifacts"][0])); second.update(id="second", output="nested/second.xlsx")
        c["artifacts"].append(second)
    contract = edit_contract(planned, add)
    original = run.shutil.copyfile
    calls = []
    def copy(src, dst):
        calls.append(1)
        if len(calls) == 2: raise KeyboardInterrupt()
        return original(src, dst)
    with patch.object(run.shutil, "copyfile", copy), pytest.raises(KeyboardInterrupt): execute(planned)
    delivery = Path(contract["delivery_root"])
    assert not (delivery / "pilot").exists()
    assert not list(delivery.glob(".preparing-*"))
    state = execute(planned, lambda *a: pytest.fail("both artifacts already verified"))
    assert state["status"] == "completed"
    assert set(state["result"]["outputs"]) == {"edited", "second"}


def test_contract_guard_rejects_a_record_for_other_contract(planned):
    contract = load_contract(planned[2]); book = openpyxl.Workbook()
    target = Path(contract["staging_root"]) / "wrong.xlsx"
    def fake(path):
        return {"status": "passed", "checks": {a: "passed" for a in ("values", "layout", "styles", "rich_text", "annotations")},
                "staged_file_sha256": file_sha256(path), "contract_content_sha256": "other"}
    try:
        with pytest.raises(ValueError, match="record"):
            save_verified_atomic(book, target, fake, contract=contract)
        assert not target.exists()
    finally: book.close()


def test_rich_target_edit_is_explicit_and_non_target_runs_survive(tmp_path):
    source = tmp_path / "rich.xlsx"
    book = openpyxl.Workbook(); sheet = book.active
    sheet["A1"] = CellRichText(TextBlock(InlineFont(b=True), "old"))
    sheet["B1"] = CellRichText(TextBlock(InlineFont(i=True), "keep"))
    book.save(source); book.close()
    edits = [{"sheet": "Sheet", "cell": "A1", "before": "old", "after": "new"}]
    preview = prepare_edit_contract(source, edits, work_root=tmp_path/"work", delivery_root=tmp_path/"out", work_id="rich")
    approval = tmp_path / "approval.json"
    atomic_json(approval, {"approved": True, "approved_by": "test", "contract_sha256": preview["contract_sha256"]})
    assert {x["axis"] for x in preview["allowed_diffs"]} == {"values", "rich_text"}
    result = apply_edits_contract(source, edits, Path(preview["contract"]), approval)
    assert result["status"] == "ok"
    book = openpyxl.load_workbook(result["revised"], rich_text=True)
    assert book.active["A1"].value == "new"
    assert book.active["B1"].value[0].font.i
    book.close()


def test_chart_anchor_change_with_same_count_is_detected(tmp_path):
    from openpyxl.chart import BarChart, Reference
    book = openpyxl.Workbook(); sheet = book.active
    sheet.append([1]); sheet.append([2])
    chart = BarChart(); chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2))
    sheet.add_chart(chart, "D1")
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    book.save(first); book.close()
    book = openpyxl.load_workbook(first)
    book.active._charts[0].anchor._from.col = 6
    book.save(second); book.close()
    assert any(x["axis"] == "layout" and x["path"][0] == "package" for x in diff_facts(facts(first), facts(second)))


def test_batch_manifest_cannot_redirect_verified_output(planned, tmp_path):
    state = execute(planned)
    manifest = Path(load_contract(planned[2])["delivery_root"]) / "pilot/batch.complete.json"
    payload = json.loads(manifest.read_text())
    payload["outputs"]["edited"] = str(tmp_path / "unverified.xlsx")
    atomic_json(manifest, payload)
    assert execute(planned)["status"] == "blocked"


def test_new_run_rejects_prepopulated_staging(planned):
    staging = Path(load_contract(planned[2])["staging_root"])
    staging.mkdir(parents=True); (staging / "unrelated.txt").write_text("keep")
    with pytest.raises(ContractError, match="empty staging"):
        execute(planned)
    assert (staging / "unrelated.txt").read_text() == "keep"


def test_column_delete_rejects_formula_dependency_on_other_sheet():
    book = openpyxl.Workbook(); book.active.title = "Source"
    book.active["C1"] = 7
    book.create_sheet("Dependent")["A1"] = "=Source!C1"
    with pytest.raises(ValueError, match="dependency"):
        delete_cols_with_manifest(book["Source"], 3, 1)
    assert book["Source"]["C1"].value == 7
    book.close()


@pytest.mark.parametrize("value", [3, 1.0, 3.5, True, None])
def test_scalar_edits_are_verified_with_excel_numeric_types(planned, tmp_path, value):
    edits = [{"sheet": "US", "cell": "C7", "after": value}]
    preview = prepare_edit_contract(planned[0], edits, work_root=tmp_path/"scalars",
                                    delivery_root=tmp_path/"scalar-output", work_id="scalars")
    approval = tmp_path / "scalar-approval.json"
    atomic_json(approval, {"approved": True, "approved_by": "test", "contract_sha256": preview["contract_sha256"]})
    result = apply_edits_contract(planned[0], edits, Path(preview["contract"]), approval)
    assert result["status"] == "ok"


def test_lossy_number_is_rejected_during_plan(planned, tmp_path):
    with pytest.raises(ContractError, match="precision"):
        prepare_edit_contract(planned[0], [{"sheet": "US", "cell": "C7", "after": 12345678901234567}],
            work_root=tmp_path/"lossy", delivery_root=tmp_path/"out", work_id="lossy")
    assert not (tmp_path/"lossy/contract.json").exists()
