from __future__ import annotations

import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from glossary_activation_candidates import (  # noqa: E402
    BASIS, find_candidates, markdown_review, to_activation_entries,
)


class _PromptBuilder:
    @staticmethod
    def is_glossary_deactivated(rule: str) -> bool:
        return "비활성화" in (rule or "")


class _Checker:
    """The GlossaryChecker members the scan uses, without loading the app."""

    def __init__(self, terms: dict[str, str], rules: dict[str, str] | None = None):
        self.glossary = {key: {"targets": {"스페인어_콜롬비아": value},
                               "rule": (rules or {}).get(key, "")}
                         for key, value in terms.items()}
        self.glossary_map = {key.lower(): key for key in terms}
        self.prompt_builder = _PromptBuilder()
        pattern = "|".join(re.escape(key) for key in sorted(terms, key=len, reverse=True))
        self.glossary_re = re.compile(pattern, re.IGNORECASE)

    def _get_target_val(self, targets, code):
        return targets.get(code)


def _workbook(tmp: Path, rows: dict[int, tuple[str, str]], name="Story_012_ES_co_v1.xlsx") -> Path:
    path = tmp / name
    wb = openpyxl.Workbook()
    source = wb.active
    source.title = "US(미국)"
    target = wb.create_sheet("CO(콜롬비아)")
    for row, (src, tgt) in rows.items():
        source.cell(row, 3).value = src
        target.cell(row, 3).value = tgt
    wb.save(path)
    return path


def _scan(tmp: Path, rows, terms=None, rules=None):
    checker = _Checker(terms or {"Safe": "Safe"}, rules)
    return find_candidates(checker, _workbook(tmp, rows), "US(미국)", "CO(콜롬비아)", "스페인어_콜롬비아")


class ActivationCandidateTests(unittest.TestCase):
    def test_a_lowercase_use_of_a_feature_term_becomes_a_candidate(self):
        """ES_CO story_012 C7: 'a safe home' matched the product term `Safe`."""
        with tempfile.TemporaryDirectory() as raw:
            found = _scan(Path(raw), {7: ("The comfort of a safe home", "La comodidad de un hogar Safe")})
            self.assertEqual(len(found), 1)
            entry = found[0]
            self.assertEqual((entry["story"], entry["cell"], entry["source_term"]), ("012", "C7", "Safe"))
            self.assertFalse(entry["active"])
            self.assertFalse(entry["confirmed"])
            self.assertEqual(entry["activation_basis"], BASIS)
            self.assertEqual(entry["evidence"]["matched_surfaces"], ["safe"])

    def test_a_capitalised_use_anywhere_in_the_cell_keeps_the_term_active(self):
        """Activation is keyed by (story, cell, term), so one real use covers the cell."""
        with tempfile.TemporaryDirectory() as raw:
            found = _scan(Path(raw), {7: ("SmartThings Safe keeps a safe home", "SmartThings Safe")})
            self.assertEqual(found, [])

    def test_a_term_the_glossary_already_switched_off_is_not_a_candidate(self):
        """Nothing to decide: the resolver never demands a globally deactivated term.

        Ten of the seventeen terms the first ES_CO scan reported were already marked
        `비활성화`, which is why 102 candidates shrank to the 25 that actually block
        the delivered text.
        """
        with tempfile.TemporaryDirectory() as raw:
            found = _scan(Path(raw), {7: ("a safe home", "un hogar seguro")},
                          rules={"Safe": "대괄호 제외, 비활성화"})
            self.assertEqual(found, [])

    def test_an_all_lowercase_glossary_key_is_never_a_candidate(self):
        """Such a key says nothing about product-vs-ordinary usage."""
        with tempfile.TemporaryDirectory() as raw:
            found = _scan(Path(raw), {7: ("turn on the widget", "el widget")}, terms={"widget": "widget"})
            self.assertEqual(found, [])

    def test_placeholder_and_empty_rows_are_skipped(self):
        with tempfile.TemporaryDirectory() as raw:
            found = _scan(Path(raw), {7: ("x", "x"), 8: ("", "")})
            self.assertEqual(found, [])

    def test_evidence_records_whether_the_translation_used_the_term(self):
        """Corroboration for the reviewer; it is never the decision by itself."""
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            used = _scan(tmp, {7: ("a safe home", "un hogar Safe")})[0]
            self.assertTrue(used["evidence"]["target_uses_glossary_term"])
            unused = _scan(tmp, {8: ("a safe home", "un hogar seguro")})[0]
            self.assertFalse(unused["evidence"]["target_uses_glossary_term"])

    def test_the_decision_does_not_depend_on_the_target_locale(self):
        """Whether the source word is the product term is a fact about the source.

        One confirmation therefore serves every target sheet in the source group;
        the target sheet only ever supplies corroboration for the reviewer.
        """
        with tempfile.TemporaryDirectory() as raw:
            tmp = Path(raw)
            rows = {7: ("The comfort of a safe home", "La comodidad de un hogar seguro")}
            checker = _Checker({"Safe": "Safe"})
            workbook = _workbook(tmp, rows)
            with_locale = find_candidates(checker, workbook, "US(미국)", "CO(콜롬비아)", "스페인어_콜롬비아")
            without = find_candidates(checker, workbook, "US(미국)")
            self.assertEqual(to_activation_entries(with_locale, confirmed_only=False),
                             to_activation_entries(without, confirmed_only=False))
            self.assertEqual(without[0]["evidence"]["glossary_target"], "")

    def test_only_confirmed_candidates_become_activation_entries(self):
        """A candidate is a question. Deactivating a term nobody checked drops translations."""
        candidates = [
            {"story": "012", "cell": "C7", "source_term": "Safe", "active": False, "confirmed": True},
            {"story": "012", "cell": "C8", "source_term": "Safe", "active": False, "confirmed": False},
        ]
        entries = to_activation_entries(candidates)
        self.assertEqual([entry["cell"] for entry in entries], ["C7"])
        self.assertNotIn("evidence", entries[0])
        self.assertNotIn("confirmed", entries[0])


class ActivationManifestContractTests(unittest.TestCase):
    def test_the_resolver_accepts_what_this_tool_emits(self):
        """The emitted file must load through OccurrenceActivationManifest unchanged."""
        sys.path.insert(0, str(ROOT.parents[1] / "src"))
        from translation_web_app.constraint_resolver import OccurrenceActivationManifest

        entries = to_activation_entries([
            {"story": "12", "cell": "c7", "source_term": "Safe", "active": False, "confirmed": True},
        ])
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "manifest.json"
            path.write_text(json.dumps({"entries": entries}, ensure_ascii=False), encoding="utf-8")
            manifest = OccurrenceActivationManifest.from_file(path)
        found = manifest.lookup("012", "C7", "safe")
        self.assertIsNotNone(found)
        self.assertFalse(found["active"])


class ReviewMarkdownTests(unittest.TestCase):
    def test_review_lists_each_candidate_under_its_term(self):
        markdown = markdown_review([{
            "story": "012", "cell": "C7", "source_term": "Safe", "active": False, "confirmed": False,
            "evidence": {"source_text": "The comfort of a safe home", "matched_surfaces": ["safe"],
                         "glossary_target": "Safe", "target_text": "un hogar seguro",
                         "target_uses_glossary_term": False},
        }])
        self.assertIn("## `Safe`", markdown)
        self.assertIn("story_012", markdown)
        self.assertIn("confirmed: true", markdown)


if __name__ == "__main__":
    unittest.main()
