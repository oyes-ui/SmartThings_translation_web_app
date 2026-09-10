from pathlib import Path
import asyncio
import json
import sys
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import workbook_batch as batch
from workbook_add_target_sheet import save_prepared
from workbook_contract import atomic_json, file_sha256
from workflow_approval import managed_write, publish
from glossary_sync import sync_default
from test_agent_staged_batch import packet, cell_review


def book(path, target="DE(독일)"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "US(미국)"
    ws["C3"] = "en_US"
    ws["C5"] = "story_001"
    ws["C7"] = "Actual"
    ws["C8"] = "x"
    ws.row_dimensions[7].height = 42
    ws.sheet_properties.tabColor = "FF0000"
    ws["C7"].fill = PatternFill("solid", fgColor="FFFF00")
    dest = wb.copy_worksheet(ws)
    dest.title = target
    wb.save(path)
    wb.close()
    return path


def config(tmp_path, **overrides):
    source = book(tmp_path / "story.xlsx")
    glossary = tmp_path / "g.csv"
    glossary.write_text("key,Rule,English\n,,en_US\nLng,,영어_미국\nA,,A\n")
    return {"app_root": str(ROOT.parents[1]), "obsidian_dir": str(tmp_path / "vault"),
            "jobs": [{"workbook": str(source), "sheet": "DE(독일)", "glossary": str(glossary),
                      "glossary_approved": True, "review": False, **overrides}]}


def fake_translate(settings, source, app):
    return {"status": "ok", "excel_path": str(source)}


def test_prepare_independent_facts_preserve_layout_and_clear_only_new(tmp_path):
    source = book(tmp_path / "input.xlsx")
    original = file_sha256(source)
    result = save_prepared(source, tmp_path / "prepared.xlsx", "DE(독일)", "CA(fr)", "fr_CA")
    wb = openpyxl.load_workbook(result["output"], rich_text=True)
    assert wb["CA(fr)"]["C7"].value is None
    assert wb["CA(fr)"]["C8"].value == "x"
    assert wb["CA(fr)"].row_dimensions[7].height == 42
    assert wb["DE(독일)"]["C7"].value == "Actual"
    assert file_sha256(source) == original
    assert Path(result["contract"]).is_file()
    wb.close()


def test_unsupported_copy_rejected_and_existing_target_skipped(tmp_path):
    source = book(tmp_path / "input.xlsx")
    wb = openpyxl.load_workbook(source)
    chart = BarChart()
    chart.add_data(Reference(wb.active, min_col=3, min_row=7, max_row=8))
    wb["DE(독일)"].add_chart(chart, "F5")
    wb.save(source)
    wb.close()
    with pytest.raises(ValueError, match="cannot preserve"):
        save_prepared(source, tmp_path / "bad.xlsx", "DE(독일)", "NEW", "test")
    assert not (tmp_path / "verified").exists()


def test_snapshots_settings_isolation_and_api_resume(tmp_path):
    c = config(tmp_path)
    other = tmp_path / "other.csv"
    other.write_text("custom glossary")
    c["jobs"].append({**c["jobs"][0], "sheet": "US(other)", "glossary": str(other),
                      "source_sheet": "DE(독일)", "sheet_langs": {"US(other)": {"code": "영어_미국", "lang": "English"}}})
    wb = openpyxl.load_workbook(c["jobs"][0]["workbook"])
    wb.copy_worksheet(wb["DE(독일)"]).title = "US(other)"
    wb.save(c["jobs"][0]["workbook"])
    wb.close()
    root = tmp_path / "work"
    m = batch.prepare(c, root)
    assert len({batch.settings_for(j)["glossary"] for j in m["jobs"]}) == 2
    Path(c["jobs"][0]["glossary"]).write_text("changed default")
    assert "changed default" not in Path(batch.settings_for(m["jobs"][0])["glossary"]).read_text()
    calls = []
    def fake(s, w, a):
        calls.append(s["sheet"])
        if s["sheet"] == "US(other)":
            raise RuntimeError("fake paid failure")
        return fake_translate(s,w,a)
    m = batch.advance(root / "workflow.json", pipeline=True, translator=fake)
    assert [j["state"] for j in m["jobs"]] == ["draft_completed", "translation_error"]
    batch.advance(root / "workflow.json", pipeline=True, translator=fake)
    assert len(calls) == 2
    with pytest.raises(ValueError, match="approval"):
        batch.advance(root / "workflow.json", retry_jobs=[m["jobs"][1]["job_id"]], translator=fake)


def test_interrupted_api_requires_recovery_no_duplicate(tmp_path):
    root = tmp_path / "work"
    m = batch.prepare(config(tmp_path), root)
    m["jobs"][0]["state"] = "translation_started"
    atomic_json(root / "workflow.json", m)
    def forbidden(*args):
        raise AssertionError("Duplicate paid call")
    m = batch.advance(root / "workflow.json", pipeline=True, translator=forbidden)
    assert m["jobs"][0]["state"] == "api_recovery_required"
    assert "api_recovery_required" in Path(m["reports"]["report"]).read_text()


def test_no_preparation_or_api_for_existing_review(tmp_path):
    root = tmp_path / "work"
    c = config(tmp_path, translate=False, prepare={"template_sheet": "MISSING", "lang_code": "de"})
    batch.prepare(c, root)
    with patch("workbook_add_target_sheet.save_prepared", side_effect=AssertionError("unnecessary prep")):
        m = batch.advance(root / "workflow.json", translator=lambda *_: pytest.fail("unnecessary API"))
    assert m["jobs"][0]["state"] == "draft_completed"


@pytest.mark.parametrize("revise", [False, True])
def test_draft_agent_stages_reports_and_notes_end_to_end(tmp_path, revise):
    c = config(tmp_path, review=True)
    root = tmp_path / "work"
    m = batch.prepare(c, root)
    async def fake_packet(workbook, sheet, **kw):
        assert kw["source_sheet"] == "US(미국)"
        return packet(sheet, "pkt")
    passing = lambda cell, after: {"status": "pass", "blocked": False, "normalized": after, "violations": [], "review": []}
    with patch("agent_staged_batch.build_packet", fake_packet), patch("agent_staged_batch.resolver_validator", return_value=passing):
        m = batch.advance(root / "workflow.json", pipeline=True, translator=fake_translate)
        job = m["jobs"][0]
        review_dir = Path(batch.read(job["review_manifest"])["jobs"][0]["work_dir"])
        cell = cell_review("pkt")
        if revise:
            cell["cells"][0].update(status="needs_revision", after="Corrected")
        atomic_json(review_dir / "cell_review.json", cell)
        m = batch.advance(root / "workflow.json")
        assert m["jobs"][0]["ready_for_agent"][0]["stage"] == "sheet"
        meta = {"packet_id": "pkt", "status": "completed", "stop_reason": "complete", "model": "fake-agent", "run_id": "stage", "executed_at": "2026-09-10T00:00:00Z"}
        atomic_json(review_dir / "sheet_review.json", {**meta, "kind": "sheet_consistency_review", "issues": []})
        m = batch.advance(root / "workflow.json")
        assert m["jobs"][0]["ready_for_agent"][0]["stage"] == "lead"
        atomic_json(review_dir / "lead_review.json", {**meta, "kind": "lead_review", "decisions": [{
            "cell": "C7", "finding_id": "lead-C7", "status": "needs_revision" if revise else "pass",
            "after": "Corrected" if revise else None, "reason": "checked", "rule_ids": [], "basis_refs": ["cell:C7"]}]})
        # Only rendering is stubbed: actual gates, merge, candidate IDs and manifest construction run.
        with patch("agent_app_report.render_through_app", return_value="# Detailed\n"):
            m = batch.advance(root / "workflow.json")
        assert m["jobs"][0]["state"] == "completed"
        index = batch.read(m["reports"]["index"])
        assert len(index["changes"]) == int(revise)
        if revise:
            entry = index["changes"][0]
            source = batch.read(entry["source_manifest"])
            assert entry["change"] == source["changes"][0]
            assert "Corrected" in Path(m["reports"]["report"]).read_text()
            finding = entry["change"]["finding_id"]
            assert finding in Path(m["reports"]["report"]).read_text()
            assert any(finding in p.read_text() for p in (root / "reports").glob("review-*.md"))
        else:
            assert "수정 후보 없음" in Path(m["reports"]["report"]).read_text()
        note = Path(m["reports"]["obsidian_report"])
        note.write_text(note.read_text() + "User decision: discuss tomorrow\n")
        m = batch.advance(root / "workflow.json")
        assert "User decision: discuss tomorrow" in note.read_text()


def test_existing_unmanaged_note_preserved(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("My memo")
    with pytest.raises(ValueError):
        managed_write(path, "generated")
    assert path.read_text() == "My memo"


def test_glossary_sync_failure_preserves_old_csv_and_recovery(tmp_path):
    dest = tmp_path / "latest_glossary.csv"
    dest.write_text("old")
    class Store:
        def export_csv(self, path):
            path.write_text("invalid")
    with pytest.raises(ValueError):
        sync_default(Store(), dest)
    assert dest.read_text() == "old"
    assert batch.read(dest.with_suffix(".sync.json"))["status"] == "failed"
    class GoodStore:
        def export_csv(self, path):
            path.write_text("key,rule,en\n,,English\nLng,,en_US\na,,a\n")
    result = sync_default(GoodStore(), dest)
    assert result["status"] == "ok"
    assert result["sha256"] == file_sha256(dest)


def test_persisted_api_receipt_recovers_without_call(tmp_path):
    root = tmp_path / "work"
    m = batch.prepare(config(tmp_path), root)
    job = m["jobs"][0]
    s = batch.settings_for(job)
    atomic_json(Path(job["settings"]).parent / "translation_result.json", {
        "settings_sha256": job["settings_sha256"], "path": s["workbook"], "sha256": file_sha256(s["workbook"])})
    job["state"] = "translation_started"
    atomic_json(root / "workflow.json", m)
    m = batch.advance(root / "workflow.json", pipeline=True, translator=lambda *_: pytest.fail("Duplicate API"))
    assert m["jobs"][0]["state"] == "draft_completed"


def test_frozen_activation_drift_prevents_paid_call(tmp_path):
    c = config(tmp_path)
    activation = tmp_path / "activation.json"
    atomic_json(activation, {"occurrences": []})
    c["jobs"][0]["activation_manifest"] = str(activation)
    root = tmp_path / "work"
    m = batch.prepare(c, root)
    s = batch.settings_for(m["jobs"][0])
    Path(s["activation_manifest"]).write_text("changed")
    m = batch.advance(root / "workflow.json", pipeline=True, translator=lambda *_: pytest.fail("Unsafe API"))
    assert "Frozen input changed" in m["jobs"][0]["error"]


def test_translation_command_keeps_custom_settings_and_skips_audit(tmp_path):
    from types import SimpleNamespace
    c = config(tmp_path)
    root = tmp_path / "work"
    m = batch.prepare(c, root)
    s = batch.settings_for(m["jobs"][0])
    s["activation_manifest"] = "/approved/activation.json"
    calls = []
    def fake(cmd, **kwargs):
        calls.append(cmd)
        return SimpleNamespace(returncode=0, stdout=json.dumps({"status": "ok", "excel_path": s["workbook"]}))
    with patch("workbook_batch.subprocess.run", fake):
        batch.translate(s, Path(s["workbook"]), c["app_root"])
    cmd = calls[0]
    assert "--translate-only" in cmd and "--with-api-audit" not in cmd
    assert "--backtranslation-sheet" not in cmd
    assert cmd[cmd.index("--glossary") + 1] == s["glossary"]
    assert cmd[cmd.index("--activation-manifest") + 1] == s["activation_manifest"]
    assert cmd[cmd.index("--source-sheet") + 1] == s["source_sheet"]


def test_custom_language_review_and_apply_use_identical_source_settings(tmp_path):
    from types import SimpleNamespace
    from agent_sheet_review import build_packet
    import workbook_delivery as delivery
    c = config(tmp_path, translate=False, source_sheet="US(미국)", sheet="Custom target",
               sheet_langs={"Custom target": {"code": "독어_독일", "lang": "German"}})
    source = Path(c["jobs"][0]["workbook"])
    wb = openpyxl.load_workbook(source)
    wb["DE(독일)"].title = "Custom target"
    wb.save(source)
    wb.close()
    Path(c["jobs"][0]["glossary"]).write_text(
        "Key,규칙,영어_미국,독어_독일\nKey,Rule,en_US,de_DE\nLng,Rule,Lng,Lng\nActual,,Actual,Actual\n")
    root = tmp_path / "work"
    m = batch.prepare(c, root)
    job = m["jobs"][0]
    s = batch.settings_for(job)
    packet = asyncio.run(build_packet(Path(s["workbook"]), s["sheet"], glossary=Path(s["glossary"]),
                                     app_root=Path(c["app_root"]), source_sheet=s["source_sheet"], sheet_langs=s["sheet_langs"]))
    assert packet["source_sheet"] == s["source_sheet"]
    assert packet["sheet_langs"][s["sheet"]]["code"] == "독어_독일"
    assert packet["deterministic_evidence"]
    args = SimpleNamespace(workbook=s["workbook"], glossary=s["glossary"], app_root=c["app_root"],
        edits=json.dumps([{"sheet": s["sheet"], "cell": "C7", "before": "Actual", "after": "Actual revised"}]),
        delivery_sheets=s["sheet"], cell_range="C7:C28", workflow_settings=job["settings"],
        dry_run=False, activation_manifest=None)
    with patch.object(delivery.ap, "maybe_reexec_with_app_venv"):
        result = asyncio.run(delivery.run_delivery(args, "story"))
    assert result["status"] == "ok", result
    wb = openpyxl.load_workbook(result["final"], rich_text=True)
    assert str(wb[s["sheet"]]["C7"].value) == "Actual revised"
    wb.close()


def test_help_covers_seven_workflows_without_skill_procedure_duplication():
    guide = (ROOT / "references/workflow-guide.md").read_text()
    for task in ("한국어 원문", "새 용어집", "영어·다국어", "번역 검수", "확정본 일부", "여러 파일", "Excel 서식"):
        assert task in guide
    assert "필요한 입력" in guide and "다음 작업" in guide
    for name in ("st-start.md", "st-help.md"):
        assert "references/workflow-guide.md" in (ROOT / "commands" / name).read_text()


def test_local_failure_after_paid_success_never_repeats_api(tmp_path):
    root = tmp_path / "work"
    batch.prepare(config(tmp_path), root)
    calls = []
    def fake(s, w, app):
        calls.append(w)
        return fake_translate(s, w, app)
    with patch("workbook_batch.validate_draft", side_effect=ValueError("Local check failed")):
        m = batch.advance(root / "workflow.json", pipeline=True, translator=fake)
    assert m["jobs"][0]["state"] == "review_error"
    m = batch.advance(root / "workflow.json", pipeline=True, translator=fake)
    assert m["jobs"][0]["state"] == "draft_completed"
    assert len(calls) == 1


@pytest.mark.parametrize("audit,backtranslation", [(False, False), (True, False), (True, True)])
def test_wrapper_fake_api_receives_activation_and_only_requested_stages(tmp_path, audit, backtranslation):
    from types import SimpleNamespace
    import workbook_translate as translator
    c = config(tmp_path)
    calls = {}
    class FakeChecker:
        def __init__(self, **kwargs):
            calls["init"] = kwargs
        def load_activation_manifest(self, path):
            calls["activation"] = path
        async def run_integrated_pipeline_generator(self, **kwargs):
            calls["pipeline"] = kwargs
            yield {"type": "complete", "excel_path": c["jobs"][0]["workbook"]}
    args = SimpleNamespace(workbook=c["jobs"][0]["workbook"], app_root=c["app_root"],
        glossary=c["jobs"][0]["glossary"], activation_manifest="approved.json", sheet_langs=None,
        sheets="DE(독일)", single_source=True, source_sheet="US(미국)", max_concurrency=1,
        backtranslation_lang=None, backtranslation_sheet=None, with_backtranslation=backtranslation,
        audit_model="unused", translation_model="fake", cell_range="C7:C28", bx=False,
        translate_only=not audit, json=True, verbose=False)
    with patch.object(translator.ap, "maybe_reexec_with_app_venv"), patch("translation_web_app.checker_service.TranslationChecker", FakeChecker):
        result = asyncio.run(translator.run_translate(args))
    assert result["status"] == "ok"
    assert calls["activation"] == "approved.json"
    assert calls["pipeline"]["skip_audit"] is (not audit)
    assert calls["init"]["no_backtranslation"] is (not backtranslation)


def test_standard_source_groups_and_explicit_override_are_config_only(tmp_path):
    c = config(tmp_path)
    c["jobs"] = [{**c["jobs"][0], "sheet": sheet} for sheet in ("US(미국)", "JA(일본)", "DE(독일)")]
    c["jobs"].append({**c["jobs"][-1], "sheet": "FR(프랑스)", "source_sheet": "KR(한국)"})
    m = batch.prepare(c, tmp_path / "work")
    assert [batch.settings_for(j)["source_sheet"] for j in m["jobs"]] == ["KR(한국)", "KR(한국)", "US(미국)", "KR(한국)"]


def test_per_job_output_directory_and_read_only_status_data(tmp_path):
    c = config(tmp_path, output_dir=str(tmp_path / "exports"))
    root = tmp_path / "work"
    batch.prepare(c, root)
    m = batch.advance(root / "workflow.json", pipeline=True, translator=fake_translate)
    output = Path(m["jobs"][0]["draft"])
    assert output.is_relative_to(tmp_path / "exports")
    before = file_sha256(root / "workflow.json")
    assert batch.read(root / "workflow.json")["jobs"][0]["state"] == "draft_completed"
    assert before == file_sha256(root / "workflow.json")


def test_corrupt_review_report_does_not_hide_other_job_states(tmp_path):
    c = config(tmp_path)
    root = tmp_path / "work"
    m = batch.prepare(c, root)
    m = batch.advance(root / "workflow.json", pipeline=True, translator=fake_translate)
    broken = root / "broken.json"
    broken.write_text("not json")
    m["jobs"][0]["review_manifest"] = str(broken)
    reports = publish(m, root)
    assert "report_error" in Path(reports["report"]).read_text()


@pytest.mark.parametrize("sheet_state", ["visible", "hidden", "veryHidden"])
def test_preparation_preserves_excel_saved_views_without_shared_objects(tmp_path, sheet_state):
    from workbook_add_target_sheet import add_target_sheet
    from openpyxl.xml.functions import tostring
    source = book(tmp_path / "excel_saved.xlsx")
    wb = openpyxl.load_workbook(source)
    wb["DE(독일)"].sheet_state = sheet_state
    view = wb["DE(독일)"].sheet_view
    view.zoomScale = 85
    view.zoomScaleNormal = 100
    view.selection[0].activeCell = "B4"
    view.selection[0].sqref = "B4:C5"
    expected = tostring(wb["DE(독일)"].views.to_tree())
    wb.save(source)
    wb.close()
    original_hash = file_sha256(source)
    result = save_prepared(source, tmp_path / "prepared.xlsx", "DE(독일)", "NEW", "de_DE", "BACK")
    saved = openpyxl.load_workbook(result["output"])
    try:
        for name in ("DE(독일)", "NEW", "BACK"):
            assert tostring(saved[name].views.to_tree()) == expected
            assert saved[name].sheet_state == sheet_state
    finally:
        saved.close()
    assert file_sha256(source) == original_hash
    result = add_target_sheet(source, "DE(독일)", "NEW", "de_DE")
    wb = result["wb"]
    try:
        wb["NEW"].sheet_view.selection[0].activeCell = "D9"
        assert wb["DE(독일)"].sheet_view.selection[0].activeCell == "B4"
    finally:
        wb.close()
