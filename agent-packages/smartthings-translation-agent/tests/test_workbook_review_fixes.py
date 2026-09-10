"""Regression cases from the post-step-5 implementation review; synthetic files only."""
import json
import sys
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from workbook_contract import capture_authority, file_sha256, atomic_json
from workbook_incremental_highlight import _save_highlight_verified
from workbook_apply_edits import prepare_edit_contract, contract_writer_identity, _build_contract_edits
from workbook_run import execute_run
import workbook_verifier as verifier
from lint_workbook_writers import scan, violations


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "문안.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "US"
    book.active["C7"] = "Use Brand now"
    book.active["D8"] = "keep"
    book.active.row_dimensions[8].height = 26
    book.save(path)
    book.close()
    return path


def highlight(source, output, mutate=None):
    book = openpyxl.load_workbook(source, rich_text=True)
    book.active["C7"] = CellRichText("Use ", TextBlock(InlineFont(color="0000FF"), "Brand"), " now")
    if mutate:
        mutate(book)
    try:
        return _save_highlight_verified(book, source, output, {"source": capture_authority(source)}, {("US", "C7")})
    finally:
        book.close()


def test_highlight_uses_verified_publication_and_preserves_source(source, tmp_path):
    original = file_sha256(source)
    output = highlight(source, tmp_path / "검수 표시.xlsx")
    assert output.exists() and output.parent.parent.name == "verified"
    assert (output.parent / "batch.complete.json").is_file()
    assert file_sha256(source) == original
    book = openpyxl.load_workbook(output, rich_text=True)
    assert isinstance(book.active["C7"].value, CellRichText)
    assert book.active.row_dimensions[8].height == 26
    book.close()


@pytest.mark.parametrize("mutate", [lambda b: setattr(b.active["D8"], "value", "bad"),
    lambda b: setattr(b.active.row_dimensions[8], "height", 50),
    lambda b: setattr(b.active["D8"], "value", CellRichText(TextBlock(InlineFont(b=True), "keep")))])
def test_highlight_plan_cannot_expand_nonrich_or_outside_scope(source, tmp_path, mutate):
    with pytest.raises(ValueError, match="scope"):
        highlight(source, tmp_path / "render.xlsx", mutate)
    assert not (tmp_path / "verified").exists()


def test_highlight_save_corruption_is_never_promoted(source, tmp_path):
    save = openpyxl.Workbook.save
    def corrupt(book, path):
        book.active.row_dimensions[8].height = 99
        save(book, path)
    with patch.object(openpyxl.Workbook, "save", corrupt):
        with pytest.raises(ValueError, match="no output published"):
            highlight(source, tmp_path / "render.xlsx")
    assert not list((tmp_path / "verified").glob("*/*.xlsx"))
    assert not (tmp_path / "render.xlsx").exists()


def make_run(source, tmp_path):
    plan = prepare_edit_contract(source, [{"sheet":"US", "cell":"C7", "before":"Use Brand now", "after":"Use Brand"}],
        work_root=tmp_path/'work', delivery_root=tmp_path/'delivery', work_id='run')
    approval = tmp_path/'approval.json'
    atomic_json(approval, {"approved":True,"approved_by":"test-user", "contract_sha256":plan['contract_sha256']})
    return Path(plan['contract']), approval


def test_published_batch_survives_verifier_writer_and_input_change(source, tmp_path):
    contract, approval = make_run(source, tmp_path)
    first = execute_run(contract, _build_contract_edits, writer=contract_writer_identity(), approval_path=approval)
    assert first['status'] == 'completed'
    source.unlink()  # history does not require reopening the original
    identity = {**contract_writer_identity(), 'version':'new-code'}
    with patch.object(verifier, 'verifier_version', return_value='new-verifier'):
        second = execute_run(contract, lambda *a: pytest.fail('must not rebuild'), writer=identity, approval_path=approval)
    assert second['status'] == 'completed'
    assert second['result'] == first['result']


def test_unpublished_old_record_is_still_rejected(source, tmp_path):
    contract, approval = make_run(source, tmp_path)
    c = json.loads(contract.read_text())
    book = _build_contract_edits(c, c['artifacts'][0], None)
    candidate = tmp_path/'candidate.xlsx'; book.save(candidate); book.close()
    record = verifier.verify_artifact(contract, 'edited', candidate)
    with patch.object(verifier, 'verifier_version', return_value='new-verifier'):
        with pytest.raises(ValueError, match='Stale'):
            verifier.validate_record(contract, 'edited', candidate, record)


def test_writer_lint_rejects_new_site_despite_guard_import(tmp_path):
    p=tmp_path/'writer.py'; p.write_text('wb.save(path)\n')
    baseline={'sites':scan(tmp_path)}
    p.write_text('from workbook_mutation_guard import save_verified_atomic\nwb.save(path)\nother.save(path)\n')
    assert len(violations(tmp_path,baseline)) == 1
    p.write_text('wb.save(path)\nwriter = wb.save\nwriter(path)\n')
    assert len(violations(tmp_path,baseline)) == 1


def test_repository_has_no_new_workbook_save_bypasses():
    root=Path(__file__).resolve().parents[1]
    baseline=json.loads((root/'docs/workbook_writer_baseline.json').read_text())
    assert violations(root/'scripts',baseline) == []
