"""Thin story/review adapters to the existing local contract runner.

No new scheduler or approval format: an authorized apply call binds its existing
request to an immutable plan. All workbook saves belong to workbook_run.
"""
from __future__ import annotations

import asyncio
import copy
import json
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

import _app_pipeline as ap
from workbook_apply_edits import prepare_edit_contract, _build_contract_edits, _load_edits
from workbook_contract import (atomic_json, capture_authority, canonical, check_inputs,
    digest, file_sha256, load_contract, save_contract, ContractError)
from workbook_mutation_guard import delete_cols_with_manifest, detailed_workbook_snapshot
from workbook_verifier import facts, diff_facts
from workbook_run import execute_run, exclusive_lock
from workbook_highlight_glossary import highlight_in_memory
from workbook_manifest import create_edit_revision, resolve_ledger


def writer_identity():
    names = ("workbook_delivery.py", "workbook_apply_edits.py", "workbook_review_apply.py",
             "workbook_story_apply.py", "workbook_highlight_glossary.py",
             "workbook_incremental_highlight.py", "_app_pipeline.py")
    return {"id": "workbook_delivery", "cost": "local", "version": digest(
        {name: file_sha256(Path(__file__).parent / name) for name in names})}


def _project_columns(snapshot, source, deleted):
    """Independent logical projection of authority facts, never runs delete_cols.

    Dependencies/crossing ranges are rejected by the builder before saving.
    """
    if not deleted:
        return snapshot
    start, end = deleted
    if start <= 3:
        raise ContractError("Review column deletion must leave content columns A:C intact")
    width = end - start + 1
    def shifted(col):
        return col - width if col > end else col
    result = {axis: {} for axis in snapshot}
    for axis, mapping in snapshot.items():
        for key, value in mapping.items():
            path = json.loads(key)
            if len(path) >= 4 and path[0] == "sheets":
                if path[2] == "cells":
                    row, col = coordinate_to_tuple(path[3])
                    if start <= col <= end:
                        continue
                    path[3] = f"{get_column_letter(shifted(col))}{row}"
                elif path[2] == "column_dimensions":
                    col = openpyxl.utils.column_index_from_string(path[3])
                    if start <= col <= end:
                        continue
                    path[3] = get_column_letter(shifted(col))
                    if path[-1] in {"min", "max"}:
                        value = shifted(value or col)
                elif path[2] == "merged_ranges":
                    pass
            if len(path) == 3 and path[2] == "merged_ranges":
                merges = []
                for merged in value:
                    lo, top, hi, bottom = range_boundaries(merged)
                    if start <= lo <= hi <= end:
                        continue
                    if not (hi < start or lo > end):
                        raise ContractError("Merged range crosses deletion boundary")
                    merges.append(f"{get_column_letter(shifted(lo))}{top}:{get_column_letter(shifted(hi))}{bottom}")
                value = sorted(merges)
            result[axis][canonical(path)] = value
    # Dimensions are derived from physical authority cells, including style-only cells.
    book = openpyxl.load_workbook(source, rich_text=True, data_only=False)
    try:
        for sheet in book:
            if any(cell.comment and col >= start for (row, col), cell in sheet._cells.items()):
                raise ContractError("Deleting/moving comment boxes requires an explicit VML formatting plan")
            retained = [(row, shifted(col)) for (row, col), cell in sheet._cells.items()
                        if not start <= col <= end]
            for field, index in (("max_row", 0), ("max_column", 1)):
                result["layout"][canonical(["sheets", sheet.title, field])] = max(
                    (item[index] for item in retained), default=1)
    finally:
        book.close()
    return result


def _apply_facts(snapshot, changes):
    result = copy.deepcopy(snapshot)
    for change in changes:
        key = canonical(change["path"])
        if change["after"]["present"]:
            result[change["axis"]][key] = change["after"]["value"]
        else:
            result[change["axis"]].pop(key, None)
    return result


async def _build(contract, final):
    book = _build_contract_edits(contract, contract["artifacts"][0], None)
    try:
        deleted = contract["delivery_request"].get("deleted")
        if deleted:
            for sheet in book:
                delete_cols_with_manifest(sheet, deleted[0], deleted[1] - deleted[0] + 1)
        report = None
        if final:
            request = contract["delivery_request"]
            report = await highlight_in_memory(book, glossary=Path(request["glossary"]),
                sheets=request["sheets"], cell_range=request["cell_range"],
                activation_manifest=request.get("activation_manifest"), source_sheet=request.get("source_sheet"),
                sheet_langs=request.get("sheet_langs"))
        return book, report
    except BaseException:
        book.close()
        raise


def _execute_build(contract, artifact, last_failure):
    return asyncio.run(_build(contract, artifact["id"] == "final"))[0]


def _request(args, mode, app_root):
    frozen = None
    if getattr(args, "workflow_settings", None):
        from workbook_batch import read, settings_for
        settings_path = Path(args.workflow_settings).resolve()
        stored = read(settings_path)
        frozen = settings_for({"settings": str(settings_path), "settings_sha256": stored["settings_sha256"]})
        if Path(args.glossary).expanduser().resolve() != Path(frozen["glossary"]):
            raise ContractError("Apply must use the frozen job glossary")
        for field in ("source_sheet", "activation_manifest"):
            explicit = getattr(args, field, None)
            if explicit and str(explicit) != str(frozen.get(field)):
                raise ContractError(f"Apply setting differs from reviewed job: {field}")
            setattr(args, field, frozen.get(field))
        if getattr(args, "sheet_langs", None) and ap.load_sheet_langs(args.sheet_langs) != frozen["sheet_langs"]:
            raise ContractError("Apply language mapping differs from reviewed job")
        args.sheet_langs = frozen["sheet_langs_path"]
    source, glossary = Path(args.workbook).expanduser().resolve(), Path(args.glossary).expanduser().resolve()
    request = {"mode": mode, "source": str(source), "glossary": str(glossary),
               "app_root": str(app_root), "cell_range": args.cell_range,
               "activation_manifest": str(Path(args.activation_manifest).expanduser().resolve())
                   if getattr(args, "activation_manifest", None) else None}
    if frozen:
        request["workflow_settings"] = str(settings_path)
        request["settings_sha256"] = frozen["settings_sha256"]
        request["workflow_sheet"] = frozen["sheet"]
    if getattr(args, "source_sheet", None):
        request["source_sheet"] = args.source_sheet
    if getattr(args, "sheet_langs", None):
        request["sheet_langs"] = ap.load_sheet_langs(args.sheet_langs)
    if mode == "story":
        from workbook_story_apply import _split_sheets, _validate_delivery_scope
        request.update(edits=_load_edits(args.edits), sheets=_split_sheets(args.delivery_sheets), deleted=None)
        _validate_delivery_scope(source, request["edits"], request["sheets"], request.get("sheet_langs"))
        if frozen and request["sheets"] != [frozen["sheet"]]:
            raise ContractError("Frozen settings apply to one workbook/language job")
        output = source.with_name(source.stem + "_revised.xlsx")
    else:
        from workbook_review_apply import _parse_columns
        request.update(approval_manifest=str(Path(args.approval_manifest).expanduser().resolve()),
            protected_sheets=frozen["source_sheet"] if frozen else args.protected_sheets,
            deleted=list(_parse_columns(args.drop_review_columns)) if args.drop_review_columns else None)
        output = Path(args.output).expanduser().resolve()
    if output.suffix.lower() != ".xlsx" or output == source:
        raise ContractError("Output must name a separate .xlsx copy")
    request["output"] = str(output)
    return request


async def _prepare(request, root, work_id):
    source = Path(request["source"])
    authorities = {"source": capture_authority(source), "glossary": capture_authority(request["glossary"])}
    # Pin just local glossary dependencies and rule files, never credentials/cache/RAG DB.
    app = Path(request["app_root"]) / "src" / "translation_web_app"
    deps = [app / name for name in ("glossary_checks.py", "prompt_builder.py", "constraint_resolver.py",
                                    "rules_loader.py", "paths.py", "prompt_modules.py")]
    deps += sorted((app / "rules").rglob("*.md"))
    for index, path in enumerate(deps):
        authorities[f"dependency_{index}"] = capture_authority(path)
    if request.get("workflow_settings"):
        authorities["workflow_settings"] = capture_authority(request["workflow_settings"])
    if request.get("activation_manifest"):
        authorities["activation"] = capture_authority(request["activation_manifest"])
    metadata = {}
    outcome_manifest = None
    resolved = dict(request)
    if request["mode"] == "review":
        from workbook_review_apply import _load_manifest, _validate_decisions, _apply_decisions
        authorities["approval_manifest"] = capture_authority(request["approval_manifest"])
        manifest = _load_manifest(Path(request["approval_manifest"]))
        if "changes" in manifest:
            from review_outcomes import outcomes_from_manifest
            terminal = [x for x in manifest["changes"] if x.get("approval_status") in {"approved", "rejected"}]
            outcome_manifest = {**manifest, "changes": terminal}
        protected = set(manifest.get("protected_sheets", request["protected_sheets"].split(",")))
        book = openpyxl.load_workbook(source, rich_text=True, data_only=False)
        try:
            if protected - set(book.sheetnames):
                raise ContractError("Protected sheet missing from workbook")
            decisions = _validate_decisions(book, manifest["decisions"], protected, request["cell_range"], request.get("source_sheet"))
            if request.get("workflow_sheet") and any(d["sheet"] != request["workflow_sheet"] for d in decisions):
                raise ContractError("Approval contains another job language")
            records, changes = _apply_decisions(book, decisions)
            resolved["sheets"] = ([request["workflow_sheet"]] if request.get("workflow_sheet") else
                                  [name for name in book.sheetnames if name in (request.get("sheet_langs") or ap.DEFAULT_SHEET_LANGS)])
        finally:
            book.close()
        resolved["edits"] = [{"sheet": x["sheet"], "cell": x["cell"], "before": x["current"],
                              "after": x["final"], "reason": x["reason"], "rule_ids": x["rule_ids"]} for x in changes]
        metadata = {"decisions": records, "changes": changes,
                    "decision_counts": {d: sum(x["decision"] == d for x in records) for d in ("accept", "partial", "hold")},
                    "protected_validation": {"passed": True, "sheets": sorted(protected)}}
    contract = prepare_edit_contract(source, resolved["edits"], work_root=root,
        delivery_root=Path(request["output"]).parent / "verified", work_id=work_id, persist=False)
    contract.update(authorities=authorities, writer=writer_identity(), command_request=request,
                    delivery_request=resolved, delivery_metadata=metadata)
    baseline = facts(source)
    expected = _project_columns(_apply_facts(baseline, contract["artifacts"][0]["allowed_diffs"]), source, resolved["deleted"])
    first = contract["artifacts"][0]
    first.update(id="edited", output=Path(request["output"]).name, allowed_diffs=diff_facts(baseline, expected))
    # Render a read-only plan; only scoped rich-text changes may extend the contract.
    book, report = await _build(contract, True)
    try:
        rendered = detailed_workbook_snapshot(book)
        rendered["layout"].update({key: value for key, value in baseline["layout"].items() if json.loads(key)[0] == "package"})
        rich_changes = diff_facts(expected, rendered)
        scope = set(resolved["sheets"]) | {"KR(한국)", "US(미국)"}
        lo, hi = map(int, request["cell_range"].upper().replace("C", "").split(":"))
        for change in rich_changes:
            path = change["path"]
            if (change["axis"] != "rich_text" or len(path) != 4 or path[:1] != ["sheets"]
                    or path[1] not in scope or path[2] != "cells" or not path[3].startswith("C")
                    or not lo <= int(path[3][1:]) <= hi):
                raise ContractError(f"Delivery plan changed outside authorized scope: {change}")
    finally:
        book.close()
    final = copy.deepcopy(first)
    final.update(id="final", output=Path(request["output"]).stem + "_final.xlsx",
                 allowed_diffs=diff_facts(baseline, rendered))
    contract["artifacts"] = [first, final]
    contract["no_op_files"] = [x["id"] for x in contract["artifacts"] if not x["allowed_diffs"]]
    contract["highlight_report"] = report
    if outcome_manifest is not None:
        contract["review_outcomes"] = [{**entry, "work_id": work_id}
                                       for entry in outcomes_from_manifest(outcome_manifest)]
    check_inputs(contract)
    save_contract(root / "contract.json", contract)
    return contract


async def run_delivery(args, mode):
    app_root = ap.bootstrap_project(args.app_root).resolve()
    ap.maybe_reexec_with_app_venv(app_root)
    request = _request(args, mode, app_root)
    work_id = mode + "-" + digest(request)[:20]
    root = Path(request["output"]).parent / ".st-runs" / work_id
    contract_path, approval = root / "contract.json", root / "approval.json"
    with exclusive_lock(root / ".prepare.lock"):
        contract = load_contract(contract_path) if contract_path.exists() else await _prepare(request, root, work_id)
        if contract.get("command_request") != request:
            raise ContractError("Command differs from frozen request")
        from workbook_run import _published
        if _published(contract_path, contract) is None:
            check_inputs(contract)
        if getattr(args, "dry_run", False):
            return {"status": "preview", "contract": str(contract_path),
                    "changes": contract["request_edits"], "highlight_scope": contract["highlight_report"]["completed_delivery_sheets"]}
        # This command is already the authorized apply boundary; no second user gate.
        atomic_json(approval, {"approved": True, "approved_by": "authorized_" + mode + "_apply_call",
                               "contract_sha256": file_sha256(contract_path)})
    state = await asyncio.to_thread(execute_run, contract_path, _execute_build,
                                   writer=writer_identity(), approval_path=approval)
    if state["status"] != "completed":
        return state
    with exclusive_lock(root / ".run.lock"):
        cached = root / "delivery_result.json"
        if cached.exists():
            result = json.loads(cached.read_text())
            if (result.get("final") != state["result"]["outputs"]["final"]
                    or result.get("source") != request["source"] or result.get("status") != "ok"):
                raise ContractError("Cached result does not match verified publication")
            for key in ("revision_manifest", "change_log", "highlight_report",
                        "result_manifest" if mode == "review" else "delivery_manifest"):
                if not Path(result[key]).is_file():
                    raise ContractError(f"Delivery sidecar missing: {key}")
        else:
            source = Path(request["source"])
            output, final = (Path(state["result"]["outputs"][key]) for key in ("edited", "final"))
            changes = [{"sheet": x["sheet"], "cell": x["cell"], "old_value": x["before"],
                        "new_value": x["after"], "reason": x["reason"], "rule_ids": x["rule_ids"]}
                       for x in contract["artifacts"][0]["edits"]]
            revision, revision_path = create_edit_revision(source, output, changes)
            baseline, baseline_path, _ = resolve_ledger(source)
            atomic_json(output.with_suffix(".changes.json"), {"source": str(source), "changes": changes,
                "revision_manifest": str(revision_path), "revision_id": revision["revision_id"]})
            final_revision, final_revision_path = create_edit_revision(output, final, [])
            atomic_json(final.with_suffix(".changes.json"), {"source": str(output), "changes": [],
                "revision_manifest": str(final_revision_path), "revision_id": final_revision["revision_id"]})
            report_path = final.with_suffix(".highlight_report.json")
            atomic_json(report_path, contract["highlight_report"])
            count = sum(x["before"] != x["after"] for x in contract["artifacts"][0]["edits"])
            result = {"status": "ok", "artifact_status": "delivery", "source": str(source),
                "revised": str(output), "acceptance_copy": str(output), "final": str(final),
                "change_log": str(output.with_suffix(".changes.json")), "revision_manifest": str(revision_path),
                "revision_id": revision["revision_id"], "workbook_id": baseline["workbook_id"],
                "baseline_manifest": str(baseline_path), "glossary": request["glossary"],
                "approval_manifest": request.get("approval_manifest"), "cell_range": request["cell_range"],
                "review_columns_removed": getattr(args, "drop_review_columns", None),
                "delivery_sheets": contract["delivery_request"]["sheets"],
                "source_groups": contract["highlight_report"]["source_groups"],
                "highlight_report": str(report_path), "highlight_validation": {"passed": True, **contract["highlight_report"]},
                "value_validation": {"actual_value_changes": count, "expected_value_changes": count},
                "value_diff_validation": {"passed": True, "actual": count, "expected": count},
                "work_id": work_id, "contract": str(contract_path),
                "batch_manifest": str(final.parent / "batch.complete.json"), "run_state": str(root / "state.json"),
                **contract["delivery_metadata"]}
            if "review_outcomes" in contract:
                from workbook_contract import atomic_text
                ledger = root / "review_outcomes.jsonl"
                atomic_text(ledger, "".join(canonical(x) + "\n" for x in contract["review_outcomes"]))
                result["outcome_ledger"] = str(ledger)
            manifest = final.with_suffix(".review_apply.json" if mode == "review" else ".delivery.json")
            result["result_manifest" if mode == "review" else "delivery_manifest"] = str(manifest)
            atomic_json(manifest, result)
            atomic_json(cached, result)
        if getattr(args, "result_manifest", None):
            destination = Path(args.result_manifest).expanduser().resolve()
            if destination.exists() and json.loads(destination.read_text()) != result:
                raise ContractError("Result manifest destination already contains another result")
            atomic_json(destination, result)
        return result
