#!/usr/bin/env python3
"""Find glossary occurrences that read as ordinary words, for human confirmation.

The activation manifest can record that a term is active at a story/cell, and the
resolver honours ``active: false`` — but nothing ever produced those entries. The
shipped ES_CO manifest states the reason outright: "Absent blue highlight creates
no inactive decision." Only activation was recordable, so an English source saying
"a safe home" kept matching the product term ``Safe``, and the resolver demanded a
glossary target in text where a plain adjective was correct.

That gap has a measured cost. Of the 21 proposals the ES_CO consensus produced,
six were blocked as ``missing_glossary_target`` even though the reviewer accepted
them verbatim into the delivered workbook. Under a fail-closed resolver those six
corrections never reach a human at all.

This emits *candidates*, never decisions. A candidate becomes an activation entry
only after someone sets ``confirmed: true``; --emit-manifest exports exactly those.
Deactivating a term the source really did use would silently drop a required
translation, which is the more expensive mistake of the two.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _app_pipeline as ap
from workbook_inspect import CONTENT_ROW_END, CONTENT_ROW_START

SCHEMA_VERSION = 1
KIND = "glossary_inactive_candidate_manifest"
BASIS = "common_noun_lowercase_in_source"
POLICY = [
    "A glossary key carrying an uppercase letter is a product/feature term.",
    "Every occurrence of that term in the cell must be lowercase to qualify — activation is keyed by (story, cell, term), so one capitalised use keeps the whole cell active.",
    "A candidate is not a decision. Only entries with confirmed: true become activation entries.",
    "The decision is a property of the source wording, not of any target language: emitted entries carry no locale, and one confirmation serves every target sheet in the same source group.",
    "It does not carry across source groups. Letter case is the signal, so a KR-source sheet (US/JA/CN/TW) needs its own pass against Korean source text.",
    "Excel and the master glossary are read-only inputs.",
]


def _story_of(name: str) -> str:
    match = re.search(r"(?:story[_ -]?)?(\d{3})(?:\D|$)", name, re.IGNORECASE)
    return match.group(1) if match else ""


def find_candidates(checker, workbook: Path, source_sheet: str, target_sheet: str | None = None,
                    target_code: str | None = None) -> list[dict[str, Any]]:
    """One candidate per (story, cell, term) whose source use is entirely lowercase.

    The target sheet is optional and only ever supplies corroborating evidence for
    the reviewer. Nothing that reaches an activation entry depends on it, which is
    what lets one confirmation serve every locale in the source group.
    """
    import openpyxl

    story = _story_of(workbook.name)
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        if source_sheet not in wb.sheetnames:
            return []
        source_ws = wb[source_sheet]
        target_ws = wb[target_sheet] if target_sheet and target_sheet in wb.sheetnames else None
        entries = []
        for row in range(CONTENT_ROW_START, CONTENT_ROW_END + 1):
            source = str(source_ws.cell(row, 3).value or "")
            if not source.strip() or source.strip().lower() == "x":
                continue
            target_text = str(target_ws.cell(row, 3).value or "") if target_ws else ""
            entries.extend(checker.pending_inactive_candidates(
                source, story=story, cell=f"C{row}", target_lang_code=target_code,
                target_text=target_text))
        return entries
    finally:
        wb.close()


def apply_confirmations(candidates: list[dict[str, Any]], *, terms: set[str] | None = None,
                        cells: set[str] | None = None) -> int:
    """Record a reviewer's verdict without hand-editing every entry.

    Reviewers judge these per term far more often than per cell — the question is
    whether an English word is ever the product name, and the answer rarely splits
    within one term.  Cells are addressed as ``story_<nnn>:<cell>`` for the
    exceptions that do split.  Anything not named stays unconfirmed, so a partial
    answer never widens into a blanket one.
    """
    marked = 0
    for item in candidates:
        key = f"story_{item['story']}:{item['cell']}"
        if (terms and item["source_term"] in terms) or (cells and key in cells):
            if not item.get("confirmed"):
                marked += 1
            item["confirmed"] = True
    return marked


def to_activation_entries(candidates: list[dict[str, Any]], *, confirmed_only: bool = True) -> list[dict]:
    """Strip candidates down to what OccurrenceActivationManifest reads."""
    rows = [item for item in candidates if item.get("confirmed")] if confirmed_only else candidates
    return [{"story": item["story"], "cell": item["cell"], "source_term": item["source_term"],
             "active": bool(item.get("active", False)),
             "activation_basis": item.get("activation_basis", BASIS)} for item in rows]


def markdown_review(candidates: list[dict[str, Any]]) -> str:
    lines = ["# 용어집 미적용 후보 검토", "",
             f"- 후보 {len(candidates)}건", "",
             "각 항목은 원문에서 용어집 단어가 **보통명사로 쓰인 것으로 보이는** 경우다.",
             "맞으면 `confirmed: true`로 바꾼다. 틀리면 그대로 둔다 — 확인되지 않은 후보는 무시된다.", ""]
    by_term: dict[str, list[dict]] = {}
    for item in candidates:
        by_term.setdefault(item["source_term"], []).append(item)
    for term in sorted(by_term):
        rows = by_term[term]
        lines.append(f"## `{term}` — {len(rows)}건")
        lines.append("")
        for item in rows:
            evidence = item["evidence"]
            used = "예" if evidence["target_uses_glossary_term"] else "아니오"
            lines.append(f"- story_{item['story']} `{item['cell']}` — 원문 표기 "
                         f"{', '.join(repr(s) for s in evidence['matched_surfaces'])}, "
                         f"용어집 target `{evidence['glossary_target']}`, 번역문이 이 용어를 씀: {used}")
            lines.append(f"  - 원문: {evidence['source_text'][:160]}")
        lines.append("")
    return "\n".join(lines)


async def collect(workbooks: list[Path], glossary: Path, app_root: Path, *,
                  source_sheet: str, target_sheet: str | None) -> list[dict[str, Any]]:
    ap.bootstrap_project(str(app_root))
    from translation_web_app.glossary_checks import GlossaryChecker
    from translation_web_app.prompt_builder import PromptBuilder

    checker = GlossaryChecker(PromptBuilder())
    source_code = ap.DEFAULT_SHEET_LANGS[source_sheet]["code"]
    loaded = await checker.load_glossary_from_file(str(glossary), source_code)
    if not loaded.startswith("✓"):
        raise RuntimeError(f"glossary 로드 실패: {loaded}")
    checker._compile_glossary_re()
    target_code = ap.DEFAULT_SHEET_LANGS[target_sheet]["code"] if target_sheet else None
    candidates: list[dict[str, Any]] = []
    for workbook in workbooks:
        candidates.extend(find_candidates(checker, workbook, source_sheet, target_sheet, target_code))
    return candidates


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbooks", nargs="+", help="검사할 workbook 경로 또는 디렉터리")
    parser.add_argument("--glossary", type=Path, required=True)
    parser.add_argument("--app-root", type=Path)
    parser.add_argument("--source-sheet", default="US(미국)")
    parser.add_argument("--target-sheet",
                        help="검토 근거로 번역문을 함께 보여줄 시트(선택). 판정에는 쓰이지 않으며 "
                             "결과는 같은 source group의 모든 로케일에 그대로 쓴다")
    parser.add_argument("--output", type=Path, help="후보 JSON 출력 경로")
    parser.add_argument("--review-markdown", type=Path, help="사람 검토용 Markdown 출력 경로")
    parser.add_argument("--emit-manifest", type=Path,
                        help="confirmed: true 후보만 activation manifest 형식으로 내보낸다")
    parser.add_argument("--from-candidates", type=Path,
                        help="새로 스캔하지 않고 기존 후보 파일을 읽어 manifest만 만든다")
    parser.add_argument("--confirm-terms",
                        help="사람이 '일반명사'로 판정한 용어를 쉼표로 나열해 confirmed 처리한다")
    parser.add_argument("--confirm-cells",
                        help="용어 단위로 갈리지 않는 예외를 story_012:C7 형식으로 쉼표 나열한다")
    args = parser.parse_args()
    try:
        if args.from_candidates:
            payload = json.loads(args.from_candidates.expanduser().read_text(encoding="utf-8"))
            candidates = payload.get("entries", [])
        else:
            paths: list[Path] = []
            for raw in args.workbooks:
                path = Path(raw).expanduser()
                paths.extend(sorted(path.glob("*.xlsx")) if path.is_dir() else [path])
            candidates = asyncio.run(collect(
                paths, args.glossary.expanduser(), args.app_root,
                source_sheet=args.source_sheet, target_sheet=args.target_sheet))
        confirmed_now = apply_confirmations(
            candidates,
            terms={value.strip() for value in (args.confirm_terms or "").split(",") if value.strip()},
            cells={value.strip() for value in (args.confirm_cells or "").split(",") if value.strip()})
        payload = {"schema_version": SCHEMA_VERSION, "kind": KIND, "policy": POLICY,
                   "scope": {"source_sheet": args.source_sheet,
                             "applies_to": f"{args.source_sheet} source group의 모든 target 시트",
                             "evidence_target_sheet": args.target_sheet or None},
                   "summary": {"candidates": len(candidates),
                               "confirmed": sum(1 for item in candidates if item.get("confirmed")),
                               "confirmed_this_run": confirmed_now,
                               "terms": len({item["source_term"] for item in candidates})},
                   "entries": candidates}
        if args.output:
            ap.write_text_atomic(args.output.expanduser(),
                                 json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
        if args.review_markdown:
            ap.write_text_atomic(args.review_markdown.expanduser(), markdown_review(candidates))
        if args.emit_manifest:
            entries = to_activation_entries(candidates)
            ap.write_text_atomic(args.emit_manifest.expanduser(), json.dumps(
                {"schema_version": SCHEMA_VERSION, "kind": "occurrence_activation_manifest",
                 "policy": POLICY, "scope": payload["scope"], "entries": entries},
                ensure_ascii=False, indent=2) + "\n")
            payload["summary"]["emitted_manifest_entries"] = len(entries)
        print(json.dumps({"status": "ok", **payload["summary"]}, ensure_ascii=False, indent=2))
    except Exception as error:
        print(json.dumps({"status": "error", "error": str(error)}, ensure_ascii=False))
        sys.exit(2)


if __name__ == "__main__":
    main()
