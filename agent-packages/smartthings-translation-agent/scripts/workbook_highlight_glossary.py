#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
workbook_highlight_glossary.py — 용어집 용어 rich text 하이라이트 적용

앱 본체의 TranslationChecker.run_highlight_only_pipeline_generator 를 호출하는
얇은 래퍼다. 하이라이트 로직은 재구현하지 않는다. 공용 부트스트랩/이벤트 헬퍼는
_app_pipeline.py 를 재사용한다.

안전 정책:
  - 원본 파일은 수정하지 않는다. 앱 파이프라인이 *_highlighted_<timestamp>.xlsx
    복사본을 만든다.
  - target 셀의 기존 텍스트 중 용어집 target term과 매칭되는 글자 조각만
    openpyxl rich text로 파란색 처리한다. 번역/수정은 하지 않는다.
  - 용어집 파일을 명시하지 않으면 app repo의 runtime/glossary/latest_glossary.csv를
    사용한다.

크레딧: 0 (LLM 호출 없음, openpyxl rich text 처리만).

사용 예:
  python scripts/workbook_highlight_glossary.py story.xlsx --sheets "BR(브라질)"
  python scripts/workbook_highlight_glossary.py story.xlsx --cell-range C7:C28 --json
  python scripts/workbook_highlight_glossary.py story.xlsx --cell-range C7:C28 --include-source-sheets
  python scripts/workbook_highlight_glossary.py story.xlsx --single-source --source-sheet "US(미국)" --sheets "BR(브라질),DE(독일)"
"""

from __future__ import annotations

import argparse
import asyncio
import contextlib
import io
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _app_pipeline as ap
from workbook_mutation_guard import clone_rich_value, rich_value_signature, save_verified_atomic


def _cell_text(value) -> str:
    return "" if value is None else str(value)


def _verify_text_preservation(source: Path, highlighted: Path) -> dict:
    """Fail closed if a highlight-only pass changes any cell's character stream."""
    before = openpyxl.load_workbook(source, data_only=False)
    after = openpyxl.load_workbook(highlighted, data_only=False, rich_text=True)
    differences: list[str] = []
    for sheet_name in before.sheetnames:
        if sheet_name not in after.sheetnames:
            differences.append(f"missing sheet: {sheet_name}")
            continue
        ws_before, ws_after = before[sheet_name], after[sheet_name]
        max_row = max(ws_before.max_row, ws_after.max_row)
        max_col = max(ws_before.max_column, ws_after.max_column)
        for row in ws_before.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                if _cell_text(cell.value) != _cell_text(ws_after[cell.coordinate].value):
                    differences.append(f"{sheet_name}!{cell.coordinate}")
                    if len(differences) >= 20:
                        raise RuntimeError(
                            "하이라이트가 문안을 변경했습니다: " + ", ".join(differences)
                        )
    if differences:
        raise RuntimeError("하이라이트가 문안을 변경했습니다: " + ", ".join(differences))
    return {"text_preserved": True, "changed_cells": 0}


def _coordinate_in_ranges(coordinate: str, range_arg: str) -> bool:
    column_letter, row = coordinate_from_string(coordinate)
    column = column_index_from_string(column_letter)
    for value in (item.strip() for item in range_arg.split(",")):
        if not value:
            continue
        min_col, min_row, max_col, max_row = range_boundaries(value)
        if (
            (min_col is None or min_col <= column)
            and (max_col is None or column <= max_col)
            and (min_row is None or min_row <= row)
            and (max_row is None or row <= max_row)
        ):
            return True
    return False


def _verify_out_of_scope_rich_text(
    source: Path, candidate: Path, cell_range: str, processed_sheets: set[str]
) -> None:
    before = openpyxl.load_workbook(source, data_only=False, rich_text=True)
    after = openpyxl.load_workbook(candidate, data_only=False, rich_text=True)
    try:
        differences: list[str] = []
        for sheet_name in before.sheetnames:
            ws_before, ws_after = before[sheet_name], after[sheet_name]
            for (_row, _column), cell in ws_before._cells.items():
                if sheet_name in processed_sheets and _coordinate_in_ranges(cell.coordinate, cell_range):
                    continue
                if rich_value_signature(cell.value) != rich_value_signature(ws_after[cell.coordinate].value):
                    differences.append(f"{sheet_name}!{cell.coordinate}")
        if differences:
            raise RuntimeError(
                "처리 범위 밖 rich text가 변경되었습니다: " + ", ".join(differences[:20])
            )
    finally:
        before.close()
        after.close()


def _restore_out_of_scope_rich_text(
    source: Path, highlighted: Path, cell_range: str, processed_sheets: set[str]
) -> dict:
    """Undo workbook-wide rich-run normalization outside the requested highlight scope."""
    before = openpyxl.load_workbook(source, data_only=False, rich_text=True)
    after = openpyxl.load_workbook(highlighted, data_only=False, rich_text=True)
    restored: list[str] = []
    try:
        for sheet_name in before.sheetnames:
            ws_before, ws_after = before[sheet_name], after[sheet_name]
            for (_row, _column), cell in ws_before._cells.items():
                if sheet_name in processed_sheets and _coordinate_in_ranges(cell.coordinate, cell_range):
                    continue
                target = ws_after[cell.coordinate]
                if rich_value_signature(cell.value) != rich_value_signature(target.value):
                    target.value = clone_rich_value(cell.value)
                    restored.append(f"{sheet_name}!{cell.coordinate}")

        save_verified_atomic(
            after,
            highlighted,
            lambda candidate: _verify_out_of_scope_rich_text(
                source, candidate, cell_range, processed_sheets
            ),
            overwrite=True,
        )
    finally:
        before.close()
        after.close()
    return {
        "out_of_scope_rich_text_preserved": True,
        "restored_cell_count": len(restored),
        "restored_cells": restored,
    }


async def run_highlight(args) -> dict:
    app_root = ap.bootstrap_project(args.app_root)
    ap.maybe_reexec_with_app_venv(app_root)

    from dotenv import load_dotenv
    from translation_web_app.checker_service import TranslationChecker

    load_dotenv(app_root / ".env")

    workbook = Path(args.workbook).expanduser()
    if not workbook.is_file():
        raise FileNotFoundError(f"워크북을 찾을 수 없습니다: {workbook}")

    glossary = (
        Path(args.glossary).expanduser()
        if args.glossary
        else app_root / "runtime" / "glossary" / "latest_glossary.csv"
    )
    if not glossary.is_file():
        raise FileNotFoundError(f"용어집 CSV를 찾을 수 없습니다: {glossary}")

    sheet_langs = ap.load_sheet_langs(args.sheet_langs)
    selected_sheets = ap.split_sheets(args.sheets)
    workbook_sheets = ap.workbook_sheetnames(workbook)

    if args.single_source:
        source_sheet = args.source_sheet
        source_lang = sheet_langs.get(source_sheet, {}).get("lang", "English")
        source_groups = None
    else:
        source_sheet = None
        source_lang = "English"
        source_groups = ap.default_source_groups(
            selected_sheets,
            workbook_sheets,
            include_source_sheets=args.include_source_sheets,
        )
        if not source_groups:
            raise ValueError("유효한 source group이 없습니다. --single-source 또는 --sheets 값을 확인하세요.")

    # JSON 모드에서는 앱 내부 print 가 stdout(=JSON 채널)을 오염시키지 않도록
    # 초기화와 파이프라인 실행 전체를 캡처한다. yield 이벤트는 그대로 수집한다.
    init_stdout = io.StringIO()
    redirect = contextlib.redirect_stdout(init_stdout) if args.json else contextlib.nullcontext()

    events: list[dict] = []
    with redirect:
        checker = TranslationChecker(max_concurrency=max(1, args.max_concurrency), no_backtranslation=True)
        if getattr(args, "activation_manifest", None):
            checker.load_activation_manifest(str(args.activation_manifest))
        async for event in checker.run_highlight_only_pipeline_generator(
            source_file_path=str(workbook),
            cell_range=args.cell_range,
            sheet_lang_map=sheet_langs,
            glossary_file_path=str(glossary),
            selected_sheets=selected_sheets,
            source_sheet_name=source_sheet,
            source_lang=source_lang,
            source_groups=source_groups,
            include_source_sheets=args.include_source_sheets,
        ):
            events.append(event)
            if not args.json:
                etype = event.get("type")
                if etype == "log":
                    print(event.get("message") or event.get("log"))
                elif etype == "progress" and args.verbose:
                    print(f"{event.get('percent', 0)}% {event.get('log', '')}")
                elif etype == "error":
                    print(f"ERROR: {event.get('message')}")

    # JSON 모드: 캡처된 앱 print 를 로그 이벤트로 흡수
    if args.json:
        for line in init_stdout.getvalue().splitlines():
            if line.strip():
                events.insert(0, {"type": "log", "message": line.strip()})

    summary = ap.event_summary(events)
    if summary["status"] == "ok" and summary.get("excel_path"):
        if source_groups:
            processed_sheets = {
                sheet
                for group in source_groups
                for sheet in group.get("target_sheets", [])
                if sheet in workbook_sheets
            }
        elif selected_sheets:
            processed_sheets = {
                sheet
                for sheet in selected_sheets
                if sheet in workbook_sheets and (args.include_source_sheets or sheet != source_sheet)
            }
        else:
            processed_sheets = {
                sheet
                for sheet in workbook_sheets
                if sheet in sheet_langs and (args.include_source_sheets or sheet != source_sheet)
            }
        summary["rich_text_preservation"] = _restore_out_of_scope_rich_text(
            workbook,
            Path(summary["excel_path"]),
            args.cell_range,
            processed_sheets,
        )
        summary["text_validation"] = _verify_text_preservation(
            workbook, Path(summary["excel_path"])
        )
    summary.update({
        "source": str(workbook),
        "glossary": str(glossary),
        "cell_range": args.cell_range,
        "selected_sheets": selected_sheets,
        "source_groups": source_groups,
        "single_source": args.single_source,
        "include_source_sheets": args.include_source_sheets,
    })
    return summary


def _report_path_for(summary: dict) -> str | None:
    """하이라이트 리포트(output_data: 불일치/괄호/대소문자 로그)를 산출본 옆에 저장한다."""
    report = summary.get("output_data")
    excel_path = summary.get("excel_path")
    if not report or not excel_path:
        return None
    out = Path(excel_path).with_suffix("")
    report_path = Path(f"{out}.highlight_report.txt")
    try:
        ap.write_text_atomic(report_path, report)
        return str(report_path)
    except OSError:
        return None


def main() -> None:
    parser = argparse.ArgumentParser(
        description="SmartThings 워크북의 glossary target term을 Excel rich text로 하이라이트 (크레딧 0)"
    )
    parser.add_argument("workbook", help="원본 .xlsx 경로 (수정되지 않음)")
    parser.add_argument("--cell-range", default="C7:C28", help="처리할 소스 셀 범위")
    parser.add_argument("--sheets", help="대상 시트 CSV. 예: 'BR(브라질),DE(독일)'")
    parser.add_argument("--activation-manifest", help="확정된 활성화 설정")
    parser.add_argument("--glossary", help="용어집 CSV 경로. 기본값: runtime/glossary/latest_glossary.csv")
    parser.add_argument("--sheet-langs", help="sheet_langs JSON 파일 경로. 기본값: 앱 표준 매핑")
    parser.add_argument("--single-source", action="store_true", help="복수 source group 대신 --source-sheet 하나만 사용")
    parser.add_argument("--source-sheet", default="US(미국)", help="--single-source 사용 시 소스 시트")
    parser.add_argument(
        "--include-source-sheets",
        action="store_true",
        help="기본 그룹 모드에서 KR/US source sheet 자체도 하이라이트 대상에 포함",
    )
    parser.add_argument("--max-concurrency", type=int, default=10)
    parser.add_argument("--app-root", help="app repo 경로 명시")
    parser.add_argument("--json", action="store_true", help="JSON 출력")
    parser.add_argument("--verbose", action="store_true", help="progress 로그도 출력")
    args = parser.parse_args()

    try:
        res = asyncio.run(run_highlight(args))
    except Exception as e:
        if args.json:
            print(json.dumps({"status": "error", "error": str(e)}, ensure_ascii=False, indent=2))
        else:
            print(f"❌ {e}")
        sys.exit(1)

    report_path = _report_path_for(res)
    if report_path:
        res["report_path"] = report_path

    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        if res["status"] == "ok":
            print("✅ 하이라이트 완료")
            print(f"   원본: {res['source']}")
            print(f"   수정본: {res.get('excel_path')}")
            if report_path:
                print(f"   리포트: {report_path} (용어집 불일치·괄호·대소문자 점검)")
        else:
            print("❌ 하이라이트 중 오류가 발생했습니다.")
            for err in res.get("errors", []):
                print(f"   - {err.get('message')}")
            sys.exit(1)





async def highlight_in_memory(wb, *, glossary: Path, sheets: list[str], cell_range: str,
                              activation_manifest: str | None = None, source_sheet: str | None = None,
                              sheet_langs: dict | None = None) -> dict:
    """Local delivery pass; no model/RAG clients and no workbook writes.

    Bootstrap the app before calling. Reuse its glossary and occurrence resolver
    and the existing rich-text renderer; inspect only existing cells.
    """
    from translation_web_app.glossary_checks import GlossaryChecker
    from translation_web_app.prompt_builder import PromptBuilder
    from workbook_incremental_highlight import _relevant_target_terms, _term_spans, render_cell
    from workbook_review_apply import _parse_c_range
    start, end = _parse_c_range(cell_range)
    mapping = sheet_langs or ap.DEFAULT_SHEET_LANGS
    required = {source_sheet} if source_sheet else {"KR(한국)", "US(미국)"}
    scope = set(sheets) | required
    if scope - set(wb.sheetnames) or scope - set(mapping):
        raise ValueError("납품 시트와 KR/US source sheet의 표준 매핑이 필요합니다.")
    groups = ([{"source_sheet": source_sheet, "target_sheets": sorted(scope)}] if source_sheet
              else ap.default_source_groups(sorted(scope), wb.sheetnames, True))
    completed, cells = {}, {}
    for group in groups:
        checker = GlossaryChecker(PromptBuilder())
        checker.load_activation_manifest(activation_manifest)
        source = group["source_sheet"]
        await checker.load_glossary_from_file(str(glossary), mapping[source]["code"])
        if not checker.glossary:
            raise ValueError(f"비어 있거나 읽을 수 없는 glossary: {source}")
        source_ws = wb[source]
        story_cell = source_ws._cells.get((5, 3))
        import re
        digits = re.findall(r"\d+", str(story_cell.value or "")) if story_cell else []
        story = digits[-1][-3:].zfill(3) if digits else None
        for target in group["target_sheets"]:
            if target in completed:
                continue
            completed[target] = 0
            for row in range(start, end + 1):
                source_cell = source_ws._cells.get((row, 3))
                cell = wb[target]._cells.get((row, 3))
                if cell is None or cell.value is None:
                    continue
                if cell.data_type == "f" or (source_cell is not None and source_cell.data_type == "f"):
                    raise ValueError(f"Formula in glossary scope requires a separate plan: {target}!C{row}")
                text = str(cell.value)
                if not text or text.lower() == "x":
                    continue
                source_text = str(source_cell.value or "") if source_cell else ""
                terms, excluded = _relevant_target_terms(checker, source_text,
                    mapping[target]["code"], story=story, cell=cell.coordinate)
                spans, missing = _term_spans(text, terms)
                cell.value = render_cell(text, base_font=cell.font, red_spans=[], blue_spans=spans)
                if str(cell.value) != text:
                    raise ValueError("Glossary renderer changed text")
                cells[f"{target}!{cell.coordinate}"] = {"blue_spans": spans, "missing": missing, "excluded": excluded}
                completed[target] += 1
    if set(completed) != scope:
        raise ValueError("Incomplete source-group highlight scope")
    return {"completed_delivery_sheets": completed, "source_groups": groups, "cells": cells}


if __name__ == "__main__":
    main()
