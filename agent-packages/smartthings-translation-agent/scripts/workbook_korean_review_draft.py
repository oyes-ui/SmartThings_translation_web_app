#!/usr/bin/env python3
"""Create Korean draft-review copies in column F with red/blue rich text.

The source wording and glossary highlighting authority is KR column C.  Only
manifest-declared KR column-F cells may change.  Spelling edits are rendered
red first and inherited/declared glossary spans are rendered blue last.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

from workbook_incremental_highlight import BLUE, RED, _merge_spans, _term_spans, render_cell
from workbook_manifest import diff_spans, file_sha256
from workbook_mutation_guard import (
    clone_rich_value,
    rich_value_signature,
    save_verified_atomic,
    semantic_workbook_snapshot,
    text,
)


KOREAN_SHEET = "KR(한국)"


def _rgb(value: Any) -> str | None:
    color = getattr(value, "color", None)
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    rgb = getattr(color, "rgb", None)
    return str(rgb).upper() if rgb else None


def _is_blue(font: Any) -> bool:
    rgb = _rgb(font)
    if not rgb:
        return False
    rgb = rgb[-6:]
    try:
        red, green, blue = int(rgb[:2], 16), int(rgb[2:4], 16), int(rgb[4:], 16)
    except ValueError:
        return False
    return blue >= 200 and blue > red * 2 and blue > green * 2


def inherited_blue_terms(value: Any) -> list[str]:
    if not isinstance(value, CellRichText):
        return []
    return [
        part.text
        for part in value
        if isinstance(part, TextBlock) and part.text and _is_blue(part.font)
    ]


def _target_rows(spec: dict[str, Any]) -> list[int]:
    rows = sorted({int(row) for row in spec.get("target_rows", [])})
    if not rows or any(row < 7 for row in rows):
        raise ValueError("target_rows는 7행 이상의 명시 좌표여야 합니다")
    return rows


def _edits(spec: dict[str, Any]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for item in spec.get("edits", []):
        cell = str(item["cell"]).upper()
        if not cell.startswith("F"):
            raise ValueError(f"F열 외 편집은 허용되지 않습니다: {cell}")
        if cell in result:
            raise ValueError(f"중복 편집 좌표: {cell}")
        result[cell] = item
    return result


def _glossary_terms(spec: dict[str, Any]) -> dict[str, list[str]]:
    rows = {f"F{row}" for row in _target_rows(spec)}
    result: dict[str, list[str]] = {}
    for cell, terms in spec.get("glossary_terms", {}).items():
        coordinate = str(cell).upper()
        if coordinate not in rows:
            raise ValueError(f"glossary_terms에 target_rows 밖 좌표가 있습니다: {coordinate}")
        if not isinstance(terms, list):
            raise ValueError(f"glossary_terms 값은 배열이어야 합니다: {coordinate}")
        result[coordinate] = sorted({text(term) for term in terms if text(term)}, key=lambda value: (-len(value), value))
    return result


def expected_values(source: Path, spec: dict[str, Any]) -> dict[str, str]:
    workbook = openpyxl.load_workbook(source, data_only=False, rich_text=True)
    try:
        sheet = workbook[KOREAN_SHEET]
        edits = _edits(spec)
        expected: dict[str, str] = {}
        for row in _target_rows(spec):
            cell = f"F{row}"
            before = text(sheet[f"C{row}"].value)
            item = edits.get(cell)
            if item:
                if text(item.get("before")) != before:
                    raise ValueError(f"원문 드리프트: {KOREAN_SHEET}!C{row}")
                expected[cell] = text(item.get("after"))
            else:
                expected[cell] = before
        if set(edits) - set(expected):
            raise ValueError("edits에 target_rows 밖 좌표가 있습니다: " + ", ".join(sorted(set(edits) - set(expected))))
        return expected
    finally:
        workbook.close()


def build_workbook(source: Path, spec: dict[str, Any]):
    workbook = openpyxl.load_workbook(source, data_only=False, rich_text=True)
    sheet = workbook[KOREAN_SHEET]
    edits = _edits(spec)
    glossary_terms = _glossary_terms(spec)
    for row in _target_rows(spec):
        source_cell = sheet[f"C{row}"]
        target_cell = sheet[f"F{row}"]
        before = text(source_cell.value)
        item = edits.get(target_cell.coordinate)
        after = text(item["after"]) if item else before
        if item and text(item.get("before")) != before:
            workbook.close()
            raise ValueError(f"원문 드리프트: {KOREAN_SHEET}!C{row}")

        blue_terms = inherited_blue_terms(source_cell.value)
        blue_terms.extend(glossary_terms.get(target_cell.coordinate, []))
        if item:
            blue_terms.extend(text(value) for value in item.get("blue_terms", []) if text(value))
        blue_spans, _unmatched = _term_spans(after, blue_terms)
        red_spans = diff_spans(before, after)["red_spans"] if item else []
        if red_spans or blue_spans:
            target_cell.value = render_cell(
                after,
                base_font=target_cell.font,
                red_spans=red_spans,
                blue_spans=blue_spans,
                blue_color="FF0000FF",
            )
        else:
            target_cell.value = clone_rich_value(source_cell.value)
        if text(target_cell.value) != after:
            workbook.close()
            raise RuntimeError(f"렌더링이 문안을 변경했습니다: {KOREAN_SHEET}!{target_cell.coordinate}")
    return workbook


def _blue_spans(value: Any) -> list[list[int]]:
    spans: list[list[int]] = []
    if not isinstance(value, CellRichText):
        return spans
    offset = 0
    for part in value:
        segment = part.text if isinstance(part, TextBlock) else str(part)
        if isinstance(part, TextBlock) and _is_blue(part.font):
            spans.append([offset, offset + len(segment)])
        offset += len(segment)
    return _merge_spans((start, end) for start, end in spans)


def _cell_map(workbook) -> dict[tuple[str, str], dict[str, Any]]:
    return {
        (sheet.title, cell.coordinate): {
            "text": text(cell.value),
            "data_type": cell.data_type,
            "rich": rich_value_signature(cell.value),
        }
        for sheet in workbook.worksheets
        for cell in sheet._cells.values()
    }


def verify_output(output: Path, source: Path, spec: dict[str, Any]) -> dict[str, Any]:
    expected = expected_values(source, spec)
    declared_glossary = _glossary_terms(spec)
    original = openpyxl.load_workbook(source, data_only=False, rich_text=True)
    revised = openpyxl.load_workbook(output, data_only=False, rich_text=True)
    try:
        before_snapshot = semantic_workbook_snapshot(original)
        after_snapshot = semantic_workbook_snapshot(revised)
        for axis in ("layout_sha256", "styles_sha256", "annotations_sha256"):
            if before_snapshot[axis] != after_snapshot[axis]:
                raise AssertionError(f"허용되지 않은 {axis.removesuffix('_sha256')} 변경")

        allowed = {(KOREAN_SHEET, cell) for cell in expected}
        before_cells, after_cells = _cell_map(original), _cell_map(revised)
        for key in sorted(set(before_cells) | set(after_cells)):
            if key in allowed:
                continue
            if before_cells.get(key) != after_cells.get(key):
                raise AssertionError(f"범위 밖 셀 변경: {key[0]}!{key[1]}")

        revised_sheet = revised[KOREAN_SHEET]
        for cell, value in expected.items():
            if text(revised_sheet[cell].value) != value:
                raise AssertionError(f"F열 결과 불일치: {KOREAN_SHEET}!{cell}")
        for cell, terms in declared_glossary.items():
            value = revised_sheet[cell].value
            desired, missing = _term_spans(text(value), terms)
            if missing:
                raise AssertionError(f"선언 용어가 F열에 없습니다: {KOREAN_SHEET}!{cell} {missing}")
            actual = _blue_spans(value)
            uncovered = [
                span for span in desired
                if not any(start <= span[0] and span[1] <= end for start, end in actual)
            ]
            if uncovered:
                raise AssertionError(f"용어집 파란 하이라이트 누락: {KOREAN_SHEET}!{cell} {uncovered}")
        return {
            "status": "verified",
            "target_cells": sorted(expected),
            "changed_value_cells": sorted(
                cell for cell, value in expected.items()
                if text(original[KOREAN_SHEET][cell].value) != value
            ),
            "glossary_verified_cells": sorted(declared_glossary),
            "source_sha256": file_sha256(source),
            "output_sha256": file_sha256(output),
            "preserved_axes": ["layout", "styles", "annotations", "all_non_target_cells"],
        }
    finally:
        revised.close()
        original.close()


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, path)


def run_manifest(manifest_path: Path, output_dir: Path) -> dict[str, Any]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("approval_status") != "approved":
        raise PermissionError("approval_status가 approved인 manifest만 Excel에 반영할 수 있습니다")
    specs = manifest.get("workbooks", [])
    if not specs:
        raise ValueError("workbooks가 비어 있습니다")
    output_dir.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    promoted: list[Path] = []
    with tempfile.TemporaryDirectory(dir=output_dir, prefix=".kr-review-staging-") as tmp:
        staging = Path(tmp)
        staged: list[tuple[Path, Path]] = []
        for spec in specs:
            source = Path(spec["source"])
            final = output_dir / spec["output_name"]
            if final.exists():
                raise FileExistsError(f"출력 파일이 이미 존재합니다: {final}")
            stage = staging / final.name
            workbook = build_workbook(source, spec)
            try:
                verification = save_verified_atomic(
                    workbook,
                    stage,
                    lambda path, src=source, current=spec: verify_output(path, src, current),
                )
            finally:
                workbook.close()
            staged.append((stage, final))
            results.append({"source": source.name, "output": str(final), **verification})
        try:
            for stage, final in staged:
                os.replace(stage, final)
                promoted.append(final)
        except Exception:
            for final in promoted:
                if final.exists():
                    final.unlink()
            raise

    result = {
        "status": "verified",
        "renderer": "red_then_inherited_glossary_blue_v1",
        "value_authority": "각 입력 workbook의 KR(한국) C열 + 승인 manifest의 맞춤법/용어집 수정",
        "structure_authority": "각 입력 workbook 전체",
        "format_authority": "각 입력 workbook 전체; 허용 변경은 KR(한국) 대상 F열 rich text만",
        "results": results,
    }
    result_manifest = output_dir / "korean_review_draft.result.json"
    _atomic_json(result_manifest, result)
    return {**result, "result_manifest": str(result_manifest)}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    result = run_manifest(args.manifest, args.output_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
