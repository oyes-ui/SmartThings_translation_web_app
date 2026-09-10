#!/usr/bin/env python3
"""Compose existing workbook preparation, draft API, agent review and approval reports.

No agent API, scheduler or automatic paid retry. Config is authored by the active
agent from the user's request; advance exposes the existing staged prompts.
"""
from __future__ import annotations
import argparse
import asyncio
import json
import os
import subprocess
import sys
from pathlib import Path

import _app_pipeline as ap
import agent_staged_batch as review
from workbook_contract import atomic_json, digest, file_sha256
from workbook_run import exclusive_lock


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def copy_file(source: Path, target: Path):
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".tmp")
    tmp.write_bytes(source.read_bytes())
    os.replace(tmp, target)


def prepare(config: dict, root: Path) -> dict:
    root = root.expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    with exclusive_lock(root / ".lock"):
        path = root / "workflow.json"
        if path.exists():
            raise FileExistsError("Existing workflow: use advance/status")
        if not config.get("jobs"):
            raise ValueError("At least one job is required")
        # Validate every configuration before capturing any work.
        specs = []
        for raw in config["jobs"]:
            job = dict(raw)
            if job.get("glossary_approved") is not True:
                raise ValueError("Confirm glossary application/exclusion settings before prepare")
            for name in ("workbook", "glossary"):
                job[name] = str(Path(job[name]).expanduser().resolve(strict=True))
            if job.get("activation_manifest"):
                job["activation_manifest"] = str(Path(job["activation_manifest"]).expanduser().resolve(strict=True))
            marker = Path(job["glossary"]).with_suffix(".sync.json")
            if marker.exists() and read(marker).get("status") != "ok":
                raise ValueError("Default glossary sync incomplete; run sync-default")
            sheet = job["sheet"]
            job["source_sheet"] = job.get("source_sheet") or (
                ap.GROUP_A_SOURCE if sheet in ap.GROUP_A_TARGETS else ap.GROUP_B_SOURCE)
            job["sheet_langs"] = {**ap.DEFAULT_SHEET_LANGS, **job.get("sheet_langs", {})}
            for name in (sheet, job["source_sheet"]):
                info = job["sheet_langs"].get(name, {})
                if not info.get("code") or not info.get("lang"):
                    raise ValueError(f"Missing language mapping: {name}")
            if sheet == job["source_sheet"]:
                raise ValueError("Translation target must differ from source")
            if (job.get("backtranslation_sheet") or job.get("backtranslation_lang") or job.get("with_backtranslation")) and not job.get("api_audit"):
                raise ValueError("Backtranslation requires explicit api_audit selection")
            if job.get("output_dir"):
                job["output_dir"] = str(Path(job["output_dir"]).expanduser().resolve())
            job.setdefault("translate", True)
            job.setdefault("review", True)
            job.setdefault("story", Path(job["workbook"]).stem)
            specs.append(job)
        keys = [digest([j["workbook"], j["sheet"]])[:16] for j in specs]
        if len(keys) != len(set(keys)):
            raise ValueError("Duplicate workbook/language job")
        jobs = []
        for key, job in zip(keys, specs):
            folder = root / "jobs" / key
            originals = {}
            for field in ("workbook", "glossary", "activation_manifest"):
                if not job.get(field):
                    continue
                source = Path(job[field])
                target = folder / "inputs" / field / source.name
                copy_file(source, target)
                originals[field] = str(source)
                job[field] = str(target)
            atomic_json(folder / "sheet_langs.json", job["sheet_langs"])
            job["sheet_langs_path"] = str(folder / "sheet_langs.json")
            job["originals"] = originals
            job["input_hashes"] = {field: file_sha256(job[field]) for field in
                                   ("workbook", "glossary", "activation_manifest", "sheet_langs_path") if job.get(field)}
            job["settings_sha256"] = digest(job)
            atomic_json(folder / "settings.json", job)
            jobs.append({"job_id": key, "settings": str(folder / "settings.json"),
                         "settings_sha256": job["settings_sha256"], "state": "pending"})
        manifest = {"schema_version": 1, "kind": "workbook_workflow", "app_root": str(Path(config["app_root"]).resolve()),
                    "obsidian_dir": str(Path(config["obsidian_dir"]).expanduser().resolve()) if config.get("obsidian_dir") else None,
                    "jobs": jobs}
        atomic_json(path, manifest)
        return manifest


def settings_for(job):
    settings = read(job["settings"])
    claimed = settings.pop("settings_sha256")
    if claimed != job["settings_sha256"] or digest(settings) != claimed:
        raise ValueError("Frozen job settings changed")
    settings["settings_sha256"] = claimed
    for field, sha in settings["input_hashes"].items():
        if file_sha256(settings[field]) != sha:
            raise ValueError(f"Frozen input changed: {field}")
    return settings


def translate(settings, workbook, app_root):
    cmd = [sys.executable, str(Path(__file__).with_name("workbook_translate.py")), str(workbook),
           "--pipeline", "--json", "--single-source", "--source-sheet", settings["source_sheet"],
           "--sheets", settings["sheet"], "--glossary", settings["glossary"],
           "--sheet-langs", settings["sheet_langs_path"], "--app-root", app_root,
           "--translation-model", settings.get("translation_model", "gemini-3.8-flash")]
    if settings.get("audit_model"):
        cmd.extend(["--audit-model", settings["audit_model"]])
    if settings.get("max_concurrency"):
        cmd.extend(["--max-concurrency", str(settings["max_concurrency"])])
    cmd.append("--with-api-audit" if settings.get("api_audit") else "--translate-only")
    if settings.get("with_backtranslation"):
        cmd.append("--with-backtranslation")
    for field in ("activation_manifest", "backtranslation_sheet", "backtranslation_lang"):
        if settings.get(field):
            cmd.extend(["--" + field.replace("_", "-"), settings[field]])
    completed = subprocess.run(cmd, capture_output=True, text=True)
    try:
        result = json.loads(completed.stdout)
    except ValueError as error:
        raise ValueError("Translation returned no readable JSON; inspect API run before retry") from error
    if completed.returncode or result.get("status") != "ok":
        raise ValueError(f"Translation did not complete: {result.get('status', 'error')}")
    return result


def validate_draft(path, settings):
    """Cheap completeness check, not a final delivery contract or language audit."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, rich_text=True)
    try:
        source, target = wb[settings["source_sheet"]], wb[settings["sheet"]]
        missing = [f"C{row}" for row in range(7, 29)
                   if str(source.cell(row, 3).value or "").strip().lower() not in {"", "x"}
                   and str(target.cell(row, 3).value or "").strip().lower() in {"", "x"}]
        if missing:
            raise ValueError(f"Incomplete draft: {missing}")
    finally:
        wb.close()


def advance(path: Path, *, pipeline=False, retry_jobs=(), translator=None) -> dict:
    path = path.resolve()
    with exclusive_lock(path.parent / ".lock"):
        manifest = read(path)
        if set(retry_jobs) - {j["job_id"] for j in manifest["jobs"]}:
            raise ValueError("Unknown retry job")
        if retry_jobs and not pipeline:
            raise ValueError("Paid retry requires explicit --pipeline approval")
        for job in manifest["jobs"]:
            folder = Path(job["settings"]).parent
            receipt = folder / "translation_result.json"
            try:
                settings = settings_for(job)
                folder = Path(job["settings"]).parent
                receipt = folder / "translation_result.json"
                if job["job_id"] in retry_jobs:
                    if job["state"] not in {"translation_error", "api_recovery_required"}:
                        raise ValueError("Only failed/uncertain API calls can be retried")
                    job["state"] = "pending"
                if job["state"] == "translation_started":
                    # A persisted receipt permits local recovery; never reissue the call.
                    job["state"] = "draft_ready" if receipt.exists() else "api_recovery_required"
                if job["state"] in {"translation_error", "api_recovery_required"}:
                    continue
                if job["state"] in {"pending", "awaiting_api_approval", "preparation_error"}:
                    workbook = Path(settings["workbook"])
                    if settings.get("prepare"):
                        from workbook_add_target_sheet import save_prepared
                        prep = settings["prepare"]
                        # Already-prepared targets need no copy or template.
                        required = [settings["sheet"]] + ([settings["backtranslation_sheet"]] if settings.get("backtranslation_sheet") else [])
                        if set(required) - set(ap.workbook_sheetnames(workbook)):
                            output = folder / "prepared" / workbook.name
                            result = save_prepared(workbook, output, prep["template_sheet"], settings["sheet"],
                                                   prep["lang_code"], settings.get("backtranslation_sheet"))
                            workbook = Path(result["output"])
                    job["prepared_workbook"] = str(workbook)
                    if settings.get("prep_only"):
                        job["state"] = "prepared"
                        continue
                    if settings["translate"]:
                        names = ap.workbook_sheetnames(workbook)
                        if settings["sheet"] not in names or settings["source_sheet"] not in names:
                            raise ValueError("Missing target/source sheet; configure workbook preparation before API")
                        if not pipeline:
                            job["state"] = "awaiting_api_approval"
                            continue
                        job["state"] = "translation_started"
                        job.pop("error", None)
                        atomic_json(path, manifest)  # Durable before spending any credits.
                        response = (translator or translate)(settings, workbook, manifest["app_root"])
                        if response.get("status") != "ok":
                            raise ValueError("Translation incomplete")
                        produced = Path(response["excel_path"]).resolve(strict=True)
                        # Receipt is persisted before further local operations.
                        atomic_json(receipt, {"settings_sha256": job["settings_sha256"],
                                              "path": str(produced), "sha256": file_sha256(produced)})
                    else:
                        atomic_json(receipt, {"settings_sha256": job["settings_sha256"],
                                              "path": str(workbook), "sha256": file_sha256(workbook)})
                    job["state"] = "draft_ready"
                if job["state"] in {"draft_ready", "review_error", "reviewing", "completed", "incomplete"}:
                    result = read(receipt)
                    if result["settings_sha256"] != job["settings_sha256"] or file_sha256(result["path"]) != result["sha256"]:
                        raise ValueError("Draft or receipt changed; cannot resume")
                    draft_root = Path(settings["output_dir"]) / job["job_id"] if settings.get("output_dir") else folder / "draft"
                    draft = draft_root / Path(settings["workbook"]).name
                    if draft.resolve().is_relative_to((folder / "inputs").resolve()) or str(draft.resolve()) in settings["originals"].values():
                        raise ValueError("Draft output overlaps input authority")
                    if draft.exists() and file_sha256(draft) != result["sha256"]:
                        raise ValueError("Captured draft changed")
                    if not draft.exists():
                        copy_file(Path(result["path"]), draft)
                    validate_draft(draft, settings)
                    job["draft"] = str(draft)
                    if not settings["review"]:
                        job["state"] = "draft_completed"
                        continue
                    review_path = folder / "review" / "batch_manifest.json"
                    if not review_path.exists():
                        asyncio.run(review.prepare_batch(draft, [settings["sheet"]], review_path.parent,
                            Path(settings["glossary"]), Path(manifest["app_root"]), max_concurrency=1,
                            activation_manifest=Path(settings["activation_manifest"]) if settings.get("activation_manifest") else None,
                            source_sheet=settings["source_sheet"], sheet_langs=settings["sheet_langs"]))
                    staged = review.advance_batch(review_path)
                    job["review_manifest"] = str(review_path)
                    counts = staged["status_counts"]
                    job["state"] = ("completed" if counts == {"completed": 1} else
                                    "incomplete" if counts.get("incomplete") else "review_error" if counts.get("error") else "reviewing")
                    job["ready_for_agent"] = staged["ready_for_agent"]
                    job["error"] = [j["error"] for j in staged["jobs"] if j.get("error")] or None
            except Exception as error:
                job["error"] = str(error)
                job["state"] = ("review_error" if receipt.exists() else
                                "translation_error" if job["state"] in {"translation_started", "translation_error", "api_recovery_required"} else
                                "review_error" if job.get("draft") else "preparation_error")
            finally:
                atomic_json(path, manifest)
        from workflow_approval import publish
        manifest["reports"] = publish(manifest, path.parent)
        atomic_json(path, manifest)
        return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("prepare")
    p.add_argument("--config", type=Path, required=True)
    p.add_argument("--work-dir", type=Path, required=True)
    for command in ("advance", "status"):
        p = sub.add_parser(command)
        p.add_argument("--manifest", type=Path, required=True)
        if command == "advance":
            p.add_argument("--pipeline", action="store_true")
            p.add_argument("--retry-job", action="append", default=[])
    args = parser.parse_args()
    if args.command == "prepare":
        result = prepare(read(args.config), args.work_dir)
    elif args.command == "advance":
        result = advance(args.manifest, pipeline=args.pipeline, retry_jobs=args.retry_job)
    else:
        result = read(args.manifest)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
