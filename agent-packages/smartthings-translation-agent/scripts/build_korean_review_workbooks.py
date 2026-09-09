#!/usr/bin/env python3
"""Build Korean review workbooks from non-standard source sheets and a review template."""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from workbook_mutation_guard import (
    clone_rich_value,
    copy_row_layout,
    rich_value_signature,
    save_verified_atomic,
    slice_rich_value,
    text,
)


KOREAN_SHEET = "KR(한국)"
SECTION_ROWS = (10, 15, 20, 25)


def _trim_rich(value: Any) -> Any:
    displayed = text(value)
    left = len(displayed) - len(displayed.lstrip())
    right = len(displayed.rstrip())
    return slice_rich_value(value, left, right)


def _disclaimer_value(value: Any) -> Any:
    displayed = text(value)
    marker = "※ Disclaimer"
    start = displayed.find(marker)
    if start < 0:
        return "x"
    extracted = _trim_rich(slice_rich_value(value, start + len(marker)))
    return extracted if text(extracted) else "x"


def _button_value(value: Any) -> Any:
    displayed = text(value)
    match = re.search(r"(?m)^-[ \t]*버튼명[ \t]*:?[ \t]*([^\r\n]*)", displayed)
    if not match:
        return "x"
    extracted = _trim_rich(slice_rich_value(value, *match.span(1)))
    return extracted if text(extracted) else "x"


def extract_source_payload(source_path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(source_path, data_only=False, rich_text=True)
    try:
        sheet = workbook.worksheets[0]
        section_starts = [
            row for row in range(1, sheet.max_row + 1)
            if text(sheet.cell(row, 2).value).strip().lower().startswith("section")
        ]
        sections: list[dict[str, Any]] = []
        for index, row in enumerate(section_starts):
            end = section_starts[index + 1] if index + 1 < len(section_starts) else sheet.max_row + 1
            disclaimer = "x"
            button = "x"
            for candidate_row in range(row + 2, end):
                candidate = sheet.cell(candidate_row, 3).value
                displayed = text(candidate).lstrip()
                if displayed.startswith("※ Disclaimer"):
                    disclaimer = _disclaimer_value(candidate)
                elif displayed.startswith("※ Deep Link"):
                    button = _button_value(candidate)
            sections.append({
                "title": clone_rich_value(sheet.cell(row, 3).value),
                "description": clone_rich_value(sheet.cell(row + 1, 3).value),
                "disclaimer": clone_rich_value(disclaimer),
                "button": clone_rich_value(button),
                "heights": {
                    "title": max(sheet.row_dimensions[row].height or 0, 22.0),
                    "description": max(sheet.row_dimensions[row + 1].height or 0, 45.0),
                    "disclaimer": max(
                        next(
                            (
                                sheet.row_dimensions[candidate_row].height or 0
                                for candidate_row in range(row + 2, end)
                                if text(sheet.cell(candidate_row, 3).value).lstrip().startswith("※ Disclaimer")
                            ),
                            0,
                        ),
                        22.0,
                    ),
                    "button": 22.0,
                },
            })
        if not 1 <= len(sections) <= len(SECTION_ROWS):
            raise ValueError(f"섹션 수를 처리할 수 없습니다: {source_path.name} ({len(sections)})")
        return {
            "title": clone_rich_value(sheet["C5"].value),
            "description": clone_rich_value(sheet["C6"].value),
            "title_height": max(sheet.row_dimensions[5].height or 0, 22.0),
            "description_height": max(sheet.row_dimensions[6].height or 0, 45.0),
            "sections": sections,
        }
    finally:
        workbook.close()


def _extend_fourth_section(workbook) -> None:
    for sheet in workbook.worksheets:
        copy_row_layout(sheet, sheet, 19, columns=range(1, 9), target_row=24)
        for source_row, target_row in zip(range(20, 24), range(25, 29)):
            copy_row_layout(sheet, sheet, source_row, columns=range(1, 9), target_row=target_row)
        suffixes = ("", "_description", "_disclaimer", "_button")
        for target_row, suffix in zip(range(25, 29), suffixes):
            sheet.cell(target_row, 2).value = f'="//section_"&RIGHT($C$5, 3)&"_4{suffix}"'
            sheet.cell(target_row, 5).value = f'="//section_"&RIGHT($C$5, 3)&"_4{suffix}"'
        sheet["F27"] = clone_rich_value(sheet["F22"].value)
        sheet["F28"] = clone_rich_value(sheet["F23"].value)


def build_workbook(template_path: Path, source_path: Path, story_number: str, update: datetime):
    payload = extract_source_payload(source_path)
    workbook = openpyxl.load_workbook(template_path, data_only=False, rich_text=True)
    if len(payload["sections"]) == 4:
        _extend_fourth_section(workbook)

    story_id = f"story_{story_number}"
    for sheet in workbook.worksheets:
        sheet["C2"] = update
        sheet["F2"] = update
        sheet["C5"] = story_id
        sheet["F5"] = story_id
        sheet.sheet_state = "visible" if sheet.title == KOREAN_SHEET else "hidden"
    workbook.active = workbook.sheetnames.index(KOREAN_SHEET)

    korean = workbook[KOREAN_SHEET]
    for row in range(7, 29):
        korean.cell(row, 3).value = None
    korean["C7"] = clone_rich_value(payload["title"])
    korean["C8"] = clone_rich_value(payload["description"])
    korean.row_dimensions[7].height = payload["title_height"]
    korean.row_dimensions[8].height = payload["description_height"]
    for target_row, section in zip(SECTION_ROWS, payload["sections"]):
        korean.cell(target_row, 3).value = clone_rich_value(section["title"])
        korean.cell(target_row + 1, 3).value = clone_rich_value(section["description"])
        korean.cell(target_row + 2, 3).value = clone_rich_value(section["disclaimer"])
        korean.cell(target_row + 3, 3).value = clone_rich_value(section["button"])
        korean.row_dimensions[target_row].height = section["heights"]["title"]
        korean.row_dimensions[target_row + 1].height = section["heights"]["description"]
        korean.row_dimensions[target_row + 2].height = section["heights"]["disclaimer"]
        korean.row_dimensions[target_row + 3].height = section["heights"]["button"]
    return workbook


def verify_output(path: Path, template_path: Path, source_path: Path, story_number: str) -> dict[str, Any]:
    expected = extract_source_payload(source_path)
    template = openpyxl.load_workbook(template_path, data_only=False, rich_text=True)
    output = openpyxl.load_workbook(path, data_only=False, rich_text=True)
    try:
        if output.sheetnames != template.sheetnames:
            raise AssertionError("시트 순서가 양식과 다릅니다")
        for sheet in output.worksheets:
            expected_state = "visible" if sheet.title == KOREAN_SHEET else "hidden"
            if sheet.sheet_state != expected_state:
                raise AssertionError(f"시트 숨김 상태 오류: {sheet.title}")
            if sheet["C5"].value != f"story_{story_number}" or sheet["F5"].value != f"story_{story_number}":
                raise AssertionError(f"story_id 오류: {sheet.title}")

        korean = output[KOREAN_SHEET]
        expected_cells: dict[str, Any] = {"C7": expected["title"], "C8": expected["description"]}
        for target_row, section in zip(SECTION_ROWS, expected["sections"]):
            expected_cells.update({
                f"C{target_row}": section["title"],
                f"C{target_row + 1}": section["description"],
                f"C{target_row + 2}": section["disclaimer"],
                f"C{target_row + 3}": section["button"],
            })
        for coordinate, expected_value in expected_cells.items():
            actual = korean[coordinate].value
            if text(actual) != text(expected_value):
                raise AssertionError(f"텍스트 불일치: {coordinate}")
            if rich_value_signature(actual) != rich_value_signature(expected_value):
                raise AssertionError(f"리치 텍스트 불일치: {coordinate}")

        expected_heights = {7: expected["title_height"], 8: expected["description_height"]}
        for target_row, section in zip(SECTION_ROWS, expected["sections"]):
            expected_heights.update({
                target_row: section["heights"]["title"],
                target_row + 1: section["heights"]["description"],
                target_row + 2: section["heights"]["disclaimer"],
                target_row + 3: section["heights"]["button"],
            })
        for row, expected_height in expected_heights.items():
            if korean.row_dimensions[row].height != expected_height:
                raise AssertionError(f"행 높이 불일치: {row}")

        section_count = len(expected["sections"])
        for index, target_row in enumerate(SECTION_ROWS[:section_count], start=1):
            expected_formula = f'="//section_"&RIGHT($C$5, 3)&"_{index}"'
            if korean.cell(target_row, 2).value != expected_formula:
                raise AssertionError(f"섹션 키 오류: B{target_row}")
        if section_count == 4:
            for sheet in output.worksheets:
                if sheet.max_row < 28:
                    raise AssertionError(f"4번 섹션 구조 누락: {sheet.title}")

        for sheet_name in template.sheetnames:
            template_sheet = template[sheet_name]
            output_sheet = output[sheet_name]
            for row in range(1, 24):
                for column in range(1, 9):
                    coordinate = output_sheet.cell(row, column).coordinate
                    allowed_value = coordinate in {"C2", "F2", "C5", "F5"} or (
                        sheet_name == KOREAN_SHEET and column == 3 and 7 <= row <= 23
                    )
                    if not allowed_value and output_sheet.cell(row, column).value != template_sheet.cell(row, column).value:
                        raise AssertionError(f"허용되지 않은 값 변경: {sheet_name}!{coordinate}")
                    if output_sheet.cell(row, column)._style != template_sheet.cell(row, column)._style:
                        raise AssertionError(f"양식 서식 변경: {sheet_name}!{coordinate}")
        return {
            "status": "verified",
            "story": story_number,
            "section_count": len(expected["sections"]),
            "content_cells": sorted(expected_cells),
            "visible_sheets": [sheet.title for sheet in output.worksheets if sheet.sheet_state == "visible"],
        }
    finally:
        output.close()
        template.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--date", default="2026-09-04")
    parser.add_argument("--source", action="append", nargs=2, metavar=("NUMBER", "PATH"), required=True)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    update = datetime.strptime(args.date, "%Y-%m-%d")
    specs = [(number, Path(path)) for number, path in args.source]
    if [number for number, _path in specs] != ["003", "013", "039"]:
        raise ValueError("이 배치는 003, 013, 039 순서로만 생성합니다")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    promoted: list[Path] = []

    with tempfile.TemporaryDirectory(dir=args.output_dir, prefix=".korean-review-staging-") as temporary:
        staging = Path(temporary)
        staged_outputs: list[tuple[Path, Path]] = []
        for number, source_path in specs:
            filename = f"(CX Center) SmartThings_2.0_Story_Contents_{number}_KR_260904.xlsx"
            staged = staging / filename
            final = args.output_dir / filename
            if final.exists() and not args.overwrite:
                raise FileExistsError(f"출력 파일이 이미 존재합니다: {final}")
            workbook = build_workbook(args.template, source_path, number, update)
            try:
                verification = save_verified_atomic(
                    workbook,
                    staged,
                    lambda path, n=number, s=source_path: verify_output(path, args.template, s, n),
                )
            finally:
                workbook.close()
            results.append({"output": str(final), "source": str(source_path), **verification})
            staged_outputs.append((staged, final))

        backups: dict[Path, Path] = {}
        try:
            for _staged, final in staged_outputs:
                if final.exists():
                    backup = staging / f"backup-{final.name}"
                    shutil.copy2(final, backup)
                    backups[final] = backup
            for staged, final in staged_outputs:
                os.replace(staged, final)
                promoted.append(final)
        except Exception:
            for final in promoted:
                backup = backups.get(final)
                if backup and backup.exists():
                    os.replace(backup, final)
                elif final.exists():
                    final.unlink()
            raise

    manifest = {
        "status": "verified",
        "value_authority": "각 원본 스토리의 Overview·Section 텍스트",
        "structure_authority": str(args.template),
        "format_authority": str(args.template),
        "rich_text_authority": "각 원본 C열의 rich-text run",
        "allowed_changes": [
            "KR(한국) C열 스토리·섹션 내용",
            "C2/F2 update date, C5/F5 story_id",
            "비한국어 시트 hidden",
            "013·039의 4번 섹션 행 25:28 확장",
            "KR(한국) 내용 행 높이는 각 원본 대응 행을 기준으로 보정",
        ],
        "results": results,
    }
    if args.manifest:
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
