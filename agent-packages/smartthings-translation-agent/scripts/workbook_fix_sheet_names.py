#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
workbook_fix_sheet_names.py — 워크북의 시트명만 교정한 사본을 생성한다 (크레딧 0).

원본은 절대 수정하지 않는다. 여러 --rename "OLD=NEW" 쌍을 받아 시트 제목만
바꾸고(내용/서식/위치는 그대로) 새 경로에 저장한다.

사용 예 (CO 배치 대상 파일의 기존 결함 보정):
  python workbook_fix_sheet_names.py 046.xlsx --rename "046=_archived_046" --out 046_fixed.xlsx
  python workbook_fix_sheet_names.py 048.xlsx \
      --rename "FR (프랑스)=FR(프랑스)" --rename "PT(포루투갈)=PT(포르투갈)" \
      --out 048_fixed.xlsx
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import openpyxl


def parse_rename(raw: str) -> tuple[str, str]:
    if "=" not in raw:
        raise ValueError(f"--rename 값은 'OLD=NEW' 형식이어야 합니다: {raw!r}")
    old, new = raw.split("=", 1)
    old, new = old.strip(), new.strip()
    if not old or not new:
        raise ValueError(f"--rename 값에 빈 이름이 있습니다: {raw!r}")
    return old, new


def fix_sheet_names(workbook: Path, renames: list[tuple[str, str]]) -> dict:
    wb = openpyxl.load_workbook(workbook, rich_text=True)
    applied = []
    skipped = []
    for old, new in renames:
        if old not in wb.sheetnames:
            skipped.append({"old": old, "new": new, "reason": "시트 없음"})
            continue
        if new in wb.sheetnames:
            skipped.append({"old": old, "new": new, "reason": "대상 이름이 이미 존재함"})
            continue
        wb[old].title = new
        applied.append({"old": old, "new": new})
    return {"wb": wb, "applied": applied, "skipped": skipped}


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("workbook", help="원본 .xlsx 경로 (수정되지 않음)")
    p.add_argument("--rename", action="append", required=True,
                    help="'OLD=NEW' 시트명 교정, 여러 번 지정 가능")
    p.add_argument("--out", required=True, help="보정된 사본을 저장할 경로")
    p.add_argument("--json", action="store_true")
    args = p.parse_args()

    workbook = Path(args.workbook).expanduser()
    if not workbook.is_file():
        msg = {"status": "error", "error": f"워크북을 찾을 수 없습니다: {workbook}"}
        print(json.dumps(msg, ensure_ascii=False) if args.json else f"❌ {msg['error']}")
        sys.exit(2)

    try:
        renames = [parse_rename(r) for r in args.rename]
    except ValueError as e:
        msg = {"status": "error", "error": str(e)}
        print(json.dumps(msg, ensure_ascii=False) if args.json else f"❌ {e}")
        sys.exit(2)

    result = fix_sheet_names(workbook, renames)

    out_path = Path(args.out).expanduser()
    tmp = out_path.with_suffix(out_path.suffix + ".tmp")
    result["wb"].save(tmp)
    os.replace(tmp, out_path)

    summary = {
        "status": "ok",
        "source": str(workbook),
        "output": str(out_path),
        "applied": result["applied"],
        "skipped": result["skipped"],
    }
    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        print(f"✅ 시트명 보정 완료: {out_path}")
        for r in result["applied"]:
            print(f"   - {r['old']!r} → {r['new']!r}")
        for r in result["skipped"]:
            print(f"   (건너뜀) {r['old']!r} → {r['new']!r}: {r['reason']}")


if __name__ == "__main__":
    main()
