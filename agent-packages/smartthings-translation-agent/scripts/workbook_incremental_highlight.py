#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Incremental review renderer: red edits first, blue glossary terms last.

The glossary resolver remains the app's ``TranslationChecker``.  This module
only composes its term matches with revision-manifest diff spans and writes one
rich-text pass, so a second save cannot discard the red review markup.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles import Font
from openpyxl.cell.text import InlineFont

import _app_pipeline as ap
from workbook_manifest import file_sha256, load_revision, text_sha256, update_render_state
from workbook_mutation_guard import detailed_workbook_snapshot


RED = "FF0000"
BLUE = "0000FF"


def _sha(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _font(base: Font | None, color: str | None) -> InlineFont:
    params: dict[str, Any] = {}
    if base:
        params = {
            "rFont": base.name, "sz": base.sz, "b": base.b, "i": base.i, "u": base.u,
            "strike": base.strike, "family": base.family, "charset": base.charset,
            "outline": base.outline, "shadow": base.shadow, "condense": base.condense,
            "extend": base.extend, "vertAlign": base.vertAlign, "scheme": base.scheme,
        }
    if color:
        params["color"] = color
    return InlineFont(**params)


def _merge_spans(spans: list[tuple[int, int]]) -> list[list[int]]:
    merged: list[list[int]] = []
    for start, end in sorted({(start, end) for start, end in spans if start < end}):
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return merged


def _term_spans(text: str, terms: list[str]) -> tuple[list[list[int]], list[str]]:
    found: list[tuple[int, int]] = []
    unmatched: list[str] = []
    for term in sorted({term.strip() for term in terms if term and term.strip()}, key=len, reverse=True):
        matches = list(re.finditer(re.escape(term), text, flags=re.IGNORECASE))
        if matches:
            found.extend(match.span() for match in matches)
        else:
            unmatched.append(term)
    return _merge_spans(found), unmatched


def _fold_bare_whitespace_runs(rt: CellRichText) -> CellRichText:
    """Fold whitespace-only rich-text runs into a neighbour.

    openpyxl's ``whitespace()`` helper only marks a run ``xml:space="preserve"``
    when it contains BOTH whitespace and non-whitespace text.  A run that is
    *only* whitespace ships without the attribute and Excel trims it on open.
    This is the same fix as ``TranslationChecker._fold_bare_whitespace_runs``
    but operates on a finished ``CellRichText`` without requiring the app import.
    """
    blocks: list[list] = []
    for part in rt:
        if isinstance(part, TextBlock):
            blocks.append([part.text or "", part.font])
        else:
            blocks.append([str(part), None])
    if not any(text and text.strip() == "" for text, _ in blocks):
        return rt  # no bare whitespace runs
    folded: list[list] = []
    for seg_text, seg_font in blocks:
        if seg_text == "":
            continue
        if folded and seg_text.strip() == "":
            folded[-1][0] += seg_text
        else:
            folded.append([seg_text, seg_font])
    if len(folded) >= 2 and folded[0][0].strip() == "":
        folded[1][0] = folded[0][0] + folded[1][0]
        folded.pop(0)
    out = CellRichText()
    for seg_text, seg_font in folded:
        out.append(seg_text if seg_font is None else TextBlock(text=seg_text, font=seg_font))
    return out


def render_cell(
    text: str, *, base_font: Font | None, red_spans: list[list[int]], blue_spans: list[list[int]],
    blue_color: str = BLUE,
) -> CellRichText | str:
    """Compose a cell in one pass. Blue is intentionally evaluated after red."""
    if not text or not (red_spans or blue_spans):
        return text
    boundaries = {0, len(text)}
    for start, end in [*red_spans, *blue_spans]:
        boundaries.add(max(0, min(len(text), start)))
        boundaries.add(max(0, min(len(text), end)))
    ordered = sorted(boundaries)
    result = CellRichText()
    for start, end in zip(ordered, ordered[1:]):
        segment = text[start:end]
        if not segment:
            continue
        # Glossary wins intentionally: red is calculated first, then blue overlays it.
        color = blue_color if any(left <= start and end <= right for left, right in blue_spans) else (
            RED if any(left <= start and end <= right for left, right in red_spans) else None
        )
        result.append(TextBlock(text=segment, font=_font(base_font, color)))
    return _fold_bare_whitespace_runs(result)


def _relevant_target_terms(checker, source_text: str, target_code: str, *, story=None, cell=None) -> tuple[list[str], list[dict]]:
    terms: list[str] = []
    excluded: list[dict] = []
    for source_term in checker._get_relevant_glossary_terms(source_text) or []:  # app canonical resolver
        metadata = checker.glossary.get(source_term)
        if not metadata:
            excluded.append({"term": source_term, "reason": "missing_glossary_metadata"})
            continue
        active = (checker.is_term_active_for_occurrence(source_term, metadata.get("rule", ""), story, cell)
                  if story is not None or cell is not None else
                  not checker.prompt_builder.is_glossary_deactivated(metadata.get("rule", "")))
        if not active:
            excluded.append({"term": source_term, "reason": "deactivated_glossary_rule"})
            continue
        target_value = checker._get_target_val(metadata["targets"], target_code)
        if not target_value:
            excluded.append({"term": source_term, "reason": "no_target_locale_value"})
            continue
        for candidate in (re.sub(r"\(.*?\)", "", target_value).strip(), target_value.strip()):
            if candidate:
                terms.append(candidate)
                terms.extend(value.strip() for value in re.split(r"[,/]", candidate) if value.strip())
    return sorted(set(terms), key=lambda value: (-len(value), value.casefold())), excluded


def _range_cells(ws, range_arg: str):
    from openpyxl.utils.cell import range_boundaries
    seen = set()
    for area in (x.strip() for x in range_arg.split(",") if x.strip()):
        left, top, right, bottom = range_boundaries(area)
        for (row, col), cell in sorted(ws._cells.items()):
            if left <= col <= right and top <= row <= bottom and cell.coordinate not in seen:
                seen.add(cell.coordinate)
                yield cell


def _red_by_cell(revision: dict) -> dict[tuple[str, str], list[list[int]]]:
    return {
        (change["sheet"], change["cell"]): change.get("diff", {}).get("red_spans", [])
        for change in revision["changes"]
    }


def _is_rich_text(value: Any) -> bool:
    return isinstance(value, CellRichText)


def _workbook_text_snapshot(wb) -> dict[tuple[str, str], str]:
    """Capture every non-empty cell's displayed character stream."""
    return {
        (ws.title, cell.coordinate): _cell_text(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }


def _harden_workbook_whitespace_runs(wb) -> list[str]:
    """Fold unsafe whitespace-only runs throughout a workbook before saving.

    ``openpyxl`` serializes every rich-text cell it loads, including cells this
    renderer did not touch. Therefore folding only the newly rendered cells is
    insufficient: an untouched cell can still lose a standalone space in
    Excel after the workbook is saved.
    """
    hardened: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not _is_rich_text(cell.value):
                    continue
                # A cell whose complete display value is whitespace has no
                # neighbouring text to absorb the run. Excel may trim it, but
                # there is no document text to protect and rewriting it cannot
                # make the serialization safer.
                if not _cell_text(cell.value).strip():
                    continue
                folded = _fold_bare_whitespace_runs(cell.value)
                if folded is not cell.value:
                    if _cell_text(folded) != _cell_text(cell.value):
                        raise RuntimeError(f"공백 hardening이 문안을 변경했습니다: {ws.title}!{cell.coordinate}")
                    cell.value = folded
                    hardened.append(f"{ws.title}!{cell.coordinate}")
    return hardened


def _bare_whitespace_run_cells(wb) -> list[str]:
    """Return cells that Excel may trim because they contain a bare space run."""
    unsafe = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if (
                    _is_rich_text(cell.value)
                    and _cell_text(cell.value).strip()
                    and any(str(part).strip() == "" and str(part) for part in cell.value)
                ):
                    unsafe.append(f"{ws.title}!{cell.coordinate}")
    return unsafe


def _verify_saved_workbook(output: Path, expected_texts: dict[tuple[str, str], str]) -> None:
    verify = openpyxl.load_workbook(output, rich_text=True, data_only=False)
    try:
        actual_texts = _workbook_text_snapshot(verify)
        changed = [
            f"{sheet}!{cell}"
            for sheet, cell in sorted(set(expected_texts) | set(actual_texts))
            if expected_texts.get((sheet, cell)) != actual_texts.get((sheet, cell))
        ]
        if changed:
            raise RuntimeError("저장 후 텍스트 무결성 실패: " + ", ".join(changed[:20]))
        unsafe = _bare_whitespace_run_cells(verify)
        if unsafe:
            raise RuntimeError("저장 후 공백-only rich-text run이 남았습니다: " + ", ".join(unsafe[:20]))
    finally:
        verify.close()


def _next_output_path(workbook: Path, timestamp: str) -> Path:
    candidate = workbook.with_name(f"{workbook.stem}_review_highlighted_{timestamp}{workbook.suffix}")
    suffix = 2
    from workbook_contract import digest
    while candidate.exists() or (workbook.parent / ".st-runs" / ("highlight-" + digest(candidate.name)[:24])).exists():
        candidate = workbook.with_name(
            f"{workbook.stem}_review_highlighted_{timestamp}_{suffix}{workbook.suffix}"
        )
        suffix += 1
    return candidate


def _save_highlight_verified(wb, workbook, output, authorities, allowed_cells):
    """One artifact through the shared runner; only scoped rich-text diffs allowed.

    Values, structure and formatting authorities are the supplied workbook.
    The approved render plan declares exact runs, disk verification is independent.
    """
    from workbook_contract import atomic_json, canonical, digest, save_contract
    from workbook_run import execute_run
    from workbook_verifier import facts, diff_facts
    source = Path(workbook).resolve()
    baseline = facts(source)
    planned = detailed_workbook_snapshot(wb)
    planned["layout"].update({key: value for key, value in baseline["layout"].items()
                              if json.loads(key)[0] == "package"})
    differences = diff_facts(baseline, planned)
    for change in differences:
        path = change["path"]
        if (change["axis"] != "rich_text" or len(path) != 4 or path[0] != "sheets"
                or path[2] != "cells" or (path[1], path[3]) not in allowed_cells):
            raise ValueError(f"Highlight plan exceeds authorized rich-text scope: {path}")
    work_id = "highlight-" + digest(output.name)[:24]
    root = output.parent / ".st-runs" / work_id
    ref = {"authority": "source", "file": "."}
    writer = {"id": "incremental_highlight", "version": file_sha256(__file__), "cost": "local"}
    contract = {"schema_version": 1, "work_id": work_id, "work_root": str(root.resolve()),
        "staging_root": str((root / "staging").resolve()),
        "delivery_root": str((output.parent / "verified").resolve()),
        "authorities": authorities, "writer": writer,
        "artifacts": [{"id": "highlight", "output": output.name,
            "authorities": {role: dict(ref) for role in ("values", "structure", "formatting")},
            "allowed_diffs": differences}], "no_op_files": [] if differences else ["highlight"],
        "recovery": {"max_attempts": 1, "same_failure_limit": 1}}
    path = root / "contract.json"
    sha = save_contract(path, contract)
    approval = root / "approval.json"
    atomic_json(approval, {"approved": True, "approved_by": "authorized_review_highlight_call",
                          "contract_sha256": sha})
    def build(*_):
        return wb
    result = execute_run(path, build, writer=writer, approval_path=approval)
    if result["status"] != "completed":
        raise ValueError(f"Highlight verification failed; no output published: {root / 'state.json'}")
    return Path(result["result"]["outputs"]["highlight"])


async def run_incremental_highlight(args) -> dict:
    app_root = ap.bootstrap_project(args.app_root)
    ap.maybe_reexec_with_app_venv(app_root)
    from dotenv import load_dotenv
    from translation_web_app.checker_service import TranslationChecker

    workbook = Path(args.workbook).expanduser()
    glossary = Path(args.glossary).expanduser()
    revision_path = Path(args.revision_manifest).expanduser()
    if not workbook.is_file() or not glossary.is_file() or not revision_path.is_file():
        raise FileNotFoundError("workbook, glossary, revision manifest 경로를 모두 확인하세요.")
    from workbook_contract import capture_authority, check_inputs
    authorities = {"source": capture_authority(workbook.resolve()),
                   "glossary": capture_authority(glossary.resolve()),
                   "revision": capture_authority(revision_path.resolve())}
    revision = load_revision(revision_path)
    wb = openpyxl.load_workbook(workbook, rich_text=True, data_only=False)
    try:
        sheet_langs = ap.load_sheet_langs(args.sheet_langs)
        groups = ap.default_source_groups(ap.split_sheets(args.sheets), wb.sheetnames, args.include_source_sheets)
        if not groups:
            raise ValueError("유효한 source group이 없습니다. --sheets 또는 --include-source-sheets를 확인하세요.")
        load_dotenv(app_root / ".env")
        checker = TranslationChecker(max_concurrency=1, no_backtranslation=True)
        if getattr(args, "activation_manifest", None):
            checker.load_activation_manifest(str(args.activation_manifest))
            authorities["activation"] = capture_authority(Path(args.activation_manifest).resolve())
        red_by_cell = _red_by_cell(revision)
        old_state = revision.get("render_state", {}).get("cells", {})
        new_state: dict[str, dict] = {}
        applied = skipped = 0
        unmatched: dict[str, list[str]] = {}

        for group in groups:
            source_sheet = group["source_sheet"]
            source_info = sheet_langs[source_sheet]
            await checker.load_glossary_from_file(str(glossary), source_info["code"])
            source_ws = wb[source_sheet]
            import re
            digits = re.findall(r"\d+", str(source_ws["C5"].value or ""))
            story = digits[-1][-3:].zfill(3) if digits else None
            source_cells = list(_range_cells(source_ws, args.cell_range))
            for target_sheet in group["target_sheets"]:
                target_info = sheet_langs[target_sheet]
                target_ws = wb[target_sheet]
                for source_cell in source_cells:
                    target_cell = target_ws._cells.get((source_cell.row, source_cell.column))
                    if target_cell is None:
                        continue
                    if target_cell.data_type == "f":
                        raise ValueError("Formula cells cannot be rendered as rich text")
                    text = _cell_text(target_cell.value)
                    if not text or text.lower() == "x":
                        continue
                    key = f"{target_sheet}!{target_cell.coordinate}"
                    terms, excluded = _relevant_target_terms(checker, _cell_text(source_cell.value), target_info["code"],
                                                             story=story, cell=source_cell.coordinate)
                    blue, missing = _term_spans(text, terms)
                    red = red_by_cell.get((target_sheet, target_cell.coordinate), [])
                    highlighted_terms = [term for term in terms if term not in missing]
                    desired = {
                        "text_sha256": text_sha256(text), "terms_sha256": _sha(terms),
                        "red_spans": red, "blue_spans": blue, "highlighted_terms": highlighted_terms,
                        "excluded_terms": excluded, "unmatched_terms": missing,
                    }
                    # Skip only when this workbook retained the prior rich-text rendering.
                    if old_state.get(key) == desired and _is_rich_text(target_cell.value):
                        new_state[key] = desired
                        skipped += 1
                        continue
                    target_cell.value = render_cell(text, base_font=target_cell.font, red_spans=red, blue_spans=blue)
                    if _cell_text(target_cell.value) != text:
                        raise RuntimeError(f"rich-text renderer가 문안을 변경했습니다: {key}")
                    new_state[key] = desired
                    if missing or excluded:
                        unmatched[key] = missing
                    applied += 1

        expected_texts = _workbook_text_snapshot(wb)
        hardened_cell_coordinates = _harden_workbook_whitespace_runs(wb)
        if _workbook_text_snapshot(wb) != expected_texts:
            raise RuntimeError("공백 hardening 후 워크북 문안이 변경되었습니다.")
        unsafe_before_save = _bare_whitespace_run_cells(wb)
        if unsafe_before_save:
            raise RuntimeError("저장 전 공백-only rich-text run이 남았습니다: " + ", ".join(unsafe_before_save[:20]))

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output = _next_output_path(workbook, timestamp)
        allowed_cells = {(sheet, cell) for key in new_state for sheet, cell in [key.rsplit("!", 1)]}
        allowed_cells.update(tuple(key.rsplit("!", 1)) for key in hardened_cell_coordinates)
        check_inputs({"authorities": authorities})
        output = _save_highlight_verified(wb, workbook, output, authorities, allowed_cells)
    finally:
        wb.close()
    state = {
        "renderer": "incremental_red_then_glossary_blue_v1",
        "glossary": {"file": glossary.name, "sha256": file_sha256(glossary)},
        "cells": new_state,
        "last_output": output.name,
        "last_output_sha256": file_sha256(output),
        "applied_cells": applied,
        "skipped_cells": skipped,
        "hardened_whitespace_cells": len(hardened_cell_coordinates),
        "hardened_whitespace_cell_coordinates": hardened_cell_coordinates,
        "unmatched_terms": unmatched,
    }
    update_render_state(revision_path, state)
    return {"status": "ok", "output": str(output), "revision_manifest": str(revision_path), **state}





def main() -> None:
    parser = argparse.ArgumentParser(description="revision manifest 기반 증분 red→blue Excel rich-text 하이라이트")
    parser.add_argument("workbook")
    parser.add_argument("--revision-manifest", required=True)
    parser.add_argument("--activation-manifest", help="확정된 활성화 설정")
    parser.add_argument("--glossary", required=True)
    parser.add_argument("--sheets", help="대상 시트 CSV")
    parser.add_argument("--cell-range", default="C7:C28")
    parser.add_argument("--include-source-sheets", action="store_true")
    parser.add_argument("--sheet-langs")
    parser.add_argument("--app-root")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    try:
        result = asyncio.run(run_incremental_highlight(args))
    except Exception as exc:
        result = {"status": "error", "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["status"] != "ok":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
