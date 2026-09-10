#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Apply an approved native-review decision manifest to a workbook copy.

This is deliberately an executor, not a reviewer.  It reads the current text
from C, reviewer proposals from F and comments from H, but changes C only for
``accept`` / ``partial`` decisions already recorded in a manifest.  The source
workbook is never modified.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

import _app_pipeline as ap


DECISIONS = {"accept", "partial", "hold"}
DEFAULT_PROTECTED = ["KR(한국)", "US(미국)", "CN(중국)", "BR(브라질)", "RU(러시아)"]


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _atomic_json(path: Path, payload: dict) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(tmp, path)


APPROVED_STATUS = "approved"


def decisions_from_report_changes(changes: list) -> list[dict]:
    """Convert a report_format_spec.md approval manifest into apply decisions.

    The two vocabularies describe different things and are deliberately not
    merged: ``approval_status`` is a gate (did a human sign off?), ``decision``
    is which text wins (F column / an override / no change). Only rows a human
    approved cross over, and they carry their own ``after`` text, so no F column
    is required -- that is what lets AI/agent proposals use this tool unchanged.
    """
    decisions: list[dict] = []
    for index, raw in enumerate(changes):
        if not isinstance(raw, dict):
            raise ValueError(f"changes[{index}]는 객체여야 합니다.")
        status = str(raw.get("approval_status", "")).strip().lower()
        if status != APPROVED_STATUS:
            continue  # pending/rejected never reach Excel
        if "after" not in raw or not _text(raw["after"]):
            raise ValueError(f"changes[{index}]: 승인 항목에는 비어 있지 않은 after가 필요합니다.")
        if "before" not in raw:
            raise ValueError(f"changes[{index}]: 승인 항목에는 드리프트 검증용 before가 필요합니다.")
        decisions.append({
            "sheet": raw.get("sheet"),
            "cell": raw.get("cell"),
            "decision": "accept",
            "final_value": _text(raw["after"]),
            "expected_before": _text(raw["before"]),
            "finding_id": raw.get("finding_id", ""),
            "rule_ids": raw.get("rule_ids", []),
            "basis": ", ".join(raw.get("rule_ids", []) or []),
        })
    return decisions


def _load_manifest(path: Path) -> dict:
    """Accept either the native decisions manifest or a report-format one.

    Auto-detecting here rather than shipping a separate adapter script means
    there is no conversion step anyone can forget to run.
    """
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("manifest는 JSON 객체여야 합니다.")
    if isinstance(data.get("decisions"), list):
        return data
    if isinstance(data.get("changes"), list):
        converted = dict(data)
        converted["decisions"] = decisions_from_report_changes(data["changes"])
        if not converted["decisions"]:
            raise ValueError(
                "report manifest에 approval_status='approved' 항목이 없습니다. "
                "승인되지 않은 제안은 Excel에 반영하지 않습니다."
            )
        return converted
    raise ValueError(
        "manifest는 decisions 배열(판정) 또는 changes 배열(report_format_spec)을 가져야 합니다."
    )


def _parse_c_range(raw: str) -> tuple[int, int]:
    match = re.fullmatch(r"C(\d+):C(\d+)", raw.strip().upper())
    if not match:
        raise ValueError("--cell-range는 C7:C28처럼 C열의 연속 범위여야 합니다.")
    start, end = map(int, match.groups())
    if start > end:
        raise ValueError("--cell-range의 시작 행이 끝 행보다 큽니다.")
    return start, end


def _parse_columns(raw: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]+):([A-Z]+)", raw.strip().upper())
    if not match:
        raise ValueError("--drop-review-columns는 E:H처럼 열 범위여야 합니다.")
    start, end = (column_index_from_string(item) for item in match.groups())
    if start > end:
        raise ValueError("삭제할 열 범위가 올바르지 않습니다.")
    return start, end


def _source_for_sheet(sheet: str) -> str:
    return "KR(한국)" if sheet in ap.GROUP_A_TARGETS else "US(미국)"


def _validate_decisions(
    wb: openpyxl.Workbook,
    decisions: list[dict],
    protected: set[str],
    cell_range: str,
    source_sheet: str | None = None,
) -> list[dict]:
    start, end = _parse_c_range(cell_range)
    seen: set[tuple[str, str]] = set()
    normalized: list[dict] = []

    for index, raw in enumerate(decisions):
        if not isinstance(raw, dict):
            raise ValueError(f"decisions[{index}]는 객체여야 합니다.")
        sheet = str(raw.get("sheet", "")).strip()
        cell = str(raw.get("cell", "")).strip().upper()
        decision = str(raw.get("decision", "")).strip().lower()
        if sheet not in wb.sheetnames:
            raise ValueError(f"decisions[{index}]: 없는 시트 {sheet!r}")
        if decision not in DECISIONS:
            raise ValueError(f"decisions[{index}]: decision은 {sorted(DECISIONS)} 중 하나여야 합니다.")
        try:
            col, row = coordinate_from_string(cell)
        except ValueError as exc:
            raise ValueError(f"decisions[{index}]: 잘못된 셀 좌표 {cell!r}") from exc
        if col != "C" or not start <= row <= end:
            raise ValueError(f"decisions[{index}]: {cell}은 허용 범위 {cell_range} 밖입니다.")
        key = (sheet, cell)
        if key in seen:
            raise ValueError(f"같은 셀의 판정이 중복되었습니다: {sheet}!{cell}")
        seen.add(key)
        if decision in {"accept", "partial"} and sheet in protected:
            raise ValueError(f"보호 언어는 수용/부분 수용할 수 없습니다: {sheet}!{cell}")
        declared_source = raw.get("source_sheet")
        if declared_source and declared_source != (source_sheet or _source_for_sheet(sheet)):
            raise ValueError(
                f"{sheet}!{cell}: source_sheet={declared_source!r}가 표준 source group "
                f"{_source_for_sheet(sheet)!r}과 다릅니다."
            )
        normalized.append({**raw, "sheet": sheet, "cell": cell, "decision": decision,
                           "source_sheet": source_sheet or _source_for_sheet(sheet)})
    return normalized


def _logical_cells(ws, deleted: tuple[int, int] | None) -> dict[tuple[int, int], tuple[str, Any]]:
    """Snapshot values/formulas after logically removing a column interval."""
    values: dict[tuple[int, int], tuple[str, Any]] = {}
    start, end = deleted or (0, -1)
    width = end - start + 1
    for row in ws.iter_rows():
        for cell in row:
            if start <= cell.column <= end:
                continue
            logical_col = cell.column - width if cell.column > end else cell.column
            if cell.value is not None:
                values[(cell.row, logical_col)] = (cell.data_type, cell.value)
    return values


def _logical_merges(ws, deleted: tuple[int, int] | None) -> set[str]:
    start, end = deleted or (0, -1)
    width = end - start + 1
    result: set[str] = set()
    for merged in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged.bounds
        if not (max_col < start or min_col > end):
            continue
        if min_col > end:
            min_col -= width
            max_col -= width
        result.add(f"{min_col}:{min_row}:{max_col}:{max_row}")
    return result


def _snapshot_protected(wb, protected: set[str], deleted: tuple[int, int] | None = None) -> dict:
    return {
        name: {
            "cells": _logical_cells(wb[name], deleted),
            "merges": _logical_merges(wb[name], deleted),
            "c_values": [_text(wb[name][f"C{row}"].value) for row in range(1, wb[name].max_row + 1)],
        }
        for name in protected
    }


def _verify_protected(before: dict, after: dict) -> dict:
    failed: list[str] = []
    for sheet, snapshot in before.items():
        if snapshot["c_values"] != after[sheet]["c_values"]:
            failed.append(f"{sheet}: C열 값")
        if snapshot["cells"] != after[sheet]["cells"]:
            failed.append(f"{sheet}: 값/수식")
        if snapshot["merges"] != after[sheet]["merges"]:
            failed.append(f"{sheet}: 병합")
    if failed:
        raise RuntimeError("보호 언어 무결성 검증 실패: " + ", ".join(failed))
    return {"passed": True, "sheets": sorted(before)}


def _apply_decisions(wb, decisions: list[dict]) -> tuple[list[dict], list[dict]]:
    records: list[dict] = []
    changes: list[dict] = []
    for item in decisions:
        ws = wb[item["sheet"]]
        current = _text(ws[item["cell"]].value)
        reviewer_cell = str(item.get("reviewer_cell") or f"F{item['cell'][1:]}").upper()
        reviewer = _text(ws[reviewer_cell].value)
        comment_cell = str(item.get("comment_cell") or f"H{item['cell'][1:]}").upper()
        comment = _text(ws[comment_cell].value)
        decision = item["decision"]

        # Drift guard: the workbook may have moved on since the review was written.
        if "expected_before" in item:
            expected_before = _text(item["expected_before"])
            if expected_before != current:
                raise ValueError(
                    f"{item['sheet']}!{item['cell']}: 워크북이 검수 이후 변경되었습니다. "
                    f"expected_before={expected_before!r} != 현재값={current!r}"
                )

        if decision == "accept":
            # final_value present -> AI/agent proposal, no F column involved.
            # absent -> the original Excel-reviewer path, F column is the proposal.
            if "final_value" in item and _text(item["final_value"]):
                final = _text(item["final_value"])
                if final == current:
                    raise ValueError(
                        f"{item['sheet']}!{item['cell']}: 수용 final_value가 현재값과 같습니다."
                    )
            elif not reviewer or reviewer == current:
                raise ValueError(f"{item['sheet']}!{item['cell']}: 수용할 F열 수정안이 없거나 현재값과 같습니다.")
            else:
                final = reviewer
        elif decision == "partial":
            if "final_value" not in item or not _text(item["final_value"]):
                raise ValueError(f"{item['sheet']}!{item['cell']}: 부분 수용에는 final_value가 필요합니다.")
            final = _text(item["final_value"])
            if final == current:
                raise ValueError(f"{item['sheet']}!{item['cell']}: 부분 수용 final_value가 현재값과 같습니다.")
        else:
            final = current
        entry = {
            "sheet": item["sheet"], "cell": item["cell"], "decision": decision,
            "current": current, "reviewer_proposal": reviewer, "reviewer_comment": comment,
            "final": final, "reason": item.get("reason", ""), "basis": item.get("basis", ""),
            "rag_basis": item.get("rag_basis", ""), "source_sheet": item.get("source_sheet", _source_for_sheet(item["sheet"])),
            "finding_id": item.get("finding_id", ""), "rule_ids": item.get("rule_ids", []),
            "before_verified": "expected_before" in item,
        }
        if decision != "hold":
            ws[item["cell"]].value = final
            changes.append(entry)
        records.append(entry)
    return records, changes


def _verify_expected_c_diff(source: Path, acceptance: Path, expected: set[tuple[str, str]]) -> dict:
    before = openpyxl.load_workbook(source, data_only=False)
    after = openpyxl.load_workbook(acceptance, data_only=False)
    actual: set[tuple[str, str]] = set()
    for name in before.sheetnames:
        for row in range(1, max(before[name].max_row, after[name].max_row) + 1):
            if _text(before[name][f"C{row}"].value) != _text(after[name][f"C{row}"].value):
                actual.add((name, f"C{row}"))
    if actual != expected:
        raise RuntimeError(f"C열 값 diff 불일치: unexpected={sorted(actual - expected)}, missing={sorted(expected - actual)}")
    return {"expected": len(expected), "actual": len(actual), "passed": True}


async def run_apply(args) -> dict:
    from workbook_delivery import run_delivery
    return await run_delivery(args, "review")


def main() -> None:
    parser = argparse.ArgumentParser(description="승인된 감수 판정만 반영해 최종 하이라이트 납품본 생성")
    parser.add_argument("workbook", help="감수본 .xlsx (원본 불변)")
    parser.add_argument("approval_manifest", help="수용/부분 수용/유지 판정 JSON")
    parser.add_argument("--output", required=True, help="하이라이트 전 1차 수용본 .xlsx")
    parser.add_argument("--glossary", required=True, help="기준 glossary CSV")
    parser.add_argument("--cell-range", default="C7:C28")
    parser.add_argument("--protected-sheets", default=",".join(DEFAULT_PROTECTED))
    parser.add_argument("--drop-review-columns", help="명시 시 감수 열 삭제. 예: E:H")
    parser.add_argument("--result-manifest", help="생성 결과 manifest 경로")
    parser.add_argument("--app-root", help="SmartThings app repo")
    parser.add_argument("--max-concurrency", type=int, default=10)
    parser.add_argument("--dry-run", action="store_true", help="계약 preview만 저장; Excel은 쓰지 않음")
    parser.add_argument("--workflow-settings", help="이 작업의 고정 settings.json; 용어집·원문·활성화 설정 재사용")
    parser.add_argument("--source-sheet", help="작업별 기준 원문 예외")
    parser.add_argument("--sheet-langs", help="작업별 언어 매핑 JSON")
    parser.add_argument("--activation-manifest", help="승인된 story/cell/term 활성화 manifest")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    try:
        result = asyncio.run(run_apply(args))
    except Exception as exc:
        result = {"status": "error", "error": str(exc)}
    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else result)
    if result["status"] not in {"ok", "preview"}:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
