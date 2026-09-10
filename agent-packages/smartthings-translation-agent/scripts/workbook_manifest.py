#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Workbook baseline and revision manifests for safe Excel review work.

The ledger is deliberately adjacent to a workbook (``.st-history``) rather
than embedded in it.  It never changes the supplied workbook and records only
relative artifact names plus fingerprints, not absolute local paths.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

from workbook_mutation_guard import semantic_workbook_snapshot


SCHEMA_VERSION = 1


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _hash_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def text_sha256(value: Any) -> str:
    return _hash_bytes(("" if value is None else str(value)).encode("utf-8"))


def _history_root(workbook: Path) -> Path:
    return workbook.parent / ".st-history"


def _atomic_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, path)


def workbook_snapshot(workbook: str | Path) -> dict:
    """Create a semantic snapshot without retaining workbook text in the ledger."""
    path = Path(workbook)
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False, rich_text=True)
    try:
        legacy_structure = {
            "sheetnames": list(wb.sheetnames),
            "sheets": {
                ws.title: {
                    "state": ws.sheet_state,
                    "merged_ranges": sorted(str(value) for value in ws.merged_cells.ranges),
                    "protection": bool(ws.protection.sheet),
                    "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                }
                for ws in wb.worksheets
            },
        }
        legacy_values = [
            f"{ws.title}\0{cell.coordinate}\0{str(cell.value)}"
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        semantic = semantic_workbook_snapshot(wb)
        semantic_values_sha256 = semantic.pop("values_sha256")
        return {
            "file_sha256": file_sha256(path),
            # Preserve schema-v1 meanings so old and new ledgers remain
            # comparable. The expanded fingerprints use distinct field names.
            "structure_sha256": _hash_bytes(
                json.dumps(legacy_structure, ensure_ascii=False, sort_keys=True).encode("utf-8")
            ),
            "values_sha256": _hash_bytes("\n".join(legacy_values).encode("utf-8")),
            "semantic_values_sha256": semantic_values_sha256,
            **semantic,
        }
    finally:
        wb.close()


def ensure_baseline(workbook: str | Path) -> tuple[dict, Path]:
    """Return the immutable baseline for a supplied workbook, creating it once."""
    path = Path(workbook).expanduser().resolve()
    snapshot = workbook_snapshot(path)
    workbook_id = snapshot["file_sha256"][:20]
    baseline_path = _history_root(path) / workbook_id / "baseline.json"
    if baseline_path.exists():
        return json.loads(baseline_path.read_text(encoding="utf-8")), baseline_path
    payload = {
        "schema_version": SCHEMA_VERSION,
        "kind": "workbook_baseline",
        "workbook_id": workbook_id,
        "source_file": path.name,
        "created_at": _utc_now(),
        "snapshot": snapshot,
    }
    _atomic_json(baseline_path, payload)
    return payload, baseline_path


def _find_parent_revision(workbook: str | Path) -> tuple[dict, Path] | None:
    """Find a prior revision whose output is the supplied input workbook."""
    path = Path(workbook)
    current_sha = file_sha256(path)
    candidates = list(_history_root(path).glob("*/revisions/*.json"))
    # Versioned publication directories carry the existing change-log pointer.
    sidecar = path.with_suffix(".changes.json")
    if sidecar.is_file():
        try:
            pointer = json.loads(sidecar.read_text()).get("revision_manifest")
            if pointer:
                candidates.insert(0, Path(pointer))
        except (OSError, ValueError):
            pass
    for candidate in candidates:
        try:
            payload = load_revision(candidate)
        except (OSError, ValueError, json.JSONDecodeError):
            continue
        if payload.get("revised_snapshot", {}).get("file_sha256") == current_sha:
            return payload, candidate
    return None


def resolve_ledger(workbook: str | Path) -> tuple[dict, Path, dict | None]:
    """Resolve an existing parent revision, otherwise establish a fresh baseline."""
    parent = _find_parent_revision(workbook)
    if parent:
        parent_payload, parent_path = parent
        baseline_path = parent_path.parents[1] / "baseline.json"
        if baseline_path.is_file():
            return json.loads(baseline_path.read_text(encoding="utf-8")), baseline_path, parent_payload
    baseline, baseline_path = ensure_baseline(workbook)
    return baseline, baseline_path, None


def diff_spans(before: Any, after: Any) -> dict:
    """Return red review spans for a ``before -> after`` string change.

    Inserts/replacements colour their final characters.  A deletion has no
    remaining character, so its immediately adjacent words (or neighbouring
    characters for no-space scripts) point reviewers to the deletion site.
    """
    from difflib import SequenceMatcher

    old, new = "" if before is None else str(before), "" if after is None else str(after)
    if old == new:
        return {"after_text_sha256": text_sha256(new), "red_spans": [], "deleted_text": []}
    red: list[tuple[int, int]] = []
    deleted: list[str] = []
    has_word_spaces = any(char.isspace() for char in old + new)

    def deletion_context(position: int) -> None:
        if not new:
            return
        if not has_word_spaces:
            if position > 0:
                red.append((position - 1, position))
            if position < len(new):
                red.append((position, position + 1))
            return
        left = position - 1
        while left >= 0 and new[left].isspace():
            left -= 1
        if left >= 0:
            start = left
            while start > 0 and not new[start - 1].isspace():
                start -= 1
            red.append((start, left + 1))
        right = position
        while right < len(new) and new[right].isspace():
            right += 1
        if right < len(new):
            end = right
            while end < len(new) and not new[end].isspace():
                end += 1
            red.append((right, end))

    # Word/token comparison avoids character-level alignment turning a removed
    # ``on `` into an unhelpful fragment such as ``n o``. Scripts without word
    # spaces still use character comparison and the neighbouring-character rule.
    if has_word_spaces:
        old_tokens = re.findall(r"\s+|\w+|[^\w\s]", old, flags=re.UNICODE)
        new_tokens = re.findall(r"\s+|\w+|[^\w\s]", new, flags=re.UNICODE)
        old_offsets, new_offsets = [], []
        offset = 0
        for token in old_tokens:
            old_offsets.append(offset); offset += len(token)
        offset = 0
        for token in new_tokens:
            new_offsets.append(offset); offset += len(token)
        opcodes = SequenceMatcher(a=old_tokens, b=new_tokens, autojunk=False).get_opcodes()
        for tag, i1, i2, j1, j2 in opcodes:
            if tag in {"insert", "replace"} and j1 != j2:
                red.append((new_offsets[j1], new_offsets[j2 - 1] + len(new_tokens[j2 - 1])))
            if tag in {"delete", "replace"} and i1 != i2:
                deleted.append("".join(old_tokens[i1:i2]))
                if tag == "delete":
                    deletion_context(new_offsets[j1] if j1 < len(new_offsets) else len(new))
    else:
        for tag, i1, i2, j1, j2 in SequenceMatcher(a=old, b=new, autojunk=False).get_opcodes():
            if tag in {"insert", "replace"} and j1 != j2:
                red.append((j1, j2))
            if tag in {"delete", "replace"} and i1 != i2:
                deleted.append(old[i1:i2])
                if tag == "delete":
                    deletion_context(j1)

    merged: list[list[int]] = []
    for start, end in sorted({(max(0, start), min(len(new), end)) for start, end in red if start < end}):
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return {"after_text_sha256": text_sha256(new), "red_spans": merged, "deleted_text": deleted}


def create_edit_revision(source: str | Path, revised: str | Path, changes: list[dict]) -> tuple[dict, Path]:
    baseline, baseline_path, parent = resolve_ledger(source)
    source_path, revised_path = Path(source), Path(revised)
    revision_key = _hash_bytes((file_sha256(revised_path) + _utc_now()).encode("utf-8"))[:16]
    normalized = []
    for change in changes:
        before = change.get("before", change.get("old_value"))
        after = change.get("after", change.get("new_value"))
        normalized.append({
            "sheet": change["sheet"], "cell": change["cell"], "before": before,
            "after": after, "reason": change.get("reason"), "rule_ids": change.get("rule_ids", []),
            "diff": diff_spans(before, after),
        })
    payload = {
        "schema_version": SCHEMA_VERSION,
        "kind": "edit_revision",
        "revision_id": f"rev-{revision_key}",
        "workbook_id": baseline["workbook_id"],
        "parent_revision": parent.get("revision_id") if parent else None,
        "created_at": _utc_now(),
        "source_file": source_path.name,
        "revised_file": revised_path.name,
        "source_snapshot": workbook_snapshot(source_path),
        "revised_snapshot": workbook_snapshot(revised_path),
        "approval": "approved",
        "changes": normalized,
        "render_state": {},
    }
    path = baseline_path.parent / "revisions" / f"{payload['revision_id']}.json"
    _atomic_json(path, payload)
    return payload, path


def load_revision(path: str | Path) -> dict:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if payload.get("kind") != "edit_revision" or not isinstance(payload.get("changes"), list):
        raise ValueError("edit revision manifest가 아닙니다.")
    return payload


def update_render_state(path: str | Path, state: dict) -> dict:
    manifest_path = Path(path)
    payload = load_revision(manifest_path)
    payload["render_state"] = state
    payload["updated_at"] = _utc_now()
    _atomic_json(manifest_path, payload)
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel baseline/revision manifest 관리")
    sub = parser.add_subparsers(dest="command", required=True)
    init = sub.add_parser("init"); init.add_argument("workbook")
    status = sub.add_parser("status"); status.add_argument("workbook")
    args = parser.parse_args()
    baseline, path = ensure_baseline(args.workbook)
    if args.command == "status":
        payload = {"status": "ok", "baseline": baseline, "path": str(path)}
    else:
        payload = {"status": "ok", "operation": "init", "workbook_id": baseline["workbook_id"], "path": str(path)}
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
