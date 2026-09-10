#!/usr/bin/env python3
"""Canonical safety primitives for SmartThings one-off Excel mutation tools.

New workbook writers should import this module instead of reimplementing rich
text cloning, cross-workbook layout copying, physical row deletion, semantic
fingerprints, or verified atomic saves.  Extend this module with tests when a
new mutation needs a capability that is not yet represented here.
"""

from __future__ import annotations

import copy
import hashlib
import json
import os
import tempfile
from pathlib import Path
from typing import Any, Callable, Iterable

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock


def text(value: Any) -> str:
    return "" if value is None else str(value)


def _hash_json(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def clone_rich_value(value: Any) -> Any:
    """Clone a cell value without flattening rich-text runs."""
    if isinstance(value, CellRichText):
        result = CellRichText()
        for part in value:
            if isinstance(part, TextBlock):
                result.append(TextBlock(text=part.text or "", font=copy.copy(part.font)))
            else:
                result.append(str(part))
        return result
    return copy.copy(value)


def slice_rich_value(value: Any, start: int = 0, end: int | None = None) -> Any:
    """Slice a cell value by displayed-character offsets without flattening runs."""
    displayed = text(value)
    limit = len(displayed)
    normalized_start, normalized_end, _step = slice(start, end).indices(limit)
    if normalized_start >= normalized_end:
        return ""
    if not isinstance(value, CellRichText):
        return displayed[normalized_start:normalized_end]

    result = CellRichText()
    cursor = 0
    for part in value:
        segment = part.text or "" if isinstance(part, TextBlock) else str(part)
        segment_end = cursor + len(segment)
        overlap_start = max(normalized_start, cursor)
        overlap_end = min(normalized_end, segment_end)
        if overlap_start < overlap_end:
            fragment = segment[overlap_start - cursor:overlap_end - cursor]
            if isinstance(part, TextBlock):
                result.append(TextBlock(text=fragment, font=copy.copy(part.font)))
            else:
                result.append(fragment)
        cursor = segment_end
    return result


def rich_value_signature(value: Any) -> dict[str, Any]:
    """Return a stable signature for displayed text and rich-text run formatting."""
    displayed = text(value)
    return {
        "kind": "rich" if isinstance(value, CellRichText) else "plain",
        "text_sha256": hashlib.sha256(displayed.encode("utf-8")).hexdigest(),
        "runs": _rich_signature(value),
    }


def copy_cell_style(source, target) -> None:
    """Copy semantic style components safely across different workbooks."""
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)
    target.number_format = source.number_format


def copy_row_layout(
    source_sheet,
    target_sheet,
    row: int,
    *,
    columns: Iterable[int] | None = None,
    target_row: int | None = None,
) -> None:
    """Copy row geometry and cell styles without copying cell values."""
    destination_row = row if target_row is None else target_row
    source_dimension = source_sheet.row_dimensions.get(row)
    if source_dimension is None:
        if destination_row in target_sheet.row_dimensions:
            del target_sheet.row_dimensions[destination_row]
    else:
        target_dimension = target_sheet.row_dimensions[destination_row]
        for attribute in (
            "height", "hidden", "outlineLevel", "collapsed", "thickTop", "thickBot",
        ):
            setattr(target_dimension, attribute, copy.copy(getattr(source_dimension, attribute)))
    selected = columns if columns is not None else range(1, source_sheet.max_column + 1)
    for column in selected:
        copy_cell_style(source_sheet.cell(row, column), target_sheet.cell(destination_row, column))


def delete_rows_with_manifest(sheet, start: int, amount: int) -> dict[str, Any]:
    """Physically delete rows and return every removed non-empty cell for audit."""
    if start < 1 or amount < 0:
        raise ValueError("행 삭제 범위가 올바르지 않습니다")
    if amount == 0:
        return {"start": start, "end": start - 1, "rows": [], "values": {}}
    end = start + amount - 1
    values = {
        cell.coordinate: text(cell.value)
        for row in sheet.iter_rows(min_row=start, max_row=end)
        for cell in row
        if cell.value is not None
    }
    styled_cells = {
        cell.coordinate: _hash_json(_style_signature(cell))
        for row in sheet.iter_rows(min_row=start, max_row=end)
        for cell in row
        if cell.has_style
    }
    deleted_dimensions = {
        str(index): _dimension_signature(dimension)
        for index, dimension in sheet.row_dimensions.items()
        if start <= index <= end
    }

    relocated_merges: list[tuple[str, str | None]] = []
    for merged in list(sheet.merged_cells.ranges):
        if merged.max_row < start:
            continue
        if merged.min_row > end:
            old = str(merged)
            new = (
                f"{sheet.cell(merged.min_row - amount, merged.min_col).coordinate}:"
                f"{sheet.cell(merged.max_row - amount, merged.max_col).coordinate}"
            )
            relocated_merges.append((old, new))
            continue
        if merged.min_row >= start and merged.max_row <= end:
            relocated_merges.append((str(merged), None))
            continue
        raise ValueError(f"삭제 행 경계를 가로지르는 병합 셀이 있습니다: {sheet.title}!{merged}")

    sheet.delete_rows(start, amount)

    # openpyxl moves cells but not RowDimension objects or merged ranges.
    dimensions = [(index, copy.copy(value)) for index, value in sheet.row_dimensions.items()]
    for index in list(sheet.row_dimensions):
        del sheet.row_dimensions[index]
    for index, dimension in dimensions:
        if start <= index <= end:
            continue
        new_index = index - amount if index > end else index
        dimension.index = new_index
        sheet.row_dimensions[new_index] = dimension

    for old, new in relocated_merges:
        sheet.unmerge_cells(old)
        if new is not None:
            sheet.merge_cells(new)

    return {
        "start": start,
        "end": end,
        "rows": list(range(start, end + 1)),
        "values": values,
        "styled_cells": styled_cells,
        "row_dimensions": deleted_dimensions,
        "merged_ranges": [{"before": old, "after": new} for old, new in relocated_merges],
    }


def _style_signature(cell) -> dict[str, Any]:
    return {
        "font": str(cell.font),
        "fill": str(cell.fill),
        "border": str(cell.border),
        "alignment": str(cell.alignment),
        "protection": str(cell.protection),
        "number_format": cell.number_format,
    }


def _rich_signature(value: Any) -> list[dict[str, str]] | None:
    if not isinstance(value, CellRichText):
        return None
    result: list[dict[str, str]] = []
    for part in value:
        if isinstance(part, TextBlock):
            result.append({
                "text_sha256": hashlib.sha256((part.text or "").encode("utf-8")).hexdigest(),
                "font": str(part.font),
            })
        else:
            result.append({
                "text_sha256": hashlib.sha256(str(part).encode("utf-8")).hexdigest(),
                "font": "",
            })
    return result


def _dimension_signature(dimension) -> dict[str, Any]:
    return {
        "height": getattr(dimension, "height", None),
        "width": getattr(dimension, "width", None),
        "hidden": getattr(dimension, "hidden", False),
        "outlineLevel": getattr(dimension, "outlineLevel", 0),
        "collapsed": getattr(dimension, "collapsed", False),
        "thickTop": getattr(dimension, "thickTop", False),
        "thickBot": getattr(dimension, "thickBot", False),
    }


def semantic_workbook_snapshot(workbook) -> dict[str, Any]:
    """Return independent hashes for values, layout, styles, and rich text."""
    values: list[dict[str, str]] = []
    styles: list[dict[str, Any]] = []
    rich_text: list[dict[str, Any]] = []
    annotations: list[dict[str, str]] = []
    sheets: dict[str, Any] = {}

    for sheet in workbook.worksheets:
        sheets[sheet.title] = {
            "state": sheet.sheet_state,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_ranges": sorted(str(value) for value in sheet.merged_cells.ranges),
            "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "auto_filter": sheet.auto_filter.ref,
            "tab_color": str(sheet.sheet_properties.tabColor) if sheet.sheet_properties.tabColor else None,
            "row_dimensions": {
                str(index): _dimension_signature(dimension)
                for index, dimension in sorted(sheet.row_dimensions.items())
            },
            "column_dimensions": {
                str(index): _dimension_signature(dimension)
                for index, dimension in sorted(sheet.column_dimensions.items())
            },
            "data_validations": sorted(str(item.sqref) for item in sheet.data_validations.dataValidation),
            "image_count": len(getattr(sheet, "_images", [])),
            "chart_count": len(getattr(sheet, "_charts", [])),
        }
        for (_row, _column), cell in sorted(sheet._cells.items()):
            coordinate = cell.coordinate
            if cell.value is not None:
                values.append({
                    "sheet": sheet.title,
                    "cell": coordinate,
                    "value_sha256": hashlib.sha256(text(cell.value).encode("utf-8")).hexdigest(),
                    "data_type": cell.data_type,
                })
            if cell.has_style:
                styles.append({"sheet": sheet.title, "cell": coordinate, **_style_signature(cell)})
            rich = _rich_signature(cell.value)
            if rich is not None:
                rich_text.append({"sheet": sheet.title, "cell": coordinate, "runs": rich})
            if cell.comment is not None:
                annotations.append({
                    "sheet": sheet.title,
                    "cell": coordinate,
                    "comment_sha256": hashlib.sha256(cell.comment.text.encode("utf-8")).hexdigest(),
                    "author_sha256": hashlib.sha256((cell.comment.author or "").encode("utf-8")).hexdigest(),
                })

    return {
        "sheet_count": len(workbook.sheetnames),
        "values_sha256": _hash_json(values),
        "layout_sha256": _hash_json({"sheetnames": list(workbook.sheetnames), "sheets": sheets}),
        "styles_sha256": _hash_json(styles),
        "rich_text_sha256": _hash_json(rich_text),
        "annotations_sha256": _hash_json(annotations),
    }


def snapshot_path(path: str | Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=False, rich_text=True)
    try:
        return semantic_workbook_snapshot(workbook)
    finally:
        workbook.close()


def snapshot_axis_diff(before: dict[str, Any], after: dict[str, Any]) -> list[str]:
    axes = ("values_sha256", "layout_sha256", "styles_sha256", "rich_text_sha256", "annotations_sha256")
    return [axis.removesuffix("_sha256") for axis in axes if before.get(axis) != after.get(axis)]


def delete_cols_with_manifest(sheet, start: int, amount: int) -> dict[str, Any]:
    """Delete physical columns with audit; reject dependencies we cannot rebase safely.

    Formula/chart/table/validation dependencies are intentionally not rewritten.
    Callers must first resolve these in a separate approved structural plan.
    """
    from openpyxl.utils import get_column_letter, column_index_from_string
    if start < 1 or amount < 0:
        raise ValueError("Invalid column deletion range")
    if amount == 0:
        return {"columns": [], "values": {}, "annotations": {}}
    if (sheet._charts or sheet._images or sheet.tables or sheet.data_validations.count
            or sheet.conditional_formatting or sheet.freeze_panes or sheet.auto_filter.ref
            or sheet.parent.defined_names
            or any(ws._charts or ws.tables or ws.data_validations.count or ws.conditional_formatting
                   or ws.defined_names for ws in sheet.parent)
            or any(cell.data_type == "f" for ws in sheet.parent for cell in ws._cells.values())):
        raise ValueError("Column deletion requires an explicit dependency-rebasing plan")
    end = start + amount - 1
    merges = []
    for merged in sheet.merged_cells.ranges:
        if merged.min_col < start <= merged.max_col or merged.min_col <= end < merged.max_col:
            raise ValueError("Merged range crosses deletion boundary")
        if merged.max_col < start:
            merges.append(str(merged))
        elif merged.min_col > end:
            merges.append(f"{get_column_letter(merged.min_col-amount)}{merged.min_row}:"
                          f"{get_column_letter(merged.max_col-amount)}{merged.max_row}")
    dimensions = []
    for key, dimension in sheet.column_dimensions.items():
        lo = dimension.min or column_index_from_string(key)
        hi = dimension.max or lo
        if lo < start <= hi or lo <= end < hi:
            raise ValueError("Column dimension group crosses deletion boundary")
        dimensions.append((lo, hi, copy.copy(dimension)))
    removed = [cell for (row, col), cell in sheet._cells.items() if start <= col <= end]
    manifest = {
        "columns": list(range(start, end+1)),
        "values": {cell.coordinate: text(cell.value) for cell in removed if cell.value is not None},
        "styles": {cell.coordinate: _style_signature(cell) for cell in removed if cell.has_style},
        "annotations": {cell.coordinate: {"text": cell.comment.text, "author": cell.comment.author}
                        for cell in removed if cell.comment},
        "column_dimensions": {str(lo): _dimension_signature(dim) for lo, hi, dim in dimensions
                              if start <= lo <= end},
        "merged_ranges": {"before": sorted(str(x) for x in sheet.merged_cells.ranges), "after": sorted(merges)},
    }
    # Remove merge metadata before moving cells, so unmerge cannot delete relocated cells.
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    sheet.delete_cols(start, amount)
    sheet.column_dimensions.clear()
    for lo, hi, dimension in dimensions:
        if start <= lo <= end:
            continue
        shift = amount if lo > end else 0
        dimension.min, dimension.max = lo-shift, hi-shift
        dimension.index = get_column_letter(lo-shift)
        sheet.column_dimensions[dimension.index] = dimension
    for merged in merges:
        sheet.merge_cells(merged)
    return manifest


def detailed_workbook_snapshot(workbook) -> dict[str, dict[str, Any]]:
    """Property-addressable five-axis facts, shared by independent disk verification.

    XML-bearing objects use full semantic XML rather than counts or style-table IDs.
    Binary drawing payloads and unsupported OOXML parts are handled by the disk verifier.
    """
    from openpyxl.xml.functions import tostring
    axes = {name: {} for name in ("values", "layout", "styles", "rich_text", "annotations")}
    def put(axis, path, value):
        axes[axis][json.dumps(path, ensure_ascii=False, separators=(",", ":"))] = value
    def xml(obj):
        return tostring(obj.to_tree()).decode() if obj is not None else None
    put("layout", ["sheetnames"], workbook.sheetnames)
    put("layout", ["defined_names"], sorted(xml(x) for x in workbook.defined_names.values()))
    put("layout", ["calculation"], xml(workbook.calculation))
    for sheet in workbook:
        base = ["sheets", sheet.title]
        for key, value in {
            "state": sheet.sheet_state, "max_row": sheet.max_row, "max_column": sheet.max_column,
            "merged_ranges": sorted(str(x) for x in sheet.merged_cells.ranges),
            "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "sheet_properties": xml(sheet.sheet_properties), "sheet_format": xml(sheet.sheet_format),
            "protection": xml(sheet.protection), "auto_filter": xml(sheet.auto_filter),
            "views": xml(sheet.views), "page_margins": xml(sheet.page_margins),
            "header_footer": xml(sheet.HeaderFooter),
            "defined_names": sorted(xml(x) for x in sheet.defined_names.values()),
            "page_setup": xml(sheet.page_setup), "print_options": xml(sheet.print_options),
            "print_area": str(sheet.print_area), "print_title_rows": sheet.print_title_rows,
            "print_title_cols": sheet.print_title_cols, "row_breaks": xml(sheet.row_breaks),
            "col_breaks": xml(sheet.col_breaks),
            "data_validations": xml(sheet.data_validations),
            "tables": sorted(xml(t) for t in sheet.tables.values()),
            "conditional_formatting": sorted((str(cf.sqref), [xml(rule) for rule in rules])
                                             for cf, rules in sheet.conditional_formatting._cf_rules.items()),
        }.items():
            put("layout", base + [key], value)
        for axis_name, dimensions in (("row_dimensions", sheet.row_dimensions), ("column_dimensions", sheet.column_dimensions)):
            for index, dimension in sorted(dimensions.items()):
                path = base + [axis_name, str(index)]
                for key, value in _dimension_signature(dimension).items():
                    put("layout", path + [key], value)
                for key in ("min", "max", "bestFit"):
                    if hasattr(dimension, key):
                        put("layout", path + [key], getattr(dimension, key))
                if dimension.has_style:
                    for key, value in _style_signature(dimension).items():
                        put("styles", path + [key], value)
        for (row, col), cell in sorted(sheet._cells.items()):
            path = base + ["cells", cell.coordinate]
            if cell.value is not None:
                put("values", path, {"text": text(cell.value), "data_type": cell.data_type})
            if cell.has_style:
                for key, value in _style_signature(cell).items():
                    put("styles", path + [key], value)
            rich = _rich_signature(cell.value)
            if rich is not None:
                put("rich_text", path, rich)
            if cell.comment:
                put("annotations", path + ["comment"], {"text": cell.comment.text, "author": cell.comment.author})
            if cell.hyperlink:
                put("annotations", path + ["hyperlink"], {key: getattr(cell.hyperlink, key)
                    for key in ("target", "location", "tooltip", "display")})
    return axes


def save_verified_atomic(
    workbook,
    output: str | Path,
    verify_path: Callable[[Path], Any],
    *,
    overwrite: bool = False,
    contract: dict | None = None,
) -> Any:
    """Save to a temp file, verify the reopened temp, then atomically promote."""
    destination = Path(output)
    if contract is None:
        import warnings
        warnings.warn("Contract-less workbook save is legacy; migrate to a declared contract",
                      DeprecationWarning, stacklevel=2)
    if contract is not None:
        from workbook_contract import validate_contract
        validate_contract(contract)
        root = Path(contract["staging_root"]).resolve()
        if not destination.resolve().is_relative_to(root) or destination.resolve() == root:
            raise ValueError("Contract saves are restricted to staging")
    if destination.exists() and not overwrite:
        raise FileExistsError(f"출력 파일이 이미 존재합니다: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=destination.parent,
        prefix=f".{destination.stem}.verify-",
        suffix=destination.suffix,
        delete=False,
    ) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        result = verify_path(temporary)
        if contract is not None:
            from workbook_contract import file_sha256, check_inputs, digest
            if (not isinstance(result, dict) or result.get("status") != "passed"
                    or result.get("staged_file_sha256") != file_sha256(temporary)
                    or result.get("contract_content_sha256") != digest(contract)
                    or set(result.get("checks", {})) != {"values", "layout", "styles", "rich_text", "annotations"}
                    or any(value != "passed" for value in result["checks"].values())):
                raise ValueError("Contract save requires a complete verification record")
            check_inputs(contract)
        os.replace(temporary, destination)
        return result
    except BaseException:
        if temporary.exists():
            temporary.unlink()
        raise
