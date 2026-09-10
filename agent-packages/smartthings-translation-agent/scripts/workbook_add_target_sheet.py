#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
workbook_add_target_sheet.py — 워크북에 새 언어 타겟 시트(+선택적 역번역 시트)를
추가한 사본을 생성한다 (크레딧 0).

원본은 절대 수정하지 않는다. 같은 워크북 안의 기존 시트(--template-sheet로 명시)를
복제해 새 시트(--new-sheet, 예: "CO(콜롬비아)")를 만든다 — 스토리마다 section 개수가 달라 행
구조가 파일마다 다르므로, 반드시 같은 파일 안의 시트를 복제해야 B열 라벨 수식과 행 구조가
그 스토리와 정확히 일치한다.

멱등적이다: 대상 시트가 이미 있으면 그 시트는 건드리지 않고 성공으로 보고한다(재실행 안전).
--new-sheet, --backtranslation-sheet 각각 독립적으로 존재 여부를 확인한다.

동작:
  1. --template-sheet 를 복제해 --new-sheet 로 이름 변경, --template-sheet 바로 뒤로 이동.
     //language 셀(행 3, C열)을 --lang-code 로 설정하고, 번역 대상 텍스트 셀(C7:C28, "x"
     placeholder·헤더 행 제외)을 비운다 — 파이프라인이 실패해도 복제 원본 언어의 텍스트가
     새 시트에 잘못 남아있는 상태를 방지한다(앱은 번역 시 어차피 무조건 덮어쓰지만, 실패 시
     빈 셀이 훨씬 눈에 잘 띄는 안전장치).
  2. --backtranslation-sheet 를 지정하면, --new-sheet 바로 뒤에 같은 방식으로 시트를 하나 더
     만든다(역번역 결과 전용 — 번역 파이프라인이 --sheets 타겟으로 취급하지 않도록 별도 이름).
     이 시트의 C열에 역번역 텍스트가 채워진다(checker_service.py 의
     TranslationChecker(backtranslation_sheet=...) 참고).

사용 예:
  python workbook_add_target_sheet.py story.xlsx --template-sheet "ES(스페인)" --new-sheet "CO(콜롬비아)" --lang-code es_CO \
      --backtranslation-sheet "CO(콜롬비아) 역번역" --out story_co_prep.xlsx
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
import os
import sys
from pathlib import Path

import openpyxl

# 파싱 위치 상수: workbook_inspect.py 와 동일 규약(rag_db_builder.py 기준, 무거운 의존성 없이도
# 동작해야 하므로 fallback 유지). ⚠ 값 변경 시 workbook_inspect.py 와 동일하게 유지할 것.
try:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import bootstrap as _bs
    _app_root, _src = _bs.resolve_app_root(_bs.cli_app_root_from_argv())
    if _app_root:
        _src_dir = Path(_app_root) / "src"
        if _src_dir.is_dir() and str(_src_dir) not in sys.path:
            sys.path.insert(0, str(_src_dir))
    from translation_web_app.rag_db_builder import (
        STORY_ID_CELL,
        CONTENT_ROW_START,
        CONTENT_ROW_END,
        CONTENT_COL,
    )
except Exception:
    STORY_ID_CELL = "C5"
    CONTENT_ROW_START = 7
    CONTENT_ROW_END = 28
    CONTENT_COL = 3  # C열

LANGUAGE_ROW = 3      # "//language" 라벨 행


def _clear_content(ws) -> list[str]:
    cleared = []
    for row in range(CONTENT_ROW_START, CONTENT_ROW_END + 1):
        cell = ws._cells.get((row, CONTENT_COL))
        if cell is None:
            continue
        val = str(cell.value).strip() if cell.value is not None else ""
        if val and val.lower() != "x":
            cell.value = None
            cleared.append(cell.coordinate)
    return cleared


def _clone_sheet(wb, template_sheet: str, title: str, insert_after: str):
    _check_copyable(wb[template_sheet])
    new_ws = wb.copy_worksheet(wb[template_sheet])
    # copy_worksheet omits views; retain zoom and selection explicitly.
    new_ws.views = deepcopy(wb[template_sheet].views)
    new_ws.sheet_state = wb[template_sheet].sheet_state
    new_ws.title = title
    after_idx = wb.sheetnames.index(insert_after)
    wb.move_sheet(new_ws, offset=(after_idx + 1) - wb.sheetnames.index(title))
    return new_ws


def add_target_sheet(
    workbook: Path,
    template_sheet: str,
    new_sheet: str,
    lang_code: str,
    backtranslation_sheet: str | None = None,
) -> dict:
    wb = openpyxl.load_workbook(workbook, rich_text=True)

    if template_sheet not in wb.sheetnames:
        return {"wb": wb, "status": "error", "reason": f"템플릿 시트 없음: '{template_sheet}'"}

    story_id = wb[template_sheet][STORY_ID_CELL].value
    result = {
        "wb": wb,
        "status": "ok",
        "story_id": story_id,
        "new_sheet_action": None,
        "backtranslation_sheet_action": None,
        "cleared_cells": [],
    }

    if new_sheet in wb.sheetnames:
        result["new_sheet_action"] = "skipped (이미 존재함)"
    else:
        new_ws = _clone_sheet(wb, template_sheet, new_sheet, insert_after=template_sheet)
        new_ws.cell(row=LANGUAGE_ROW, column=CONTENT_COL).value = lang_code
        result["cleared_cells"] = _clear_content(new_ws)
        result["new_sheet_action"] = "created"

    if backtranslation_sheet:
        if backtranslation_sheet in wb.sheetnames:
            result["backtranslation_sheet_action"] = "skipped (이미 존재함)"
        else:
            bt_ws = _clone_sheet(wb, template_sheet, backtranslation_sheet, insert_after=new_sheet)
            bt_ws.cell(row=LANGUAGE_ROW, column=CONTENT_COL).value = (
                f"{lang_code} 역번역 참고용 (배포 대상 아님)"
            )
            _clear_content(bt_ws)
            result["backtranslation_sheet_action"] = "created"

    result["position"] = wb.sheetnames.index(new_sheet)
    return result


def _check_copyable(ws):
    """Reject objects copy_worksheet cannot preserve before performing a copy."""
    from openpyxl.worksheet.worksheet import Worksheet
    from openpyxl.xml.functions import tostring
    blank = Worksheet(ws.parent)
    unsupported = ("protection", "auto_filter", "HeaderFooter", "row_breaks", "col_breaks")
    if (ws._charts or ws._images or ws.tables or ws.data_validations.count
            or ws.conditional_formatting or ws.freeze_panes or ws.defined_names
            or ws.print_area or ws.print_title_rows or ws.print_title_cols
            or any(c.comment for c in ws._cells.values())
            or any(tostring(getattr(ws, k).to_tree()) != tostring(getattr(blank, k).to_tree())
                   for k in unsupported)):
        raise ValueError("Template contains objects copy_worksheet cannot preserve; use an explicit layout plan")


def save_prepared(workbook: Path, output: Path, template_sheet: str, new_sheet: str,
                  lang_code: str, backtranslation_sheet: str | None = None) -> dict:
    """All preparation callers use the same source-derived five-axis contract."""
    import copy
    from workbook_contract import atomic_json, canonical, capture_authority, digest, file_sha256, save_contract
    from workbook_run import execute_run
    from workbook_verifier import facts, diff_facts
    # Import the shared writer explicitly; execute_run owns its invocation.
    from workbook_mutation_guard import save_verified_atomic
    workbook, output = workbook.resolve(), output.resolve()
    if workbook == output:
        raise ValueError("Preparation output must differ from source")
    before = facts(workbook)
    expected = copy.deepcopy(before)
    names = expected["layout"][canonical(["sheetnames"])]
    additions = [(new_sheet, lang_code, template_sheet)]
    if backtranslation_sheet:
        additions.append((backtranslation_sheet, f"{lang_code} 역번역 참고용 (배포 대상 아님)", new_sheet))
    for title, code, after in additions:
        if title in names:
            continue
        names.insert(names.index(after) + 1, title)
        for axis, entries in before.items():
            for key, value in entries.items():
                path = json.loads(key)
                if path[:2] == ["sheets", template_sheet]:
                    expected[axis][canonical(["sheets", title, *path[2:]])] = copy.deepcopy(value)
        for row in range(CONTENT_ROW_START, CONTENT_ROW_END + 1):
            key = canonical(["sheets", title, "cells", f"C{row}"])
            value = expected["values"].get(key)
            if value and value["text"].strip() and value["text"].strip().lower() != "x":
                expected["values"].pop(key)
                expected["rich_text"].pop(key, None)
        key = canonical(["sheets", title, "cells", "C3"])
        expected["values"][key] = {"text": code, "data_type": "s"}
        expected["rich_text"].pop(key, None)
        for dim, minimum in (("max_row", 3), ("max_column", 3)):
            key = canonical(["sheets", title, dim])
            expected["layout"][key] = max(expected["layout"][key], minimum)
    result = add_target_sheet(workbook, template_sheet, new_sheet, lang_code, backtranslation_sheet)
    if result["status"] != "ok":
        result["wb"].close()
        raise ValueError(result["reason"])
    changes = diff_facts(before, expected)
    work_id = "prepare-" + digest([str(workbook), file_sha256(workbook), str(output), additions])[:24]
    root = output.parent / ".st-runs" / work_id
    ref = {"authority": "source", "file": "."}
    writer = {"id": "workbook_prepare", "version": file_sha256(__file__), "cost": "local"}
    contract = {"schema_version": 1, "work_id": work_id, "work_root": str(root),
        "staging_root": str(root / "staging"), "delivery_root": str(output.parent / "verified"),
        "authorities": {"source": capture_authority(workbook)}, "writer": writer,
        "artifacts": [{"id": "prepared", "output": output.name,
            "authorities": {role: dict(ref) for role in ("values", "structure", "formatting")},
            "allowed_diffs": changes}], "no_op_files": [] if changes else ["prepared"],
        "recovery": {"max_attempts": 1, "same_failure_limit": 1}}
    path, approval = root / "contract.json", root / "approval.json"
    if not path.exists():
        sha = save_contract(path, contract)
        atomic_json(approval, {"approved": True, "approved_by": "authorized_workbook_preparation",
                              "contract_sha256": sha})
    try:
        state = execute_run(path, lambda *_: result["wb"], writer=writer, approval_path=approval)
        if state["status"] != "completed":
            raise ValueError(f"Workbook preparation failed verification: {root / 'state.json'}")
        return {**{k: v for k, v in result.items() if k != "wb"},
                "output": state["result"]["outputs"]["prepared"], "contract": str(path)}
    finally:
        result["wb"].close()


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("workbook", help="원본 .xlsx 경로 (수정되지 않음)")
    p.add_argument("--template-sheet", required=True, help="복제할 기존 시트명")
    p.add_argument("--new-sheet", required=True, help="새로 만들 시트명 (예: 'CO(콜롬비아)')")
    p.add_argument("--lang-code", required=True, help="//language 셀에 넣을 값 (예: 'es_CO')")
    p.add_argument("--backtranslation-sheet",
                    help="역번역 결과를 담을 별도 시트명 (예: 'CO(콜롬비아) 역번역'). "
                         "미지정 시 생성하지 않음")
    p.add_argument("--out", required=True, help="결과를 저장할 경로")
    p.add_argument("--json", action="store_true")
    args = p.parse_args()

    workbook = Path(args.workbook).expanduser()
    if not workbook.is_file():
        msg = {"status": "error", "error": f"워크북을 찾을 수 없습니다: {workbook}"}
        print(json.dumps(msg, ensure_ascii=False) if args.json else f"❌ {msg['error']}")
        sys.exit(2)

    summary = save_prepared(workbook, Path(args.out).expanduser(), args.template_sheet,
                            args.new_sheet, args.lang_code, args.backtranslation_sheet)
    out_path = Path(summary["output"])
    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        print(f"✅ 완료 → {out_path}")
        print(f"   story_id={summary.get('story_id')}")
        print(f"   '{args.new_sheet}': {summary['new_sheet_action']}")
        if args.backtranslation_sheet:
            print(f"   '{args.backtranslation_sheet}': {summary['backtranslation_sheet_action']}")


if __name__ == "__main__":
    main()
